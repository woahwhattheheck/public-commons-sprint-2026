"""Native editable Office charts and matching vector PDF plots.

Only the retained disclosure-safe report view enters this module. Every chart
panel embeds its complete category and numeric data in an editable XLSX part.
"""
from __future__ import annotations

import io
import math
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal


NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)
PALETTE = ["404040", "808080", "B8B8B8", "5C5C5C", "9A9A9A"]


def element(parent, tag_name, text=None, **attributes):
    prefix, local = tag_name.split(":")
    item = ET.SubElement(parent, "{" + NS[prefix] + "}" + local, {k.replace("_", ":"):str(v) for k,v in attributes.items()})
    if text is not None:
        item.text = str(text)
    return item


def _xml(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _column(number):
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def chart_panels(spec, categories_per_panel=12):
    categories = spec.get("categories", [])
    series = spec.get("series", [])
    if spec.get("kind") not in {"bar", "line"} or not categories or not series:
        return []
    for item in series:
        if len(item.get("values", [])) != len(categories):
            raise ValueError("Chart categories and retained value counts disagree")
        for value in item["values"]:
            if value is not None and not math.isfinite(float(Decimal(str(value)))):
                raise ValueError("Chart values must be finite numbers")
    result = []
    for series_start in range(0,len(series),4):
        series_end = min(len(series),series_start+4)
        for start in range(0, len(categories), categories_per_panel):
            end = min(len(categories), start + categories_per_panel)
            suffix = "" if len(categories) <= categories_per_panel else "  Rows " + str(start + 1) + " to " + str(end)
            if len(series)>4: suffix += "  Group " + str(series_end) if series_end == series_start+1 else "  Groups " + str(series_start+1) + " to " + str(series_end)
            result.append({
                "kind":spec["kind"], "title":str(spec.get("title", "Report view")) + suffix,
                "categories":[str(x) for x in categories[start:end]],
                "series":[{"name":str(x.get("name", "Values")), "values":x["values"][start:end]} for x in series[series_start:series_end]],
            })
    return result


def _workbook(spec):
    workbook = ET.Element("{" + NS["s"] + "}workbook")
    sheets = element(workbook, "s:sheets")
    sheet = element(sheets, "s:sheet", name="Data", sheetId="1")
    sheet.set("{" + NS["r"] + "}id", "rId1")
    worksheet = ET.Element("{" + NS["s"] + "}worksheet")
    data = element(worksheet, "s:sheetData")
    rows = [["Category"] + [x["name"] for x in spec["series"]]]
    rows += [[label] + [x["values"][index] for x in spec["series"]] for index,label in enumerate(spec["categories"])]
    for row_index, values in enumerate(rows, 1):
        row = element(data, "s:row", r=row_index)
        for column_index, value in enumerate(values, 1):
            if value is None:
                continue
            cell = element(row, "s:c", r=_column(column_index) + str(row_index))
            if row_index == 1 or column_index == 1:
                cell.set("t", "inlineStr")
                string = element(cell, "s:is")
                element(string, "s:t", value)
            else:
                element(cell, "s:v", str(value))
    parts = {
        "[Content_Types].xml":b'<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
        "_rels/.rels":b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/_rels/workbook.xml.rels":b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/workbook.xml":_xml(workbook), "xl/worksheets/sheet1.xml":_xml(worksheet),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, value in sorted(parts.items()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, value)
    return output.getvalue()


def _fill(parent, color):
    fill = element(parent, "a:solidFill")
    element(fill, "a:srgbClr", val=color)


def _line(parent, color, width=9525):
    line = element(parent, "a:ln", w=width)
    _fill(line, color)


def _reference(parent, tag, formula, values, numeric=False):
    holder = element(parent, tag)
    reference = element(holder, "c:numRef" if numeric else "c:strRef")
    element(reference, "c:f", formula)
    cache = element(reference, "c:numCache" if numeric else "c:strCache")
    if numeric:
        element(cache, "c:formatCode", "General")
    element(cache, "c:ptCount", val=len(values))
    for index, value in enumerate(values):
        if value is not None:
            point = element(cache, "c:pt", idx=index)
            element(point, "c:v", str(value))


def _chart_text(parent):
    properties = element(parent, "c:txPr")
    element(properties, "a:bodyPr")
    element(properties, "a:lstStyle")
    paragraph = element(properties, "a:p")
    default = element(element(paragraph, "a:pPr"), "a:defRPr", sz="900")
    _fill(default, "000000")
    element(default, "a:latin", typeface="DejaVu Sans")
    element(paragraph, "a:endParaRPr", lang="en-CA")


def _chart_xml(spec, workbook_relationship):
    root = ET.Element("{" + NS["c"] + "}chartSpace")
    element(root, "c:date1904", val="0")
    element(root, "c:lang", val="en-CA")
    chart = element(root, "c:chart")
    element(chart, "c:autoTitleDeleted", val="1")
    plot = element(chart, "c:plotArea")
    element(plot, "c:layout")
    line_chart = spec["kind"] == "line"
    series_group = element(plot, "c:lineChart" if line_chart else "c:barChart")
    if not line_chart:
        element(series_group, "c:barDir", val="col")
    element(series_group, "c:grouping", val="standard" if line_chart else "clustered")
    element(series_group, "c:varyColors", val="0")
    count = len(spec["categories"])
    for index, series in enumerate(spec["series"]):
        item = element(series_group, "c:ser")
        element(item, "c:idx", val=index)
        element(item, "c:order", val=index)
        column = _column(index + 2)
        _reference(item, "c:tx", "'Data'!$" + column + "$1", [series["name"]])
        properties = element(item, "c:spPr")
        color = PALETTE[index % len(PALETTE)]
        if line_chart:
            _line(properties, color, 25400)
            marker = element(item, "c:marker")
            element(marker, "c:symbol", val=["circle", "square", "diamond", "triangle", "x"][index % 5])
            element(marker, "c:size", val="5")
        else:
            _fill(properties, color)
            _line(properties, "FFFFFF", 3810)
            element(item, "c:invertIfNegative", val="0")
        _reference(item, "c:cat", "'Data'!$A$2:$A$" + str(count + 1), spec["categories"])
        _reference(item, "c:val", "'Data'!$" + column + "$2:$" + column + "$" + str(count + 1), series["values"], numeric=True)
        if line_chart:
            element(item, "c:smooth", val="0")
    if line_chart:
        element(series_group, "c:marker", val="1")
        element(series_group, "c:smooth", val="0")
    else:
        element(series_group, "c:gapWidth", val="80")
    for axis_id in (310001, 310002):
        element(series_group, "c:axId", val=axis_id)
    for category, own, cross, position in ((True,310001,310002,"b"),(False,310002,310001,"l")):
        axis = element(plot, "c:catAx" if category else "c:valAx")
        element(axis, "c:axId", val=own)
        scaling = element(axis, "c:scaling")
        element(scaling, "c:orientation", val="minMax")
        element(axis, "c:delete", val="0")
        element(axis, "c:axPos", val=position)
        if not category:
            grid = element(axis, "c:majorGridlines")
            _line(element(grid, "c:spPr"), "E0E0E0")
        element(axis, "c:numFmt", formatCode="General", sourceLinked="0")
        element(axis, "c:majorTickMark", val="none")
        element(axis, "c:minorTickMark", val="none")
        element(axis, "c:tickLblPos", val="low" if category else "nextTo")
        _line(element(axis, "c:spPr"), "808080")
        _chart_text(axis)
        element(axis, "c:crossAx", val=cross)
        element(axis, "c:crosses", val="autoZero")
        if category:
            element(axis, "c:auto", val="1")
            element(axis, "c:lblAlgn", val="ctr")
            element(axis, "c:lblOffset", val="100")
        else:
            element(axis, "c:crossBetween", val="between")
    if spec["series"]:
        legend = element(chart, "c:legend")
        element(legend, "c:legendPos", val="b")
        element(legend, "c:layout")
        element(legend, "c:overlay", val="0")
        _chart_text(legend)
    element(chart, "c:plotVisOnly", val="1")
    element(chart, "c:dispBlanksAs", val="gap")
    properties = element(root, "c:spPr")
    _fill(properties, "FFFFFF")
    _line(properties, "D9D9D9")
    _chart_text(root)
    data = element(root, "c:externalData")
    data.set("{" + NS["r"] + "}id", workbook_relationship)
    element(data, "c:autoUpdate", val="0")
    return _xml(root)


def append_word_chart(document, spec, width=7.1, height=3.2):
    """Insert an editable chart with an embedded full panel workbook."""
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml
    package = document.part.package
    existing = {str(x.partname) for x in package.iter_parts()}
    number = 1
    while "/word/charts/chart" + str(number) + ".xml" in existing:
        number += 1
    chart = Part(PackURI("/word/charts/chart" + str(number) + ".xml"), "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", b"", package)
    workbook = Part(PackURI("/word/embeddings/chart" + str(number) + ".xlsx"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _workbook(spec), package)
    relationship = chart.relate_to(workbook, RT.PACKAGE)
    chart._blob = _chart_xml(spec, relationship)
    chart_relationship = document.part.relate_to(chart, RT.CHART)
    drawing = ET.Element("{" + NS["w"] + "}drawing")
    inline = element(drawing, "wp:inline", distT="0", distB="0", distL="0", distR="0")
    element(inline, "wp:extent", cx=int(width * 914400), cy=int(height * 914400))
    element(inline, "wp:effectExtent", l="0", t="0", r="0", b="0")
    element(inline, "wp:docPr", id=document.part.next_id, name=spec["title"], descr="Editable retained report chart")
    properties = element(inline, "wp:cNvGraphicFramePr")
    element(properties, "a:graphicFrameLocks", noChangeAspect="1")
    graphic = element(inline, "a:graphic")
    data = element(graphic, "a:graphicData", uri=NS["c"])
    ref = element(data, "c:chart")
    ref.set("{" + NS["r"] + "}id", chart_relationship)
    paragraph = document.add_paragraph()
    paragraph.add_run()._r.append(parse_xml(_xml(drawing)))
    return paragraph


def pdf_chart(spec, width=511.2, height=220):
    """Plot every panel value as a grayscale vector ReportLab chart."""
    from reportlab.graphics.shapes import Drawing, String, Rect
    from reportlab.graphics.charts.barcharts import VerticalBarChart, HorizontalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase.pdfmetrics import stringWidth
    drawing = Drawing(width, height)
    horizontal = spec["kind"] == "bar" and any(len(x)>24 for x in spec["categories"])
    chart = HorizontalLineChart() if spec["kind"] == "line" else HorizontalBarChart() if horizontal else VerticalBarChart()
    legend_columns = max(1,min(len(spec["series"]),4,int((width-68)/115)))
    legend_width = (width-68)/legend_columns
    legend_entries = []
    if spec["series"]:
        for index,series in enumerate(spec["series"]):
            lines, line = [], ""
            for char in series["name"]:
                if line and stringWidth(line+char,"QuestionnaireSans",8)>legend_width-14:
                    lines.append(line); line = ""
                line += char
            lines.append(line)
            legend_entries.append((index,lines))
    legend_rows = []
    for start in range(0,len(legend_entries),legend_columns):
        group = legend_entries[start:start+legend_columns]
        legend_rows.append((group,max(len(lines) for _,lines in group)*10+5))
    legend_height = sum(h for _,h in legend_rows)
    chart.x, chart.y = 52, max(46,legend_height+27)
    chart.width, chart.height = width-68, height-chart.y-15
    category_names = spec["categories"]
    if horizontal:
        label_width = min(width*.43,240)
        wrapped = []
        for text in category_names:
            lines, line = [], ""
            for word in text.split(" "):
                candidate = (line+" "+word).strip()
                if line and stringWidth(candidate,"QuestionnaireSans",8)>label_width:
                    lines.append(line); line = word
                else: line = candidate
            lines.append(line); wrapped.append("\n".join(lines))
        category_names = wrapped
        chart.x, chart.width = label_width+12,width-label_width-24
    original_data = [[None if x is None else float(Decimal(str(x))) for x in s["values"]] for s in spec["series"]]
    style_sources = list(range(len(original_data)))
    if spec["kind"] == "line":
        # HorizontalLineChart joins across None by default. Separate retained
        # contiguous segments so a withheld/missing year stays a real gap,
        # matching the native Office chart's dispBlanksAs=gap.
        chart_data, style_sources = [], []
        for source_index, row in enumerate(original_data):
            segment = [None] * len(row)
            active = False
            for value_index, value in enumerate(row + [None]):
                if value is None:
                    if active:
                        chart_data.append(segment)
                        style_sources.append(source_index)
                        segment, active = [None] * len(row), False
                else:
                    segment[value_index], active = value, True
        chart.data = chart_data or original_data
    else:
        chart.data = original_data
    chart.categoryAxis.categoryNames = category_names
    chart.categoryAxis.labels.fontName = "QuestionnaireSans"
    chart.categoryAxis.labels.fontSize = 8
    chart.categoryAxis.labels.boxAnchor = "e" if horizontal else "ne"
    chart.categoryAxis.labels.angle = 0
    chart.categoryAxis.strokeColor = HexColor("#808080")
    chart.valueAxis.labels.fontName = "QuestionnaireSans"
    chart.valueAxis.labels.fontSize = 8
    chart.valueAxis.strokeColor = HexColor("#808080")
    chart.valueAxis.gridStrokeColor = HexColor("#E0E0E0")
    chart.valueAxis.visibleGrid = True
    values = [x for row in chart.data for x in row if x is not None]
    if not values:
        drawing.add(String(52, 110, "No visible numeric answers", fontName="QuestionnaireSans", fontSize=10))
        return drawing
    minimum = min(0, min(values))
    maximum = max(0, max(values))
    span = maximum - minimum or 1
    chart.valueAxis.valueMin = minimum - (.08 * span if minimum < 0 else 0)
    chart.valueAxis.valueMax = maximum + .08 * span
    if horizontal:
        chart.categoryAxis.labels.dx = -6 - max(0,-chart.valueAxis.valueMin)/(chart.valueAxis.valueMax-chart.valueAxis.valueMin)*chart.width
        chart.categoryAxis.labels.dy = 0
    else:
        chart.categoryAxis.labels.dy = -6 - max(0, -chart.valueAxis.valueMin) / (chart.valueAxis.valueMax - chart.valueAxis.valueMin) * chart.height
    for index, source_index in enumerate(style_sources):
        color = HexColor("#" + PALETTE[source_index % len(PALETTE)])
        if spec["kind"] == "line":
            chart.lines[index].strokeColor = color
            chart.lines[index].strokeWidth = 1.7
            marker = makeMarker("FilledCircle")
            marker.fillColor = marker.strokeColor = color
            marker.size = 4
            chart.lines[index].symbol = marker
        else:
            chart.bars[index].fillColor = color
            chart.bars[index].strokeColor = HexColor("#FFFFFF")
    y = legend_height+5
    for group,row_height in legend_rows:
        for column,(index,lines) in enumerate(group):
            color = HexColor("#"+PALETTE[index%len(PALETTE)])
            x = 52+column*legend_width
            drawing.add(Rect(x,y-6,6,6,fillColor=color,strokeColor=color))
            for line_index,line in enumerate(lines):
                drawing.add(String(x+10,y-6-line_index*10,line,fontName="QuestionnaireSans",fontSize=8,fillColor=HexColor("#000000")))
        y -= row_height
    drawing.add(chart)
    return drawing
