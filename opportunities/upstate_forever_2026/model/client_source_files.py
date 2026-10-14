"""Original-file inspection and explicit native source projections.

TokenJunkieLabs. Reads source files without altering them. Scientific preparation,
alignment, validation and native execution remain in their original modules.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, time
import itertools
import json
import math
from pathlib import Path
import re

PREVIEW_ROWS = 30
ENCODINGS = {"utf-8-sig", "utf-16", "cp1252"}
DELIMITERS = {",", ";", "\t", "|"}
TABLE_FIELDS = ("lucode", "load_n", "load_p", "load_type_n", "load_type_p",
                "eff_n", "eff_p", "crit_len_n", "crit_len_p",
                "proportion_subsurface_n", "usle_c", "usle_p")
ROLES = {
    "raster": ("CONTEXT", "DEM", "NLCD", "PRECIPITATION", "R_FACTOR", "K_FACTOR"),
    "vector": ("CONTEXT", "MODEL_DOMAIN", "REPORTING_AREA", "WATERSHEDS", "STATIONS"),
    "table": ("CONTEXT", "BIOPHYSICAL", "MEASUREMENTS"),
    "document": ("CONTEXT",),
}
TARGET_UNITS = {"DEM": "m", "NLCD": "class", "R_FACTOR": "MJ*mm/(ha*h*year)",
                "K_FACTOR": "t*ha*h/(ha*MJ*mm)"}


def need(condition, message):
    if not condition:
        raise ValueError(message)


def text(value, name, required=False, maximum=8000):
    need(isinstance(value, str) and len(value) <= maximum, "Enter a valid " + name + ".")
    need(not required or bool(value.strip()), "Enter " + name + ".")
    return value


def scalar(value):
    if type(value) is int and abs(value) > 9007199254740991:
        return str(value)
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def options(value):
    value = {} if value is None else value
    need(isinstance(value, dict), "Choose original-file reading options.")
    unknown = set(value) - {"encoding", "delimiter", "header_row", "formula_mode"}
    need(not unknown, "Unsupported reading option: " + ", ".join(sorted(unknown)))
    result = dict(value)
    result.setdefault("encoding", "utf-8-sig")
    need(result["encoding"] in ENCODINGS, "Choose UTF-8, UTF-16 or Windows-1252 explicitly.")
    if result.get("delimiter") in (None, ""):
        result.pop("delimiter", None)
    else:
        need(result["delimiter"] in DELIMITERS, "Choose comma, tab, semicolon or pipe.")
    result.setdefault("header_row", 1)
    need(type(result["header_row"]) is int and 1 <= result["header_row"] <= 10000,
         "Choose the original 1-based header row.")
    result.setdefault("formula_mode", "formulas")
    need(result["formula_mode"] in {"formulas", "cached"},
         "Choose preserved formulas or the original cached formula results.")
    return result


def headers(values):
    raw = ["" if v is None else str(scalar(v)) for v in values]
    return [{"id": "c" + str(i), "index": i, "label": name,
             "display": (name if name.strip() else "Unnamed") + " · column " + str(i + 1)}
            for i, name in enumerate(raw)]


def table_records(path, selected, sheet=None):
    """Yield original record and physical row/line positions through the same readers."""
    path = Path(path)
    opts = options(selected)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as error:
            raise ValueError("This installation needs openpyxl to read this workbook. Keep the original file and install the declared workbook dependency.") from error
        book = openpyxl.load_workbook(path, read_only=True,
                                     data_only=opts["formula_mode"] == "cached",
                                     keep_links=False)
        try:
            need(sheet in book.sheetnames, "Choose an original worksheet.")
            for number, row in enumerate(book[sheet].iter_rows(values_only=False), 1):
                yield {"row": number, "line_start": number, "line_end": number,
                       "values": [cell.value for cell in row],
                       "cell_metadata": [{"data_type": getattr(cell, "data_type", None),
                                          "number_format": getattr(cell, "number_format", None)}
                                         for cell in row]}
        finally:
            book.close()
        return
    need(path.suffix.lower() in {".csv", ".tsv", ".txt"}, "Choose a CSV, TSV or modern Excel workbook.")
    with path.open("r", encoding=opts["encoding"], newline="") as stream:
        sample = stream.read(65536)
        stream.seek(0)
        delimiter = opts.get("delimiter")
        if not delimiter:
            if path.suffix.lower() == ".tsv":
                delimiter = "\t"
            else:
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
        reader = csv.reader(stream, delimiter=delimiter, strict=True)
        previous_line = 0
        for number, row in enumerate(reader, 1):
            yield {"row": number, "line_start": previous_line + 1,
                   "line_end": reader.line_num, "values": row}
            previous_line = reader.line_num


def table_rows(path, selected, sheet=None):
    """Backward-compatible row view of the original native reader."""
    source = table_records(path, selected, sheet)
    try:
        for record in source:
            yield record["row"], record["values"]
    finally:
        source.close()


def table_preview(path, opts, sheet=None):
    iterator = table_records(path, opts, sheet)
    header_number = opts["header_row"]
    try:
        first = list(itertools.islice(iterator, header_number))
        need(len(first) == header_number, "The selected header row is beyond this " +
             ("worksheet." if sheet is not None else "file."))
        fields = headers(first[-1]["values"])
        def preview_row(record):
            value = {"row": record["row"], "cells": [scalar(x) for x in record["values"]]}
            if "cell_metadata" in record:
                value["cell_metadata"] = record["cell_metadata"]
            return value
        rows = list(itertools.islice(iterator, PREVIEW_ROWS + 1))
        return {"kind": "table", "name": sheet or Path(path).name, "sheet": sheet,
                "fields": fields, "header_row": header_number,
                "suggested_mapping": {name: matches[0]["id"] for name in TABLE_FIELDS
                                      if len(matches := [field for field in fields
                                          if re.sub(r"[\s-]+", "_", field["label"].strip().lower()) == name]) == 1},
                "preamble": [preview_row(record) for record in first[:-1]],
                "rows": [preview_row(record) for record in rows[:PREVIEW_ROWS]],
                "more_rows": len(rows) > PREVIEW_ROWS,
                "sample_limit": PREVIEW_ROWS,
                "formula_mode": opts["formula_mode"],
                "questions": ["Confirm the selected header row and original column meanings.",
                              "Confirm units, analytical fractions, dates/time zone and source period; these are not inferred from sample values."]}
    finally:
        iterator.close()


def inspect_file(path, supplied=None):
    path = Path(path)
    opts = options(supplied)
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return {"format": suffix[1:], "options": opts, "parts": [table_preview(path, opts)]}
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as error:
            raise ValueError("Install openpyxl in the original application runtime to inspect this workbook. Its raw original remains retained.") from error
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            names = list(book.sheetnames)
        finally:
            book.close()
        parts = []
        for name in names:
            try:
                parts.append(table_preview(path, opts, name))
            except ValueError as error:
                # A cover or empty sheet must not hide usable scientific sheets.
                # Keep its original position so reviewed part identities are stable.
                parts.append({"kind": "table", "name": name, "sheet": name,
                              "header_row": opts["header_row"],
                              "error": {"type": type(error).__name__, "message": str(error)}})
        return {"format": suffix[1:], "options": opts, "parts": parts,
                "questions": ["The complete original workbook is retained, including unselected worksheets and formulas. Cached formula values are used only if explicitly selected. Macros are never executed."]}
    if suffix in {".tif", ".tiff"}:
        try:
            import rasterio
        except ImportError as error:
            raise ValueError("Install the original GIS runtime with rasterio to inspect this raster; the uploaded file is retained.") from error
        with rasterio.open(path) as src:
            parts = []
            for band in range(1, src.count + 1):
                parts.append({"kind": "raster", "name": src.descriptions[band-1] or "Band " + str(band),
                              "band": band, "width": src.width, "height": src.height,
                              "crs": src.crs.to_string() if src.crs else None,
                              "bounds": [scalar(x) for x in src.bounds],
                              "transform": [scalar(x) for x in list(src.transform)[:6]],
                              "dtype": src.dtypes[band-1], "nodata": scalar(src.nodatavals[band-1]),
                              "declared_unit": src.units[band-1], "scale": scalar(src.scales[band-1]),
                              "offset": scalar(src.offsets[band-1]), "tags": src.tags(band),
                              "questions": ["Confirm physical units and any explicit conversion. No scale, offset or resampling is silently applied.",
                                            "Confirm the band, source period and scientific role. A geographic CRS is not a metre grid."]})
            return {"format": "geotiff", "options": opts, "parts": parts}
    if suffix in {".gpkg", ".geojson", ".json", ".shp"}:
        try:
            import fiona
            from model.rasters import vector_environment
        except ImportError as error:
            raise ValueError("Install the original GIS runtime with Fiona to inspect these layers; the raw files remain retained.") from error
        with vector_environment():
            names = fiona.listlayers(path)
            parts = []
            for layer in names:
                with fiona.open(path, layer=layer) as src:
                    props = dict(src.schema["properties"])
                    rows = []
                    for index, feature in enumerate(itertools.islice(src, PREVIEW_ROWS)):
                        rows.append({"row": index + 1, "feature_id": str(feature.id),
                                     "cells": [scalar(feature["properties"].get(k)) for k in props]})
                    parts.append({"kind": "vector", "name": layer, "layer": layer,
                                  "feature_count": len(src), "geometry": src.schema["geometry"],
                                  "crs": src.crs_wkt or (str(src.crs) if src.crs else None),
                                  "bounds": [scalar(x) for x in src.bounds],
                                  "fields": [{**field, "native_type": props[field["label"]]}
                                             for field in headers(props)],
                                  "rows": rows, "more_rows": len(src) > PREVIEW_ROWS,
                                  "questions": ["Confirm the original layer and identifier field. Original geometry and attributes are preserved.",
                                                "Select model domain and reporting area separately; station points are not polygon catchments."]})
            return {"format": suffix[1:], "options": opts, "parts": parts}
    return {"format": suffix[1:] or "unknown", "options": opts,
            "parts": [{"kind": "document", "name": path.name,
                       "questions": ["This format is retained as an original document. Open it with its normal application; select a supported scientific file for native calculation."]}]}


def checked_selection(value, inspection):
    need(isinstance(value, dict), "Review an original source selection.")
    result = dict(value)
    need(type(result.get("file_index")) is int and 0 <= result["file_index"] < len(inspection["files"]),
         "Choose an inspected original file.")
    file = inspection["files"][result["file_index"]]
    need(not file.get("error"), "Resolve this original file's inspection error before using it.")
    parts = file["inspection"]["parts"]
    need(type(result.get("part_index")) is int and 0 <= result["part_index"] < len(parts),
         "Choose an original worksheet, layer or raster band.")
    part = parts[result["part_index"]]
    need(not part.get("error"), "Correct this worksheet's reading choices and inspect again before selecting it.")
    role = result.get("role", "CONTEXT")
    need(role in ROLES[part["kind"]], "Choose a scientific role compatible with this original file.")
    result["role"] = role
    for name in ("title", "vintage", "uri", "units", "period", "timezone", "notes", "reviewer", "reference"):
        result[name] = text(result.get(name, ""), name, name in {"title", "vintage", "uri", "reviewer"}, 16000)
    if role in TARGET_UNITS or role == "PRECIPITATION":
        need(bool(result["units"].strip()), "Enter the original raster's physical units.")
    result["intake_status"] = "context-retained-awaiting-scientific-review" if role == "CONTEXT" else "reviewed-for-selected-role"
    if role != "CONTEXT" and part["kind"] in {"raster", "vector"}:
        need(bool(part.get("crs")), "The original spatial file has no CRS. Correct its source CRS explicitly before native use.")
    if role in {"MODEL_DOMAIN", "REPORTING_AREA", "WATERSHEDS"}:
        need(part.get("geometry") in {"Polygon", "MultiPolygon", "3D Polygon", "3D MultiPolygon"},
             "The selected scientific boundary requires a polygon layer.")
    if role != "CONTEXT" and part["kind"] == "raster":
        need(part.get("scale") == 1 and part.get("offset") == 0,
             "This raster declares a nontrivial native scale/offset. Prepare an explicit physical raster before model use; do not relabel its pixels.")
    result["resampling"] = result.get("resampling", "nearest" if role == "NLCD" else "bilinear")
    need(result["resampling"] in {"nearest", "bilinear", "average", "mode"}, "Choose the later native grid resampling method.")
    need(role != "NLCD" or result["resampling"] in {"nearest", "mode"}, "Categorical land-cover requires nearest or mode resampling.")
    conversion = result.get("conversion")
    if conversion is None:
        result.pop("conversion", None)
    if conversion is not None:
        need(part["kind"] == "raster" and role not in {"NLCD", "CONTEXT"}, "Physical scale conversion is available only for a reviewed continuous scientific raster, not context or classes.")
        need(isinstance(conversion, dict), "Review the physical conversion.")
        factor = conversion.get("factor")
        need(type(factor) in (int, float) and math.isfinite(factor) and factor > 0,
             "Enter a finite positive physical conversion factor.")
        target = text(conversion.get("target_units"), "target units", True)
        ref = text(conversion.get("reference"), "conversion source/reference", True)
        result["conversion"] = {"factor": factor, "target_units": target, "reference": ref,
                                "original_units": result["units"]}
    actual_units = result.get("conversion", {}).get("target_units", result["units"])
    if role in TARGET_UNITS:
        need(actual_units == TARGET_UNITS[role], role + " needs " + TARGET_UNITS[role] +
             ". Enter the correct original unit and an explicit sourced conversion if needed.")
    fields = part.get("fields", [])
    field_ids = {x["id"] for x in fields}
    mapping = result.get("mapping", {})
    need(isinstance(mapping, dict) and all(isinstance(k, str) and k in TABLE_FIELDS and v in field_ids
                                         for k, v in mapping.items()),
         "Map native biophysical fields to original column selections.")
    need(not mapping or part["kind"] == "table" and role == "BIOPHYSICAL",
         "Native table field mapping belongs to a selected biophysical table.")
    result["mapping"] = mapping
    models = result.get("models", ["ndr", "sdr"])
    need(isinstance(models, list) and models and len(set(models)) == len(models)
         and all(x in {"ndr", "sdr"} for x in models), "Select NDR, SDR or both for table validation.")
    result["models"] = models
    for name in ("identifier_field", "label_field"):
        selected = result.get(name, "")
        need(selected == "" or selected in field_ids, "Choose an actual original " + name.replace("_", " ") + ".")
        result[name] = selected
    need(result.get("confirmed") is True, "Confirm this explicit source interpretation before saving.")
    return result, file, part


def project_table(path, target, selected, part, reading):
    """One explicit column projection; complete original remains a companion."""
    field_map = selected["mapping"]
    need(bool(field_map), "Map the biophysical fields before producing its native CSV.")
    indexes = {f["id"]: f["index"] for f in part["fields"]}
    names = [name for name in TABLE_FIELDS if name in field_map]
    source = table_rows(path, reading, part.get("sheet"))
    # A cached formula result must not be silently replaced by missing data.
    formulas = table_rows(path, {**reading, "formula_mode": "formulas"}, part.get("sheet")) \
        if reading.get("formula_mode") == "cached" and Path(path).suffix.lower() in {".xlsx", ".xlsm"} else None
    count = 0
    try:
        with Path(target).open("w", encoding="utf-8", newline="") as out:
            writer = csv.writer(out)
            writer.writerow(names)
            for number, row in source:
                original = next(formulas)[1] if formulas else row
                if number <= part["header_row"]:
                    continue
                if not any(value is not None and str(value) != "" for value in original):
                    continue
                values = []
                for name in names:
                    index = indexes[field_map[name]]
                    need(index < len(row), "Original row " + str(number) + " has no mapped " + name + " cell.")
                    value = row[index]
                    raw = original[index] if index < len(original) else None
                    if isinstance(raw, str) and raw.startswith("="):
                        need(reading["formula_mode"] == "cached" and value is not None,
                             "Original row " + str(number) + " contains a formula without an explicitly selected cached result. Recalculate/save the original workbook or supply reviewed values.")
                    need(not (isinstance(value, str) and value.startswith("=")),
                         "A native parameter cell contains an unevaluated formula.")
                    values.append("" if value is None else scalar(value))
                writer.writerow(values)
                count += 1
    finally:
        source.close()
        if formulas:
            formulas.close()
    from model.profiles import inspect_table
    inspected = inspect_table(target, selected_models=selected["models"])
    need(len(inspected["rows"]) == count, "The complete native table row count differs.")
    return {"row_count": count, "columns": inspected["columns"],
            "lucodes": inspected["lucodes"], "original_sheet": part.get("sheet"),
            "original_header_row": part["header_row"], "mapping": field_map}


def select_band(path, target, band):
    import rasterio
    with rasterio.open(path) as src:
        need(type(band) is int and 1 <= band <= src.count, "Choose an original raster band.")
        profile = src.profile.copy()
        profile.update(driver="GTiff", count=1)
        with rasterio.Env(GDAL_TIFF_INTERNAL_MASK=True), rasterio.open(target, "w", **profile) as out:
            for _, window in src.block_windows(band):
                out.write(src.read(band, window=window), 1, window=window)
                if src.mask_flag_enums[band-1]:
                    out.write_mask(src.read_masks(band, window=window), window=window)
            out.update_tags(**src.tags())
            out.update_tags(1, **src.tags(band))
            if src.descriptions[band-1]:
                out.set_band_description(1, src.descriptions[band-1])
            if src.units[band-1]:
                out.set_band_unit(1, src.units[band-1])
    return {"original_band": band, "operation": "Exact selected band; no unit conversion, resampling or CRS change"}
