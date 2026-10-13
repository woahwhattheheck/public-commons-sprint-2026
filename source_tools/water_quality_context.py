"""Official water-quality evidence joined to the maintained Upper Savannah model.

The module reads publisher files, retains their distinct editions and meanings,
and joins their HUC identifiers and original geometries to a selected study.
It does not change InVEST inputs, model parameters, or physical export results.
Authored for TokenJunkieLabs.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import html
import io
import json
import math
import re
import shutil
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from itertools import islice
from pathlib import Path
from typing import Any

SCHEMA = "upstate.water-quality-context/v1"
CATALOG_SCHEMA = "upstate.water-quality-catalog/v1"
METHOD_VERSION = "water-quality-context-native-identity-causes-class-labels-20261013"
FOCUS = {
    "030601010703": {"name": "Upper Six and Twenty Creek", "utility": "ARJWS"},
    "030601010704": {"name": "Lower Six and Twenty Creek", "utility": "ARJWS"},
    "030601020502": {"name": "Choestoea Creek", "utility": "PRWD"},
    "030601020507": {"name": "Lower Tugaloo River", "utility": "PRWD"},
}
CLASS_GROUPS = {
    "ONRW": "outstanding_national_resource",
    "ORW": "outstanding_resource",
    "TN": "trout_natural",
    "TPGT": "trout_put_grow_take",
    "TPT": "trout_put_take",
    "FW": "freshwater",
    "SA": "saltwater_sa",
    "SB": "saltwater_sb",
    "SFH": "shellfish_harvesting",
}
CLASS_LABELS = {
    "ONRW": "Outstanding National Resource Waters",
    "ORW": "Outstanding Resource Waters",
    "TN": "Trout Natural",
    "TPGT": "Trout Put, Grow, and Take",
    "TPT": "Trout Put and Take",
    "FW": "Freshwaters",
    "SA": "Saltwaters SA",
    "SB": "Saltwaters SB",
    "SFH": "Shellfish Harvesting Waters"
}
CLASS_LABEL_SOURCE = "https://des.sc.gov/sites/des/files/Library/Regulations/R.61-68.pdf"
HIGH_QUALITY = {"ONRW", "ORW", "TN", "TPGT", "TPT"}
CAUSES = {
    "ECOLI": ("bacteria", "E. coli"),
    "E COLI": ("bacteria", "E. coli"),
    "ESCHERICHIA COLI": ("bacteria", "E. coli"),
    "FC": ("bacteria", "Fecal coliform"),
    "FECAL COLIFORM": ("bacteria", "Fecal coliform"),
    "ENTERO": ("bacteria", "Enterococci"),
    "ENTEROCOCCI": ("bacteria", "Enterococci"),
    "TN": ("nitrogen", "Total nitrogen"),
    "TOTAL NITROGEN": ("nitrogen", "Total nitrogen"),
    "NITROGEN": ("nitrogen", "Nitrogen"),
    "TP": ("phosphorus", "Total phosphorus"),
    "TOTAL PHOSPHORUS": ("phosphorus", "Total phosphorus"),
    "PHOSPHORUS": ("phosphorus", "Phosphorus"),
    "TURB": ("turbidity", "Turbidity"),
    "TURBIDITY": ("turbidity", "Turbidity"),
    "TSS": ("suspended_solids", "Total suspended solids"),
    "TOTAL SUSPENDED SOLIDS": ("suspended_solids", "Total suspended solids"),
    "SEDIMENT": ("sediment", "Sediment"),
    "SEDIMENTATION": ("sediment", "Sedimentation"),
    "DO": ("dissolved_oxygen", "Dissolved oxygen"),
    "DISSOLVED OXYGEN": ("dissolved_oxygen", "Dissolved oxygen"),
    "CHLA": ("chlorophyll_a", "Chlorophyll a"),
    "CHLOROPHYLL A": ("chlorophyll_a", "Chlorophyll a"),
    "BIO": ("biological", "Biological assessment"),
    "BIOLOGICAL": ("biological", "Biological assessment"),
    "PH": ("ph", "pH"),
    "HG": ("mercury", "Mercury in water column"),
    "MERCURY": ("mercury", "Mercury"),
    "PCBS": ("pcbs", "PCBs"),
    "LEAD": ("lead", "Lead"),
    "PB": ("lead", "Lead"),
    "PCB": ("pcbs", "Polychlorinated biphenyls in fish tissue"),
    "HGF": ("mercury", "Mercury in fish tissue"),
    "NH3": ("ammonia_nitrogen", "Ammonia nitrogen"),
    "CD": ("cadmium", "Cadmium"),
    "CR": ("chromium", "Chromium"),
    "CU": ("copper", "Copper"),
    "NI": ("nickel", "Nickel"),
    "ZN": ("zinc", "Zinc"),
}
SCDES_CAUSE_CODES = {
    "NH3", "CD", "CHLA", "CR", "CU", "DO", "ENTERO", "ECOLI", "FC", "PH",
    "PB", "BIO", "HG", "HGF", "NI", "PCB", "TN", "TP", "TURB", "TURBIDITY", "ZN",
}
SCDES_CAUSE_GUIDE = "https://www.des.sc.gov/sites/des/files/Documents/BOW/WaterQuality/303dList2024.pdf"
SOURCE_INDEX = "https://www.des.sc.gov/programs/bureau-water/south-carolina-303d-list-impaired-waters-tmdls"
STANDARDS_INDEX = "https://des.sc.gov/programs/bureau-water/water-quality-standards"
TERMINALS = {"", "0", "000000000000", "CLOSED BASIN", "CLOSED", "OCEAN", "CANADA", "MEXICO"}
CONTEXT_FILES = (
    "context.json", "watershed-context.geojson", "classified-reaches.geojson",
    "assessments.json", "watershed-context.csv", "assessment-evidence.csv",
    "classified-reach-huc-crosswalk.csv", "water-quality-context.html",
)


def utcnow() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def read_json(path: str | Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: str | Path, value: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = (json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n").encode()
    temporary = path.with_suffix(path.suffix + ".partial")
    temporary.write_bytes(data)
    temporary.replace(path)


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Non-finite source value")
    return str(value).strip()


def huc12(value: Any, *, required: bool = True) -> str | None:
    # Never repair a HUC by padding a spreadsheet number whose leading zero was lost.
    value = text(value)
    if re.fullmatch(r"\d{12}", value):
        return value
    if not required and not value:
        return None
    raise ValueError(f"Expected original 12-character HUC text; received {value!r}")


def cause_info(value: Any) -> dict:
    raw = text(value)
    key = re.sub(r"[._-]", " ", raw.upper())
    key = re.sub(r"\s+", " ", key).strip()
    family, label = CAUSES.get(key, ("other", raw or "Unspecified cause"))
    return {"cause_raw": raw, "cause_family": family, "cause_label": label}


def class_info(value: Any) -> dict:
    raw = text(value)
    # Preserve source notations separately; do not infer different reach boundaries.
    codes = re.findall(r"(?<![A-Z])(?:ONRW|ORW|TPGT|TPT|TN|FW|SFH|SA|SB)(?![A-Z])", raw.upper())
    codes = sorted(set(codes))
    return {
        "classification_raw": raw,
        "classification_codes": codes,
        "classification_groups": [CLASS_GROUPS[c] for c in codes],
        "classification_labels": [CLASS_LABELS[c] for c in codes],
        "high_quality_designation": bool(HIGH_QUALITY.intersection(codes)),
        "unrecognized_classification": bool(raw and not codes),
    }


def use_support(value: Any) -> str:
    key = re.sub(r"\s+", " ", text(value).lower())
    if key in {"not supported", "not supporting", "impaired"}:
        return "not_supported"
    if key in {"fully supported", "fully supporting", "supported"}:
        return "fully_supported"
    if key in {"not assessed", "not evaluated", "insufficient data"}:
        return "not_assessed"
    return "unspecified"


def _safe_name(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", value):
        raise ValueError(f"Unsafe source identifier: {value!r}")
    return value


def _request(url: str, output: Path, *, timeout: float = 45, limit: int = 100_000_000) -> dict:
    if urllib.parse.urlsplit(url).scheme != "https":
        raise ValueError("Publisher acquisition requires an HTTPS URL")
    request = urllib.request.Request(url, headers={"User-Agent": "TokenJunkieLabs watershed source receiver"})
    partial = output.with_suffix(output.suffix + ".partial")
    output.parent.mkdir(parents=True, exist_ok=True)
    total = 0
    sha = hashlib.sha256()
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response, partial.open("wb") as stream:
            for block in iter(lambda: response.read(1024 * 1024), b""):
                total += len(block)
                if total > limit:
                    raise ValueError(f"Source exceeds declared {limit}-byte acquisition limit")
                sha.update(block)
                stream.write(block)
            receipt = {
                "requested_url": url, "response_url": response.geturl(),
                "status": response.status, "captured_at": utcnow(),
                "content_type": response.headers.get("Content-Type"),
                "etag": response.headers.get("ETag"),
                "last_modified": response.headers.get("Last-Modified"),
                "bytes": total, "sha256": sha.hexdigest(),
            }
        partial.replace(output)
        write_json(output.with_suffix(output.suffix + ".capture.json"), receipt)
        return receipt
    except BaseException:
        if partial.exists():
            partial.unlink()
        raise


def acquire_document(item: dict, root: Path, *, resume: bool = True) -> dict:
    sid = _safe_name(item["id"])
    suffix = Path(urllib.parse.urlsplit(item["uri"]).path).suffix or ".body"
    output = root / f"{sid}{suffix}"
    receipt_path = output.with_suffix(output.suffix + ".capture.json")
    if resume and output.exists() and receipt_path.exists():
        receipt = read_json(receipt_path)
        actual = digest(output.read_bytes())
        if actual != receipt["sha256"] or output.stat().st_size != receipt["bytes"]:
            raise ValueError(f"Retained source bytes changed: {sid}")
        if receipt["requested_url"] != item["uri"]:
            raise ValueError(f"Retained source URL differs: {sid}; select a new source id")
    else:
        if output.exists() or receipt_path.exists():
            raise ValueError(f"Incomplete existing source pair for {sid}; preserve it under a separate id")
        receipt = _request(item["uri"], output, limit=int(item.get("max_bytes", 100_000_000)))
    expected = item.get("sha256")
    if expected and receipt["sha256"] != expected:
        raise ValueError(f"Publisher source changed for {sid}: expected {expected}, got {receipt['sha256']}")
    return {
        **{k: v for k, v in item.items() if k not in {"max_bytes"}},
        "path": output.name, "sha256": receipt["sha256"],
        "companions": [{"path": receipt_path.name, "sha256": digest(receipt_path.read_bytes())}],
        "acquisition": receipt,
    }


def _arcgis_json(url: str, root: Path, name: str, parameters: dict, *, resume: bool) -> tuple[dict, dict]:
    encoded = urllib.parse.urlencode(parameters)
    address = url.rstrip("/") + ("&" if "?" in url else "?") + encoded
    item = {"id": name, "uri": address, "title": name, "kind": "document", "vintage": "capture"}
    entry = acquire_document(item, root, resume=resume)
    value = read_json(root / entry["path"])
    if "error" in value:
        raise ValueError(f"ArcGIS provider error: {json.dumps(value['error'])}")
    return value, entry


def acquire_classification(item: dict, root: Path, *, resume: bool = True) -> dict:
    """Receive exact object-id pages; keep all original response bytes and field metadata."""
    sid = _safe_name(item["id"])
    bank = root / sid
    bank.mkdir(parents=True, exist_ok=True)
    metadata, meta_entry = _arcgis_json(item["uri"], bank, "metadata", {"f": "json"}, resume=resume)
    oid = metadata.get("objectIdField") or metadata.get("objectIdFieldName")
    if not oid:
        oid = next((x["name"] for x in metadata.get("fields", []) if x.get("type") == "esriFieldTypeOID"), None)
    if not oid:
        raise ValueError("Classification service has no declared object-id field")
    names = {f["name"] for f in metadata.get("fields", [])}
    for required in ("H2OCLASS", oid):
        if required not in names:
            raise ValueError(f"Classification source lacks {required}")
    query = {"f": "json", "where": item.get("where", "1=1"), "returnIdsOnly": "true"}
    ids_body, ids_entry = _arcgis_json(item["uri"] + "/query", bank, "object-ids", query, resume=resume)
    ids = ids_body.get("objectIds")
    if not isinstance(ids, list):
        raise ValueError("ArcGIS did not return its original object-id inventory")
    if len(set(ids)) != len(ids):
        raise ValueError("ArcGIS object-id inventory contains duplicates")
    ids = sorted(ids)
    entries = [meta_entry, ids_entry]
    features = []
    returned_ids = set()
    source_reference = None
    batch_size = max(1, min(int(item.get("page_size", 100)), int(metadata.get("maxRecordCount", 1000))))
    for start in range(0, len(ids), batch_size):
        selected = ids[start:start + batch_size]
        page, entry = _arcgis_json(
            item["uri"] + "/query", bank, f"features-{start:08d}",
            {"f": "json", "objectIds": ",".join(map(str, selected)), "outFields": "*",
             "returnGeometry": "true", "returnTrueCurves": "false"}, resume=resume,
        )
        if page.get("exceededTransferLimit"):
            raise ValueError("Classification page exceeded transfer limit; resume under a new id with a smaller page size")
        actual_ids = [feature["attributes"][oid] for feature in page.get("features", [])]
        if set(actual_ids) != set(selected) or len(actual_ids) != len(selected):
            raise ValueError(f"Classification page {start} differs from requested original object ids")
        if returned_ids.intersection(actual_ids):
            raise ValueError("Classification capture contains duplicate object ids")
        returned_ids.update(actual_ids)
        if page.get("spatialReference"):
            if source_reference and source_reference != page["spatialReference"]:
                raise ValueError("Classification pages have different spatial references")
            source_reference = page["spatialReference"]
        features.extend(page["features"])
        entries.append(entry)
    if returned_ids != set(ids):
        raise ValueError("Classification capture is incomplete")
    source_reference = source_reference or metadata.get("extent", {}).get("spatialReference")
    if not source_reference:
        raise ValueError("Classification source omitted its spatial reference")
    combined = {
        "displayFieldName": metadata.get("displayField"),
        "objectIdFieldName": oid, "geometryType": metadata.get("geometryType"),
        "spatialReference": source_reference, "fields": metadata.get("fields", []),
        "features": features,
    }
    vector_path = root / f"{sid}.esri.json"
    write_json(vector_path, combined)
    companions = []
    for entry in entries:
        for local in [bank / entry["path"], bank / entry["companions"][0]["path"]]:
            companions.append({"path": str(local.relative_to(root)), "sha256": digest(local.read_bytes())})
    return {
        **item, "kind": "vector", "role": "WATER_CLASSIFICATION",
        "path": vector_path.name, "sha256": digest(vector_path.read_bytes()),
        "companions": companions, "object_id_field": oid, "feature_count": len(features),
        "classification_status": item.get("classification_status", "provisional"),
        "native_spatial_reference": source_reference,
    }


def acquire_catalog(catalog_path: str | Path, output: str | Path, *, resume: bool = True, selected_ids: list[str] | None = None) -> dict:
    catalog = read_json(catalog_path)
    if catalog.get("schema") != CATALOG_SCHEMA:
        raise ValueError(f"Expected {CATALOG_SCHEMA}")
    root = Path(output)
    root.mkdir(parents=True, exist_ok=True)
    entries = []
    known_ids = {item["id"] for item in catalog["entries"]}
    if selected_ids and set(selected_ids) - known_ids:
        raise ValueError(f"Unknown catalogue source ids: {sorted(set(selected_ids) - known_ids)}")
    for item in catalog["entries"]:
        if selected_ids:
            if item["id"] not in selected_ids:
                continue
        elif not item.get("acquire_by_default", True):
            continue
        if item.get("transport") == "arcgis":
            entries.append(acquire_classification(item, root, resume=resume))
        else:
            entries.append(acquire_document(item, root, resume=resume))
    manifest = {
        "schema": "upstate.sources/v1", "title": catalog["title"],
        "captured_at": utcnow(), "scope": catalog.get("scope", {}),
        "notes": "Publisher editions and current/draft status retained on each entry.",
        "entries": entries,
    }
    write_json(root / "sources.json", manifest)
    return manifest


def inspect_table(path: str | Path, max_rows: int = 12) -> dict:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return {"sheets": [
                {"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                 "preview": [[text(v) for v in row] for row in islice(sheet.iter_rows(values_only=True), max_rows)]}
                for sheet in book.worksheets
            ]}
        finally:
            book.close()
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        rows = []
        for _, row in zip(range(max_rows), reader):
            rows.append(row)
    return {"preview": rows}


def read_table(path: str | Path, *, sheet: str | None = None, header_row: int = 1) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = book[sheet] if sheet else book.worksheets[0]
            iterator = ws.iter_rows(values_only=True)
            for _ in range(header_row - 1):
                next(iterator)
            columns = [text(v) for v in next(iterator)]
            if not all(columns) or len(columns) != len(set(columns)):
                raise ValueError("Select the actual unique table header row")
            return [{column: text(value) for column, value in zip(columns, row)}
                    for row in iterator if any(value is not None for value in row)]
        finally:
            book.close()
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("Table requires unique column names")
        return list(reader)


def _source_meta(source: dict, path: str | Path) -> dict:
    path = Path(path)
    actual = digest(path.read_bytes())
    expected = source.get("sha256")
    if expected and expected != actual:
        raise ValueError(f"Source file differs from supplied metadata: {path}")
    return {
        "source_id": source["id"], "source_sha256": actual,
        "source_uri": source["uri"], "source_vintage": source["vintage"],
        "publication_status": source.get("publication_status", "unspecified"),
        "assessment_period": source.get("assessment_period"),
        "status_date": source.get("status_date"),
        "cause_encoding": source.get("cause_encoding", (
            "scdes_303d_codes" if urllib.parse.urlparse(source["uri"]).hostname in
            {"des.sc.gov", "www.des.sc.gov"} and
            urllib.parse.urlparse(source["uri"]).path.endswith("/303dList2024.xlsx")
            else "single")),
    }


def _assessment_causes(value: str, source: dict) -> list[dict]:
    """Expand the publisher's documented cause codes while retaining the whole cell."""
    if source.get("cause_encoding") != "scdes_303d_codes":
        return [cause_info(value)]
    definitions = source.get("annotation_definitions", {})
    markers = sorted({marker for marker in definitions if marker in value})
    clean = value
    for marker in markers:
        clean = clean.replace(marker, "")
    # The native 2024 workbook contains comma-separated codes plus the original
    # punctuation variants "TN. TP, ECOLI" and "HGF PCB".
    codes = [code for code in re.split(r"[\s,;.]+", clean.strip().upper()) if code]
    unknown = sorted(set(codes).difference(SCDES_CAUSE_CODES))
    if not codes or unknown:
        raise ValueError(f"Unmapped SCDES cause code(s) {unknown} in original cell {value!r}")
    return [{
        **cause_info(code), "cause_code": code, "cause_source_cell": value,
        "cause_annotations": [{"marker": marker, "text": definitions[marker]} for marker in markers],
        "cause_code_guide_uri": SCDES_CAUSE_GUIDE, "cause_code_guide_pages": [8, 9],
    } for code in dict.fromkeys(codes)]


def normalize_assessments(rows: list[dict], bindings: dict, source: dict) -> list[dict]:
    required = {"site", "huc12", "cause"}
    if not required.issubset(bindings):
        raise ValueError("Explicit site, huc12 and cause field bindings are required")
    results, identities = [], set()
    for index, raw in enumerate(rows):
        def get(key: str) -> str:
            field = bindings.get(key)
            if not field:
                return ""
            if field not in raw:
                raise ValueError(f"Source row lacks configured field {field!r}")
            return text(raw[field])
        site, huc, cause = get("site"), get("huc12"), get("cause")
        if not any((site, huc, cause)):
            continue
        if not site or not cause:
            raise ValueError(f"Incomplete assessment identity at row {index + 1}")
        huc = huc12(huc)
        support_raw = get("use_support")
        kind = source.get("evidence_kind", "303d_listing")
        normalized_support = use_support(support_raw)
        if kind == "303d_listing" and not support_raw:
            normalized_support = "not_supported"
        for cause_data in _assessment_causes(cause, source):
            identity_cause = cause_data.get("cause_code", cause_data["cause_raw"])
            identity = (source["source_id"], site, get("use"), identity_cause, huc, kind)
            if identity in identities:
                continue
            identities.add(identity)
            row_id = digest(json.dumps(identity, ensure_ascii=False).encode())[:24]
            result = {
                "id": row_id, **source, "evidence_kind": kind, "site": site, "huc12": huc,
                "description": get("description"), "county": get("county"), "use": get("use"),
                **cause_data, "use_support": normalized_support, "use_support_raw": support_raw,
                "tmdl_relation": get("tmdl_relation"), "tmdl_report": get("tmdl_report"),
                "approval_date": get("approval_date"), "source_row": index + 1,
                "raw": raw,
            }
            if source.get("header_row"):
                result["source_table_row"] = index + int(source["header_row"]) + 1
            if get("source_page"):
                result["source_page"] = get("source_page")
            results.append(result)
    return results


def parse_tmdl_pdf(path: str | Path, source: dict) -> dict:
    """Read the publisher's complete current table, retaining page and original row text."""
    import fitz
    meta = _source_meta(source, path)
    meta["evidence_kind"] = "tmdl_attainment"
    basin = r"BROAD|CATAWBA|EDISTO|PEE\s*DEE|SALUDA|SALKEHATCHIE|SAVANNAH|SANTEE"
    start_pattern = re.compile(rf"\b({basin})\s+(\d{{12}})\s+", re.I)
    tail_pattern = re.compile(
        r"(?P<body>.*?)\s*(?P<support>Fully\s+Supported|Not\s+Supported|Not\s+Assessed)"
        r"\s+(?P<relation>In\s*TMDL|Wn\s*TMDL)\s+(?P<report>.+?)"
        r"\s+(?P<approval>\d{1,2}/\d{1,2}/\d{4})", re.I
    )
    body_pattern = re.compile(
        r"(?P<description>.+?)\s+(?P<site>[A-Z0-9][A-Z0-9._-]*)"
        r"\s+(?P<county>[A-Z]+)\s+(?P<use>[A-Z]{2,12})\s+(?P<cause>\S+)$", re.I
    )
    records, unparsed, status_dates = [], [], set()
    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            variants = []
            for sorted_text in (False, True):
                page_text = re.sub(r"\s+", " ", page.get_text("text", sort=sorted_text))
                for date in re.findall(r"as of (\d{1,2}/\d{1,2}/\d{4})", page_text, re.I):
                    status_dates.add(dt.datetime.strptime(date, "%m/%d/%Y").date().isoformat())
                anchors = list(start_pattern.finditer(page_text))
                page_records, page_unparsed = [], []
                source_hucs = re.findall(r"\b\d{12}\b", page_text)
                if len(source_hucs) != len(anchors):
                    page_unparsed.append({"page": page_index + 1, "text": page_text,
                                          "reason": "HUC identifiers outside recognized table rows"})
                if not anchors and "Sites Covered Under" in page_text:
                    page_unparsed.append({"page": page_index + 1, "text": page_text,
                                          "reason": "Table page contains no readable source rows"})
                for i, anchor in enumerate(anchors):
                    raw = page_text[anchor.end():anchors[i + 1].start() if i + 1 < len(anchors) else len(page_text)]
                    match = tail_pattern.search(raw)
                    body = body_pattern.fullmatch(match.group("body").strip()) if match else None
                    if not match or not body:
                        page_unparsed.append({"page": page_index + 1, "huc12": anchor.group(2), "text": raw})
                        continue
                    fields = {
                        "BASIN": re.sub(r"\s+", " ", anchor.group(1).upper()),
                        "HUC12": anchor.group(2), "DESCRIPTION": body.group("description"),
                        "SITE": body.group("site"), "COUNTY": body.group("county"),
                        "USE": body.group("use"), "CAUSE": body.group("cause"),
                        "USE_SUPPORT": re.sub(r"\s+", " ", match.group("support")),
                        "TMDL": re.sub(r"\s+", "", match.group("relation")), "REPORT": match.group("report"),
                        "APPROVAL_DATE": match.group("approval"),
                    }
                    binding = {"huc12": "HUC12", "site": "SITE", "cause": "CAUSE", "use": "USE",
                               "description": "DESCRIPTION", "county": "COUNTY",
                               "use_support": "USE_SUPPORT", "tmdl_relation": "TMDL",
                               "tmdl_report": "REPORT", "approval_date": "APPROVAL_DATE"}
                    row = normalize_assessments([fields], binding, meta)[0]
                    row.update({"source_page": page_index + 1, "source_row_text": anchor.group(0) + raw[:match.end()]})
                    page_records.append(row)
                variants.append((page_records, page_unparsed, len(anchors)))
            chosen = max(variants, key=lambda v: (len(v[0]), -len(v[1])))
            records.extend(chosen[0])
            unparsed.extend(chosen[1])
    if not records:
        raise ValueError("No complete publisher TMDL rows parsed; inspect the current source layout")
    if len(status_dates) != 1:
        raise ValueError(f"Publisher table status-date inventory is ambiguous: {sorted(status_dates)}")
    actual_date = next(iter(status_dates))
    expected_date = source.get("status_date")
    if expected_date and expected_date != actual_date:
        raise ValueError(f"TMDL edition changed: expected {expected_date}, received {actual_date}")
    for record in records:
        record["status_date"] = actual_date
    return {
        "schema": "upstate.water-quality-assessments/v1", "source": {**meta, "status_date": actual_date},
        "records": records, "unparsed_rows": unparsed,
        "row_count": len(records), "complete": not unparsed,
    }


def normalize_table(path: str | Path, source: dict, bindings: dict, **kwargs) -> dict:
    meta = _source_meta(source, path)
    meta["evidence_kind"] = source.get("evidence_kind", "303d_listing")
    meta["header_row"] = kwargs.get("header_row", 1)
    if kwargs.get("sheet"):
        meta["sheet"] = kwargs["sheet"]
    rows = read_table(path, **kwargs)
    notes = []
    for index, row in enumerate(rows):
        if not any(text(row.get(bindings.get(key, ""))) for key in ("site", "huc12", "cause")):
            if any(text(value) for value in row.values()):
                notes.append({"source_table_row": index + int(meta["header_row"]) + 1, "raw": row})
    if meta["cause_encoding"] == "scdes_303d_codes":
        definitions = {}
        for note in notes:
            content = " ".join(text(value) for value in note["raw"].values() if text(value))
            match = re.match(r"^(†|‡|\*+)\s*(.+)$", content)
            if match:
                definitions[match.group(1)] = match.group(2)
        meta["annotation_definitions"] = definitions
        meta["cause_code_guide_uri"] = SCDES_CAUSE_GUIDE
        meta["cause_code_guide_pages"] = [8, 9]
    records = normalize_assessments(rows, bindings, meta)
    return {
        "schema": "upstate.water-quality-assessments/v1", "source": meta,
        "records": records, "unparsed_rows": [], "row_count": len(records), "complete": True,
        "source_table_row_count": len(rows), "source_assessment_row_count": len(rows) - len(notes),
        "source_notes": notes,
    }


def _geometry_rows(path: str | Path, *, layer=None, target_crs="EPSG:5070"):
    from model.rasters import vector_geometry
    return vector_geometry(str(path), layer=layer, target_crs=target_crs)


def _classification_identity(rows: list[dict], path: str | Path, source: dict,
                             explicit_field: str | None = None) -> tuple[list[dict], dict]:
    """Bind feature identity to the retained source declaration, never its alias."""
    import gzip

    path = Path(path)
    with path.open("rb") as stream:
        compressed = stream.read(2) == b"\x1f\x8b"
    opener = gzip.open if compressed else open
    with opener(path, "rb") as stream:
        prefix = stream.read(256).lstrip(b"\xef\xbb\xbf \t\r\n")
    native = {}
    if prefix.startswith(b"{"):
        with opener(path, "rt", encoding="utf-8-sig") as stream:
            value = json.load(stream)
        if isinstance(value, dict):
            native = value

    def declaration(value):
        fields = value.get("fields") or value.get("field_metadata") or []
        typed = [field["name"] for field in fields
                 if field.get("type") == "esriFieldTypeOID"]
        named = value.get("objectIdFieldName") or value.get("objectIdField")
        if len(set(typed)) > 1 or (named and typed and named not in typed):
            raise ValueError("Original classification metadata declares conflicting native ID fields")
        return named or (typed[0] if typed else None)

    original_field = declaration(native)
    metadata = source.get("metadata") if isinstance(source.get("metadata"), dict) else {}
    source_fields = source.get("source_fields") if isinstance(source.get("source_fields"), dict) else {}
    source_field = (source.get("object_id_field") or declaration(source)
                    or metadata.get("object_id_field") or declaration(metadata)
                    or source_fields.get("object_id"))
    selected = explicit_field or original_field or source_field
    if not isinstance(selected, str) or not selected.strip():
        raise ValueError("Select the original classification ID field or retain its typed source metadata")
    selected = selected.strip()
    if original_field and selected != original_field:
        raise ValueError(f"Classification ID {selected} differs from the original native ID "
                         f"{original_field}; use the field name, not its display alias")

    originals = {}
    for feature in native.get("features", []):
        properties = feature.get("attributes", feature.get("properties", {}))
        value = properties.get(selected)
        if value is None:
            continue
        key = text(value)
        if not key or key in originals:
            raise ValueError("Original classification response contains an empty or duplicate native ID")
        originals[key] = value

    received, recovered = [], 0
    for row in rows:
        properties = dict(row["properties"])
        if selected not in properties:
            native_id = text(row.get("id"))
            if native_id not in originals:
                raise ValueError(f"Classification source lacks {selected} and its driver feature "
                                 "ID is absent from the original response")
            # Some vector drivers expose the declared original OID as the feature
            # ID. Recover it only after matching the original response's IDs.
            properties[selected] = originals[native_id]
            recovered += 1
        received.append({**row, "properties": properties})
    fields = native.get("fields") or native.get("field_metadata") or []
    return received, {
        "object_id_field": selected,
        "basis": ("explicit original field" if explicit_field else
                  "original Esri field declaration" if original_field else
                  "retained source entry declaration"),
        "original_declared_field": original_field,
        "source_entry_field": source_field,
        "field_definition": next((field for field in fields if field.get("name") == selected), None),
        "native_response_id_count": len(originals),
        "driver_feature_ids_recovered": recovered,
    }


def _topology(domain_rows: list[dict], huc_field: str, to_field: str | None,
              bindings_path: str | Path | None) -> tuple[dict, dict]:
    graph, attributes, source_edges = {}, {}, set()
    for row in domain_rows:
        props = row["properties"]
        hid = huc12(props[huc_field])
        if hid in attributes:
            raise ValueError(f"Duplicate domain HUC12: {hid}")
        attributes[hid] = row
        if to_field and to_field in props:
            source_edges.add(hid)
        elif not bindings_path:
            raise ValueError(f"Domain lacks configured topology field {to_field!r}; supply original bindings")
        value = text(props.get(to_field)) if to_field else ""
        graph[hid] = None if value.upper() in TERMINALS else huc12(value)
    if bindings_path:
        value = read_json(bindings_path)
        bindings = value if isinstance(value, list) else value.get("features", value.get("bindings", value.get("feature_bindings", [])))
        if not bindings:
            raise ValueError("Topology binding file has no feature bindings")
        seen = set()
        for binding in bindings:
            hid = huc12(binding["huc12"])
            if hid not in graph:
                continue
            value = text(binding.get("tohuc"))
            downstream = None if value.upper() in TERMINALS else huc12(value)
            if hid in seen:
                raise ValueError(f"Duplicate topology binding: {hid}")
            if hid in source_edges and graph[hid] != downstream:
                raise ValueError(f"Topology source disagreement for {hid}")
            graph[hid] = downstream
            seen.add(hid)
        if seen != set(graph):
            raise ValueError("Topology bindings do not cover the complete selected HUC domain")
    for start in graph:
        path, active = [], set()
        current = start
        while current in graph:
            if current in active:
                raise ValueError("Domain topology cycle: " + " -> ".join(path + [current]))
            active.add(current)
            path.append(current)
            current = graph[current]
    return graph, attributes


def _ancestors(graph: dict, target: str) -> set:
    reverse = defaultdict(set)
    for origin, downstream in graph.items():
        if downstream:
            reverse[downstream].add(origin)
    selected, queue = {target}, [target]
    while queue:
        current = queue.pop()
        for parent in reverse.get(current, ()):
            if parent not in selected:
                selected.add(parent)
                queue.append(parent)
    return selected


def _compact_assessment(record: dict) -> dict:
    return {key: value for key, value in record.items() if key not in {"raw", "source_row_text"}}


def build_context(
    domain_path: str | Path,
    assessment_paths: list[str | Path],
    output: str | Path,
    *,
    domain_source: dict,
    huc_field: str = "HUC12",
    to_field: str | None = "ToHUC",
    topology_path: str | Path | None = None,
    domain_layer: str | None = None,
    classification_path: str | Path | None = None,
    classification_source: dict | None = None,
    classification_layer: str | None = None,
    class_field: str = "H2OCLASS",
    class_id_field: str | None = None,
    class_name_field: str = "HYDNAME",
    include_draft: bool = False,
) -> dict:
    """Produce joined layers, original evidence tables and source-bound watershed context."""
    from shapely.geometry import mapping
    from shapely.ops import transform, unary_union
    from shapely.strtree import STRtree
    from pyproj import Transformer

    root = Path(output)
    root.mkdir(parents=True, exist_ok=True)
    if (root / "artifacts-complete.json").exists():
        raise ValueError("Context output is complete; reopen it or choose a new context directory")
    domain_meta = _source_meta(domain_source, domain_path)
    domain_rows, crs = _geometry_rows(domain_path, layer=domain_layer)
    graph, domain = _topology(domain_rows, huc_field, to_field, topology_path)
    ids = list(domain)
    polygons = [domain[key]["geometry"] for key in ids]
    if any(p.geom_type not in {"Polygon", "MultiPolygon"} for p in polygons):
        raise ValueError("Selected HUC domain contains non-polygon geometry")
    tree = STRtree(polygons)
    region = unary_union(polygons)
    to_wgs84 = Transformer.from_crs(crs, "EPSG:4326", always_xy=True).transform
    assessment_sources, records, exclusions, unmatched = [], [], [], []
    seen_records = set()
    for path in assessment_paths:
        bank = read_json(path)
        if bank.get("schema") != "upstate.water-quality-assessments/v1":
            raise ValueError("Assessment bank has an incompatible schema")
        if not bank.get("complete", False):
            raise ValueError(f"Assessment source contains unparsed rows: {path}")
        assessment_sources.append({**bank["source"], "normalized_sha256": digest(Path(path).read_bytes())})
        for record in bank["records"]:
            if record["id"] in seen_records:
                raise ValueError(f"Duplicate retained assessment identity: {record['id']}")
            seen_records.add(record["id"])
            if not include_draft and record.get("publication_status") in {"draft", "proposed"}:
                exclusions.append({"id": record["id"], "reason": "draft_edition", "source_id": record["source_id"]})
                continue
            huc = huc12(record["huc12"])
            if huc not in domain:
                unmatched.append(record["id"])
                continue
            records.append(record)

    by_huc = defaultdict(list)
    for record in records:
        by_huc[record["huc12"]].append(record)
    class_features, class_links = [], defaultdict(list)
    class_meta = None
    if classification_path:
        if not classification_source:
            raise ValueError("Classification source metadata is required with its vector file")
        class_meta = _source_meta(classification_source, classification_path)
        class_meta["classification_status"] = classification_source.get(
            "classification_status", classification_source.get("mapped_classification_status", "provisional"))
        rows, _ = _geometry_rows(classification_path, layer=classification_layer)
        rows, identity_binding = _classification_identity(
            rows, classification_path, classification_source, class_id_field)
        class_id_field = identity_binding["object_id_field"]
        class_meta["object_id_field"] = class_id_field
        class_meta["identity_binding"] = identity_binding
        seen_class = set()
        for row in rows:
            props = row["properties"]
            if class_id_field not in props:
                raise ValueError(f"Classified-reach source lacks native identity field {class_id_field}")
            native_id = text(props[class_id_field])
            if not native_id:
                raise ValueError("Classified-reach native identity is empty")
            if native_id in seen_class:
                raise ValueError(f"Duplicate classified-reach native id: {native_id}")
            seen_class.add(native_id)
            if class_field not in props:
                raise ValueError(f"Classified-reach source lacks {class_field}")
            geometry = row["geometry"]
            dimension = (1 if geometry.geom_type in {"LineString", "MultiLineString"} else
                         2 if geometry.geom_type in {"Polygon", "MultiPolygon"} else None)
            if dimension is None:
                raise ValueError(f"Classified reach {native_id} has unsupported geometry {geometry.geom_type}")
            clip = geometry.intersection(region)
            if clip.is_empty or (clip.length if dimension == 1 else clip.area) <= 0:
                continue
            info = class_info(props[class_field])
            linked = []
            for index in tree.query(geometry, predicate="intersects"):
                hid = ids[int(index)]
                cut = geometry.intersection(polygons[int(index)])
                # A point touch of a reach, or boundary touch of a waterbody,
                # does not transfer its classification to an adjacent HUC.
                if cut.is_empty or (cut.length if dimension == 1 else cut.area) <= 0:
                    continue
                relation = {
                    "huc12": hid, "source_feature_id": native_id,
                    "source_geometry_dimension": dimension,
                    "intersection_length_m": cut.length if dimension == 1 else None,
                    "intersection_area_m2": cut.area if dimension == 2 else None,
                    **info,
                }
                class_links[hid].append(relation)
                linked.append(hid)
            class_features.append({
                "type": "Feature", "id": native_id, "geometry": mapping(transform(to_wgs84, clip)),
                "properties": {
                    **props, **info, "source_feature_id": native_id,
                    "source_id": class_meta["source_id"], "source_sha256": class_meta["source_sha256"],
                    "classification_status": class_meta["classification_status"],
                    "waterbody_name": text(props.get(class_name_field)), "huc12_intersections": sorted(linked),
                },
            })

    rows, watershed_features = [], []
    upstream_classes = defaultdict(set)
    for target, links in class_links.items():
        codes = {code for link in links for code in link["classification_codes"] if code in HIGH_QUALITY}
        if codes:
            for ancestor in _ancestors(graph, target):
                upstream_classes[ancestor].update(codes)
    for hid in ids:
        local = by_huc.get(hid, [])
        upstream = _ancestors(graph, hid)
        upstream_records = [record for ancestor in upstream for record in by_huc.get(ancestor, [])]
        counts = defaultdict(lambda: {"listed_303d": 0, "tmdl_not_supported": 0, "tmdl_fully_supported": 0, "other": 0})
        for record in local:
            bucket = counts[record["cause_family"]]
            if record["evidence_kind"] == "303d_listing":
                bucket["listed_303d"] += 1
            elif record["evidence_kind"] == "tmdl_attainment":
                bucket["tmdl_not_supported" if record["use_support"] == "not_supported" else
                       "tmdl_fully_supported" if record["use_support"] == "fully_supported" else "other"] += 1
            else:
                bucket["other"] += 1
        focus = FOCUS.get(hid, {})
        row = {
            "huc12": hid, "name": focus.get("name", text(domain[hid]["properties"].get("Name"))),
            "focus_utility": focus.get("utility"), "area_km2": domain[hid]["geometry"].area / 1e6,
            "downstream_huc12": graph[hid],
            "downstream_outside_selection": bool(graph[hid] and graph[hid] not in graph),
            "upstream_huc12_count": len(upstream),
            "assessment_record_count": len(local), "assessment_site_count": len({r["site"] for r in local}),
            "nonattaining_site_count": len({r["site"] for r in local if r["use_support"] == "not_supported"}),
            "tmdl_nonattaining_site_count": len({r["site"] for r in local
                                                if r["evidence_kind"] == "tmdl_attainment" and r["use_support"] == "not_supported"}),
            "assessment_cause_counts": dict(counts),
            "classified_feature_count": len({r["source_feature_id"] for r in class_links.get(hid, [])}),
            "direct_classification_codes": sorted({c for r in class_links.get(hid, []) for c in r["classification_codes"]}),
            "upstream_of_high_quality_codes": sorted(upstream_classes.get(hid, set())),
            "upstream_assessment_record_count": len(upstream_records),
            "upstream_nonattaining_site_count": len({r["site"] for r in upstream_records if r["use_support"] == "not_supported"}),
            "assessment_ids": sorted(r["id"] for r in local),
            "classified_feature_ids": sorted({r["source_feature_id"] for r in class_links.get(hid, [])}),
        }
        row["direct_classification_labels"] = [CLASS_LABELS[code] for code in row["direct_classification_codes"]]
        row["upstream_of_high_quality_labels"] = [CLASS_LABELS[code] for code in row["upstream_of_high_quality_codes"]]
        rows.append(row)
        watershed_features.append({
            "type": "Feature", "id": hid,
            "geometry": mapping(transform(to_wgs84, domain[hid]["geometry"])),
            "properties": row,
        })
    assessments = [_compact_assessment(r) for r in records]
    context = {
        "schema": SCHEMA, "created_at": utcnow(), "title": "Upper Savannah water-quality context",
        "domain_source": domain_meta, "assessment_sources": assessment_sources,
        "classification_source": class_meta,
        "topology_source": ({"path": Path(topology_path).name, "sha256": digest(Path(topology_path).read_bytes())} if topology_path else
                            {"source": domain_meta["source_id"], "huc_field": huc_field, "to_field": to_field}),
        "method": {
            "version": METHOD_VERSION,
            "spatial_measurement_crs": "EPSG:5070", "export_crs": "EPSG:4326",
            "assessment_join": "Publisher HUC12 identifier; station-level evidence retained.",
            "classified_reach_join": "Original geometry intersection with selected HUC polygons.",
            "classification_boundary_rule": "Positive length for source lines or positive area for source polygons; boundary-only touches excluded.",
            "upstream_context": "Selected source ToHUC graph; context at HUC resolution.",
            "high_quality_codes": sorted(HIGH_QUALITY), "draft_editions_included": include_draft,
            "classification_labels": CLASS_LABELS,
            "classification_label_source_uri": CLASS_LABEL_SOURCE,
            "classification_label_source_section": "R.61-68, Section G; publisher copy with May 23, 2025 amendment",
            "code_meaning": "Classification TN means Trout Natural; assessment-cause TN means Total nitrogen.",
            "absence_meaning": "No record in the selected sources; use support remains unspecified.",
            "model_use": "Geographic context; model N/P/sediment mass remains in its original run.",
        },
        "watersheds": rows, "assessments": assessments,
        "classification_links": [link for hid in ids for link in class_links.get(hid, [])],
        "excluded_records": exclusions, "outside_selected_domain_record_ids": unmatched,
        "summary": {
            "huc12_count": len(ids), "focus_huc12_present": sorted(set(ids).intersection(FOCUS)),
            "assessment_record_count": len(records), "assessment_site_count": len({r["site"] for r in records}),
            "nonattaining_site_count": len({r["site"] for r in records if r["use_support"] == "not_supported"}),
            "classified_feature_count": len(class_features),
        },
    }
    write_json(root / "context.json", context)
    write_json(root / "watershed-context.geojson", {"type": "FeatureCollection", "features": watershed_features})
    write_json(root / "classified-reaches.geojson", {"type": "FeatureCollection", "features": class_features})
    write_json(root / "assessments.json", {"schema": "upstate.water-quality-assessments/v1", "records": records, "complete": True,
                                          "source": {"id": "joined-context", "sources": assessment_sources}})
    flat_rows = [{k: v for k, v in row.items() if not isinstance(v, (dict, list))} for row in rows]
    _write_csv(root / "watershed-context.csv", flat_rows)
    _write_csv(root / "assessment-evidence.csv", assessments)
    _write_csv(root / "classified-reach-huc-crosswalk.csv", context["classification_links"])
    _write_report(root / "water-quality-context.html", context)
    exports = []
    for filename in CONTEXT_FILES:
        file = root / filename
        if file.is_file():
            role = "WATER_QUALITY_CONTEXT" if file.suffix == ".geojson" else "WATER_QUALITY_EVIDENCE"
            exports.append({
                "id": _safe_name(file.name.replace(".", "-").replace("_", "-")), "title": file.name,
                "uri": SOURCE_INDEX, "vintage": context["created_at"],
                "kind": "vector" if file.suffix == ".geojson" else "table" if file.suffix in {".csv", ".json"} else "document",
                "path": file.name, "sha256": digest(file.read_bytes()), "role": role,
                "derived_from": [domain_meta] + assessment_sources + ([class_meta] if class_meta else []),
            })
    manifest = {
        "schema": "upstate.sources/v1", "title": context["title"], "captured_at": context["created_at"],
        "scope": {"huc12": ids, "kind": "geographic_water_quality_context"},
        "notes": "Derived context and full evidence are linked to retained publisher source identities.",
        "entries": exports,
    }
    write_json(root / "sources.json", manifest)
    write_json(root / "artifacts-complete.json", {
        "schema": "upstate.water-quality-context-completion/v1",
        "files": [{"path": name, "sha256": digest((root / name).read_bytes()),
                   "bytes": (root / name).stat().st_size} for name in (*CONTEXT_FILES, "sources.json")],
    })
    return context


def _write_csv(path: Path, rows: list[dict]) -> None:
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
                             for key, value in row.items()})


def _write_report(path: Path, context: dict) -> None:
    e = lambda v: html.escape(text(v), quote=True)
    rows = []
    for row in context["watersheds"]:
        rows.append("<tr>" + "".join("<td>" + e(v) + "</td>" for v in [
            row["huc12"], row["name"], row["focus_utility"], row["assessment_site_count"],
            row["nonattaining_site_count"], row["tmdl_nonattaining_site_count"],
            "; ".join(CLASS_LABELS.get(code, code) + " (" + code + ")" for code in row["direct_classification_codes"]),
            "; ".join(CLASS_LABELS.get(code, code) + " (" + code + ")" for code in row["upstream_of_high_quality_codes"]),
        ]) + "</tr>")
    sources = [context["domain_source"]] + context["assessment_sources"] + ([context["classification_source"]] if context["classification_source"] else [])
    source_rows = "".join("<li><a href='" + e(s["source_uri"]) + "'>" + e(s["source_id"]) + "</a> · " +
                          e(s.get("status_date") or s["source_vintage"]) + " · " + e(s.get("publication_status")) + "</li>" for s in sources)
    path.write_text(
        "<!doctype html><html lang='en'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<title>Upper Savannah water-quality context</title><style>"
        "body{font:15px system-ui,sans-serif;color:#1e2226;background:#fff;margin:28px;max-width:1500px}"
        "h1{font-size:27px}h2{font-size:20px}p{max-width:95ch;line-height:1.5}"
        "table{border-collapse:collapse;width:100%;font-size:13px}th,td{border:1px solid #cbd0d4;padding:8px;text-align:left}"
        "th{background:#edf0f2}thead{position:sticky;top:0}.table{overflow:auto}a{color:#233d55}"
        "@media print{thead{display:table-header-group}tr{break-inside:avoid}body{margin:12mm}.table{overflow:visible}}</style>"
        "<h1>Upper Savannah water-quality context</h1>"
        "<p>Selected source editions connect official assessments and classified reaches to the study's watersheds. "
        "The table shows evidence at its published site and HUC geography, alongside the retained upstream connections.</p>"
        "<p><strong>Reading the table.</strong> Nonattaining sites include listed 303(d) records and current TMDL attainment records. "
        "Sites can leave the 303(d) list when a TMDL is approved while remaining nonattaining. An empty count means no record "
        "in these selected sources. High-quality-water columns preserve the official designation codes and selected upstream HUC context.</p>"
        "<p>Download: <a href='watershed-context.csv'>watershed table</a> · <a href='assessment-evidence.csv'>site evidence</a> · "
        "<a href='watershed-context.geojson'>watershed layer</a> · <a href='classified-reaches.geojson'>classified reaches</a> · "
        "<a href='context.json'>complete editable context</a> · <a href='sources.json'>Workspace import manifest</a></p>"
        "<h2>Watershed context</h2><div class='table'><table><thead><tr>"
        "<th>HUC12</th><th>Watershed</th><th>Focus utility</th><th>Assessment sites</th><th>Nonattaining sites</th>"
        "<th>TMDL nonattaining sites</th><th>Intersecting classes</th><th>Downstream high-quality context</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
        "<h2>Reading classification and assessment codes</h2><p>Classification TN means Trout Natural; "
        "assessment-cause TN means Total nitrogen. Display names follow "
        "<a href='" + CLASS_LABEL_SOURCE + "'>SCDES R.61-68, Section G</a>. "
        "Raw classification codes, original reach identities and the selected source editions remain in the exports.</p>"
        "<h2>Source editions</h2><ul>" + source_rows + "</ul>"
        "<h2>Method</h2><p>Original classified-reach geometry is intersected with selected HUC boundaries. "
        "Area and length use EPSG:5070; portable geographic layers use EPSG:4326. Upstream relationships use the selected "
        "source ToHUC graph. The original nutrient and sediment model runs retain their physical units and outputs. "
        "Turbidity, bacteria, dissolved oxygen and other assessment causes remain separate evidence families.</p>"
        "<p>Prepared for TokenJunkieLabs · " + e(context["created_at"]) + "</p></html>", encoding="utf-8"
    )



def build_workspace_context(workspace, selection: dict) -> dict:
    """Build/reopen a complete context using current content-addressed Workspace sources.

    Selection accepts domain/classifications references and any number of assessment
    references. Each reference contains collection_id and source_id. Assessments
    additionally select format=tmdl_pdf|table|normalized and explicit field bindings.
    Returned file descriptors plug directly into the original operator download API.
    """
    from model.common import describe_file, verify_file

    def receive(reference: dict) -> tuple[Path, dict]:
        path, entry = workspace.source_file(reference["collection_id"], reference["source_id"])
        metadata = dict(entry)
        metadata["sha256"] = entry["body"]["sha256"]
        return Path(path), metadata

    domain_ref = selection["domain"]
    domain_path, domain_source = receive(domain_ref)
    classification_ref = selection.get("classifications")
    classification_path, classification_source = receive(classification_ref) if classification_ref else (None, None)
    acquired = [(reference, *receive(reference)) for reference in selection.get("assessments", [])]
    identity = {
        "schema": SCHEMA, "method_version": METHOD_VERSION, "selection": selection,
        "domain_sha256": domain_source["sha256"],
        "classification_sha256": classification_source["sha256"] if classification_source else None,
        "assessment_sha256": [metadata["sha256"] for _, _, metadata in acquired],
    }
    selection_id = digest(json.dumps(identity, sort_keys=True, separators=(",", ":"), allow_nan=False).encode())
    output = workspace.root / "water-quality-context-artifacts" / selection_id
    completed = output / "workspace-record.json"
    if completed.exists():
        existing = read_json(completed)
        for descriptor in existing["files"].values():
            verify_file(workspace.root, descriptor)
        return existing
    output.mkdir(parents=True, exist_ok=True)
    banks = output / "input-banks"
    banks.mkdir(exist_ok=True)
    paths = []
    for index, (reference, path, metadata) in enumerate(acquired):
        metadata["evidence_kind"] = reference.get("evidence_kind", metadata.get("evidence_kind", "303d_listing"))
        fmt = reference.get("format", "table")
        if fmt == "normalized":
            bank = read_json(path)
            if bank.get("schema") != "upstate.water-quality-assessments/v1":
                raise ValueError("Selected normalized assessment source has an incompatible schema")
        elif fmt == "tmdl_pdf":
            bank = parse_tmdl_pdf(path, metadata)
        elif fmt == "table":
            bank = normalize_table(
                path, metadata, reference["bindings"], sheet=reference.get("sheet"),
                header_row=int(reference.get("header_row", 1)),
            )
        else:
            raise ValueError(f"Unknown assessment input format: {fmt}")
        target = banks / f"{index:03d}-{_safe_name(metadata['id'])}.json"
        write_json(target, bank)
        if not bank.get("complete"):
            raise ValueError(f"Source parsing needs a layout correction; retained rows at {target}")
        paths.append(target)
    topology = None
    if domain_ref.get("topology_companion"):
        relative = Path(domain_ref["topology_companion"])
        if relative.is_absolute() or ".." in relative.parts:
            raise ValueError("Topology companion must be relative to the reconstructed domain source")
        topology = domain_path.parent / relative
        if not topology.is_file():
            raise ValueError(f"Original topology companion is unavailable: {relative}")
    if (output / "artifacts-complete.json").exists():
        # Reuse only the complete artifact set. A partial export resumes from
        # the retained original inputs and rebuilds its derived files.
        for descriptor in read_json(output / "artifacts-complete.json")["files"]:
            file = output / descriptor["path"]
            if not file.is_file() or file.stat().st_size != descriptor["bytes"] or digest(file.read_bytes()) != descriptor["sha256"]:
                raise ValueError(f"Completed context file changed: {descriptor['path']}")
        context = read_json(output / "context.json")
        if context["domain_source"]["source_sha256"] != domain_source["sha256"]:
            raise ValueError("Retained context input identity differs from its selection")
    else:
        context = build_context(
            domain_path, paths, output, domain_source=domain_source,
            huc_field=domain_ref.get("huc_field", "HUC12"),
            to_field=domain_ref.get("to_field", "ToHUC"),
            topology_path=topology, domain_layer=domain_ref.get("layer"),
            classification_path=classification_path, classification_source=classification_source,
            classification_layer=classification_ref.get("layer") if classification_ref else None,
            class_field=classification_ref.get("class_field", "H2OCLASS") if classification_ref else "H2OCLASS",
            class_id_field=classification_ref.get("id_field") if classification_ref else None,
            class_name_field=classification_ref.get("name_field", "HYDNAME") if classification_ref else "HYDNAME",
            include_draft=bool(selection.get("include_draft", False)),
        )
    imported = workspace.import_sources(str(output / "sources.json"))
    source_collection_id = imported.get("id") if isinstance(imported, dict) else imported
    files = {
        str(path.relative_to(output)): describe_file(workspace.root, path, title=path.name)
        for path in sorted(output.rglob("*")) if path.is_file() and path.name != "workspace-record.json"
    }
    record = {
        "schema": SCHEMA, "title": selection.get("title", context["title"]), "selection_id": selection_id,
        "selection": selection, "source_collection_id": source_collection_id,
        "summary": context["summary"], "method": context["method"],
        "domain_source": context["domain_source"], "assessment_sources": context["assessment_sources"],
        "classification_source": context["classification_source"], "files": files,
    }
    saved = workspace._save("water-quality-contexts", record)
    record["id"] = saved.get("id") if isinstance(saved, dict) else saved
    write_json(completed, record)
    return record


def import_context(workspace, context_directory: str | Path) -> dict:
    """Use the maintained local-source importer and immutable Workspace collection."""
    root = Path(context_directory)
    context = read_json(root / "context.json")
    if context.get("schema") != SCHEMA:
        raise ValueError("Context schema mismatch")
    collection = workspace.import_sources(str(root / "sources.json"))
    record = workspace._save("water-quality-contexts", {
        "schema": SCHEMA, "title": context["title"],
        "source_collection": collection, "context_sha256": digest((root / "context.json").read_bytes()),
        "summary": context["summary"], "method": context["method"],
        "publisher_sources": context["assessment_sources"],
        "domain_source": context["domain_source"],
        "classification_source": context["classification_source"],
    })
    return {"context_id": record, "source_collection": collection, "summary": context["summary"]}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    acq = sub.add_parser("acquire")
    acq.add_argument("--catalog", required=True)
    acq.add_argument("--output", required=True)
    acq.add_argument("--source-id", action="append", help="Acquire only these explicitly selected source ids")
    inspect = sub.add_parser("inspect-table")
    inspect.add_argument("path")
    normal = sub.add_parser("normalize")
    normal.add_argument("--input", required=True)
    normal.add_argument("--source", required=True, help="JSON metadata for this exact publisher file")
    normal.add_argument("--output", required=True)
    normal.add_argument("--bindings", help="Explicit source-column bindings JSON; required for CSV/XLSX")
    normal.add_argument("--sheet")
    normal.add_argument("--header-row", type=int, default=1)
    normal.add_argument("--tmdl-pdf", action="store_true")
    build = sub.add_parser("build")
    build.add_argument("--domain", required=True)
    build.add_argument("--domain-source", required=True)
    build.add_argument("--domain-layer")
    build.add_argument("--huc-field", default="HUC12")
    build.add_argument("--to-field", default="ToHUC")
    build.add_argument("--topology")
    build.add_argument("--assessments", action="append", default=[])
    build.add_argument("--classifications")
    build.add_argument("--classification-source")
    build.add_argument("--classification-layer")
    build.add_argument("--class-id-field", help="Original native ID field name; defaults to retained source metadata")
    build.add_argument("--class-field", default="H2OCLASS")
    build.add_argument("--class-name-field", default="HYDNAME")
    build.add_argument("--include-draft", action="store_true")
    build.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    if args.command == "acquire":
        result = acquire_catalog(args.catalog, args.output, selected_ids=args.source_id)
        print(json.dumps({"manifest": str(Path(args.output) / "sources.json"), "sources": len(result["entries"])}))
    elif args.command == "inspect-table":
        print(json.dumps(inspect_table(args.path), indent=2))
    elif args.command == "normalize":
        source = read_json(args.source)
        if args.tmdl_pdf:
            result = parse_tmdl_pdf(args.input, source)
        else:
            if not args.bindings:
                parser.error("--bindings is required for a table source")
            result = normalize_table(args.input, source, read_json(args.bindings),
                                     sheet=args.sheet, header_row=args.header_row)
        write_json(args.output, result)
        print(json.dumps({"path": args.output, "rows": result["row_count"], "complete": result["complete"]}))
        if not result["complete"]:
            return 2
    else:
        result = build_context(
            args.domain, args.assessments, args.output, domain_source=read_json(args.domain_source),
            huc_field=args.huc_field, to_field=args.to_field or None, topology_path=args.topology,
            domain_layer=args.domain_layer, classification_path=args.classifications,
            classification_source=read_json(args.classification_source) if args.classification_source else None,
            classification_layer=args.classification_layer, class_id_field=args.class_id_field,
            class_field=args.class_field, class_name_field=args.class_name_field,
            include_draft=args.include_draft,
        )
        print(json.dumps({"path": str(Path(args.output) / "context.json"), "summary": result["summary"]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
