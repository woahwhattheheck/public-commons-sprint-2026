"""Editable bilingual standards and selected SCC delivery workpapers.

The caller owns selection, archive custody and original attachments. These
functions consume one captured state/edition and return safe relative file
names mapped to bytes. They never publish, fetch, alter state or infer approval.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal
import difflib
from hashlib import sha256
from html import escape
from io import BytesIO
import json
from pathlib import Path
import re
from threading import Lock
from urllib.parse import urlsplit
import zipfile

try:
    from . import engine
except ImportError:
    import engine

SCHEMA = "scc-document-publication/v1"
FONT_DIR = Path(__file__).resolve().parent.parent / "templates" / "fonts"
FONT_LOCK = Lock()
FONT_NAMES = ("SCCVera", "SCCVeraBold", "SCCVeraItalic", "SCCVeraBoldItalic")
PUBLIC_SOURCE_FIELDS = ("id", "title", "url", "publisher", "version", "language",
                        "representation", "license", "sha256", "citation_visibility")
PUBLIC_CLAUSE_FIELDS = ("id", "number", "parent_id", "kind", "title_en", "body_en",
                        "table_en", "title_fr", "body_fr", "table_fr", "source_ids", "case_ids")
PUBLIC_CASE_FIELDS = ("id", "title", "title_fr", "authored", "source_ids", "source_rows",
                      "review_questions", "review_questions_fr", "expected_observations",
                      "expected_observations_fr")
PUBLIC_SCOPE_FIELDS = ("title_en", "title_fr", "deliverable_type", "purpose", "purpose_fr",
                       "audience", "audience_fr", "exclusions", "exclusions_fr",
                       "preface_en", "preface_fr", "roadmap_issues",
                       "national_international_relationship", "source_ids")
PUBLIC_REFERENCE_FIELDS = ("id", "identifier", "title", "owner", "deliverable_type",
                           "stage", "published_date", "maintenance_date", "source_period",
                           "relationship", "rationale", "source_ids", "source_locator", "url")
PUBLIC_RELEASE_FIELDS = ("id", "edition_id", "version", "status", "issued", "url_en", "url_fr",
                         "authority", "access", "license_notice", "scc_acknowledgment", "maintenance_due")
WORDS = {
    "en": {
        "edition": "Edition", "prepared": "Prepared", "released": "Released", "withdrawn": "Withdrawn",
        "frozen": "Edition captured", "purpose": "Purpose and scope", "audience": "Intended audience",
        "exclusions": "Exclusions", "preface": "Preface", "contents": "Contents",
        "sources": "Sources", "cases": "Worked cases", "references": "Related standards and references",
        "record": "Original record", "case": "Case", "questions": "Questions for review",
        "observations": "Recorded expected observations", "review_status": "Language review",
        "normative": "Normative", "informative": "Informative", "annex": "Annex", "table": "Table",
        "no_text": "English text is not recorded for this clause.",
        "no_title": "English title is not recorded", "source_register": "Source register",
        "case_records": "Original case cells can be read in case-records-en.html and are retained in case-records.json.",
        "public_record": "Publication record", "notice": "Recorded publication notice",
        "no_publication": "Edition for review", "language": "English",
        "edition_id": "Edition identifier", "digest": "Edition SHA256", "version": "Version",
        "issued": "Issued", "authority": "Recorded issuing authority", "maintenance": "Maintenance due",
        "license": "Recorded licence notice", "acknowledgment": "Recorded acknowledgment",
        "kind": "Clause type", "no_cases": "No worked cases are linked to this edition.",
    },
    "fr": {
        "edition": "Édition", "prepared": "Préparée", "released": "Publiée", "withdrawn": "Retirée",
        "frozen": "Édition consignée", "purpose": "Objet et domaine d’application",
        "audience": "Public visé", "exclusions": "Exclusions", "preface": "Préface", "contents": "Sommaire",
        "sources": "Sources", "cases": "Cas pratiques", "references": "Normes et références connexes",
        "record": "Enregistrement original", "case": "Cas", "questions": "Questions pour l’examen",
        "observations": "Observations attendues consignées", "review_status": "Examen linguistique",
        "normative": "Normative", "informative": "Informative", "annex": "Annexe", "table": "Tableau",
        "no_text": "Le texte français de cette clause n’est pas consigné.",
        "no_title": "Le titre français n’est pas consigné", "source_register": "Registre des sources",
        "case_records": "Les cellules originales des cas sont consultables dans case-records-fr.html et conservées dans case-records.json.",
        "public_record": "Dossier de publication", "notice": "Avis de publication consigné",
        "no_publication": "Édition pour examen", "language": "Français",
        "edition_id": "Identifiant de l’édition", "digest": "SHA256 de l’édition", "version": "Version",
        "issued": "Date de publication", "authority": "Autorité de publication consignée",
        "maintenance": "Échéance de maintien", "license": "Avis de licence consigné",
        "acknowledgment": "Mention de reconnaissance consignée", "kind": "Type de clause",
        "no_cases": "Aucun cas pratique n’est lié à cette édition.",
    },
}
CSS = """@page{size:letter;margin:19mm 18mm}*{box-sizing:border-box}
html{color:#111;background:#fff;font-family:Arial,Helvetica,sans-serif;font-size:11pt;line-height:1.5}
body{max-width:920px;margin:36px auto;padding:0 30px 50px}h1,h2,h3,h4{color:#000;line-height:1.2}
h1{font-size:25px;font-weight:600;margin:0 0 14px}h2{font-size:18px;margin:27px 0 12px}
h3{font-size:15px;margin:21px 0 9px}h4{font-size:13px;margin:16px 0 8px}
p{margin:0 0 10px;overflow-wrap:anywhere}a{color:#111;text-decoration:underline;overflow-wrap:anywhere}
.subtitle{font-size:12pt;margin:0 0 20px}.small{font-size:9pt;color:#444;line-height:1.45}
table{border-collapse:collapse;width:100%;margin:10px 0 20px;table-layout:fixed}
th,td{border:1px solid #d9d9d9;padding:7px 9px;text-align:left;vertical-align:middle;overflow-wrap:anywhere}
th{background:#e9e9e9;font-weight:600}td.number,th.number{text-align:right;font-variant-numeric:tabular-nums}
thead{display:table-header-group}tr{break-inside:avoid}ul,ol{padding-left:24px;margin:6px 0 16px}
.toc{padding-left:0;list-style:none}.toc li{margin:4px 0}.metadata{margin:12px 0 20px}
del{text-decoration:line-through}ins{text-decoration:underline;font-weight:600}
.footer{font-size:9pt;margin-top:30px;color:#444}
@media print{body{max-width:none;margin:0;padding:0}h1{font-size:20pt}h2,h3,h4{break-after:avoid}
a{color:#000}.screen-only{display:none}.page-break{break-before:page}table{font-size:9pt}}
@media(max-width:600px){body{padding:0 16px;margin:22px auto}table{font-size:9pt}th,td{padding:6px}}
"""


class PublicationError(ValueError):
    pass


def _text(value):
    value = "" if value is None else str(value)
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise PublicationError("A selected text field contains an unsupported document control character.")
    return value


def _json(value):
    return (engine.canonical(value) + "\n").encode("utf-8")


def _pick(row, fields):
    return {key: deepcopy(row[key]) for key in fields if key in row}


def _anchor(value):
    return "s" + sha256(_text(value).encode()).hexdigest()[:24]


def _money(value):
    if value is None or value == "":
        return "Not recorded"
    return format(Decimal(str(value)), ",.2f")


def _amount(rows, field, **where):
    return sum((Decimal(str(row.get(field, "0.00"))) for row in rows
                if all(row.get(k) == v for k, v in where.items())), Decimal("0"))


def _url(value):
    value = _text(value)
    try:
        parsed = urlsplit(value)
        return value if parsed.scheme in ("http", "https") and parsed.netloc else None
    except ValueError:
        return None


def _paragraphs(value):
    return [part.strip() for part in re.split(r"\n\s*\n", _text(value)) if part.strip()]


def _model(title, subtitle="", language="en", timestamp=None):
    return {"title": _text(title), "subtitle": _text(subtitle), "language": language,
            "timestamp": timestamp or "2000-01-01T00:00:00Z", "blocks": []}


def _heading(model, text, level=1, anchor=None):
    model["blocks"].append({"kind": "heading", "text": _text(text),
                            "level": max(1, min(3, level)), "anchor": anchor})


def _para(model, text, small=False):
    for paragraph in _paragraphs(text):
        model["blocks"].append({"kind": "paragraph", "text": paragraph, "small": small})


def _list(model, items):
    if items:
        model["blocks"].append({"kind": "list", "items": [_text(x) for x in items]})


def _table(model, headings, rows, weights=None, numeric=()):
    if not rows:
        return
    width = len(headings)
    normalized = []
    for row in rows:
        if len(row) != width:
            raise PublicationError("A selected table has inconsistent row widths.")
        normalized.append([_text(x) for x in row])
    model["blocks"].append({"kind": "table", "headings": [_text(x) for x in headings],
                            "rows": normalized, "weights": weights or [1] * width,
                            "numeric": list(numeric)})


def _link(model, label, url):
    model["blocks"].append({"kind": "link", "text": _text(label), "url": _url(url)})


def _frozen(edition):
    if not isinstance(edition, dict) or not isinstance(edition.get("snapshot"), dict):
        raise PublicationError("Select one retained edition.")
    if engine.digest(edition["snapshot"]) != edition.get("sha256"):
        raise PublicationError("The selected edition differs from its retained snapshot hash.")
    if edition.get("projection_sha256") is not None and edition["projection_sha256"] != edition["sha256"]:
        raise PublicationError("The selected public projection differs from its recorded projection hash.")
    if edition.get("source_edition_sha256") is not None and not re.fullmatch(r"[0-9a-f]{64}", edition["source_edition_sha256"]):
        raise PublicationError("The original frozen edition hash is malformed.")
    result = deepcopy(edition)
    snapshot = result["snapshot"]
    for name in ("clauses", "references", "cases", "sources", "language_reviews"):
        if not isinstance(snapshot.get(name), list):
            raise PublicationError("The edition is missing its " + name + " snapshot.")
    if not isinstance(snapshot.get("scope"), dict):
        raise PublicationError("The edition has no retained scope.")
    ids = [row["id"] for row in snapshot["clauses"]]
    if len(ids) != len(set(ids)):
        raise PublicationError("The edition contains duplicate clause identifiers.")
    return result


def _public_snapshot(snapshot):
    return {
        "scope": _pick(snapshot["scope"], PUBLIC_SCOPE_FIELDS),
        "clauses": [_pick(row, PUBLIC_CLAUSE_FIELDS) for row in snapshot["clauses"]],
        "references": [_pick(row, PUBLIC_REFERENCE_FIELDS) for row in snapshot["references"]],
        "cases": [_pick(row, PUBLIC_CASE_FIELDS) for row in snapshot["cases"]],
        "sources": [_pick(row, PUBLIC_SOURCE_FIELDS) for row in snapshot["sources"]],
        "language_reviews": deepcopy(snapshot["language_reviews"]),
    }


def _metadata(edition, publication=None):
    return {"schema": SCHEMA, "edition_id": edition["id"], "label": edition["label"],
            "frozen_at": edition["frozen_at"], "edition_sha256": edition["sha256"],
            "source_edition_sha256": edition.get("source_edition_sha256", edition["sha256"]),
            "projection_sha256": edition.get("projection_sha256"),
            "publication": _pick(publication, PUBLIC_RELEASE_FIELDS) if publication else None}


def _clause_level(clause, clauses):
    seen = {clause["id"]}
    parent = clause.get("parent_id")
    level = 1
    while parent:
        if parent in seen or parent not in clauses:
            raise PublicationError("The edition contains an unresolved clause hierarchy.")
        seen.add(parent)
        parent = clauses[parent].get("parent_id")
        level += 1
    return min(level, 3)


def _source_citations(ids, source_numbers):
    return ", ".join("[" + str(source_numbers[x]) + "]" if x in source_numbers else "[" + _text(x) + "]"
                     for x in ids)


def _standard_model(edition, publication, language):
    snap = edition["snapshot"]
    scope = snap["scope"]
    words = WORDS[language]
    title = scope.get("title_" + language) or words["no_title"]
    state = words.get(publication.get("status"), publication.get("status", "")) if publication else words["no_publication"]
    model = _model(title, words["edition"] + " " + edition["label"] + " · " + words["language"] +
                   (" · " + state if state else ""), language, edition["frozen_at"])
    _para(model, words["frozen"] + ": " + edition["frozen_at"], True)
    preface = scope.get("preface_" + language)
    if preface:
        _heading(model, words["preface"])
        _para(model, preface)
    purpose = scope.get("purpose_fr" if language == "fr" else "purpose")
    if purpose:
        _heading(model, words["purpose"])
        _para(model, purpose)
    for key in ("audience", "exclusions"):
        text = scope.get(key + "_fr" if language == "fr" else key)
        if text:
            _heading(model, words[key], 2)
            _para(model, text)
    if publication:
        _heading(model, words["public_record"], 2)
        rows = []
        for key, label in (("version", "version"), ("issued", "issued"),
                           ("authority", "authority"), ("maintenance_due", "maintenance")):
            if publication.get(key):
                rows.append([words[label], publication[key]])
        _table(model, [words["public_record"], words["notice"]], rows, [1, 3])
        for key in ("license_notice", "scc_acknowledgment"):
            if publication.get(key):
                _para(model, publication[key])
    clauses = {c["id"]: c for c in snap["clauses"]}
    source_numbers = {s["id"]: i + 1 for i, s in enumerate(snap["sources"])}
    case_numbers = {c["id"]: i + 1 for i, c in enumerate(snap["cases"])}
    _heading(model, words["contents"])
    model["blocks"].append({"kind": "toc", "items": [
        {"text": c["number"] + "  " + (c.get("title_" + language) or words["no_title"]),
         "anchor": _anchor(c["id"]), "level": _clause_level(c, clauses)} for c in snap["clauses"]]})
    for clause in snap["clauses"]:
        _heading(model, clause["number"] + "  " + (clause.get("title_" + language) or words["no_title"]),
                 _clause_level(clause, clauses), _anchor(clause["id"]))
        _para(model, words["kind"] + ": " + words.get(clause["kind"], clause["kind"]), True)
        body, table = clause.get("body_" + language), clause.get("table_" + language, [])
        if body:
            _para(model, body)
        if table:
            if len(table) == 1:
                _table(model, table[0], [[""] * len(table[0])])
            else:
                _table(model, table[0], table[1:])
        if not body and not table:
            _para(model, words["no_text"])
        if clause.get("source_ids"):
            _para(model, words["sources"] + ": " + _source_citations(clause["source_ids"], source_numbers), True)
        if clause.get("case_ids"):
            labels = [words["case"] + " " + str(case_numbers.get(x, x)) for x in clause["case_ids"]]
            _para(model, words["cases"] + ": " + "; ".join(labels), True)
    if snap["references"]:
        _heading(model, words["references"])
        for reference in snap["references"]:
            _heading(model, reference["identifier"] + "  " + reference["title"], 2)
            details = [reference.get(k) for k in ("owner", "deliverable_type", "stage", "source_period") if reference.get(k)]
            _para(model, " · ".join(details), True)
            _para(model, reference.get("rationale", ""))
            _link(model, reference.get("url", ""), reference.get("url"))
    if snap["cases"]:
        _heading(model, words["cases"])
        _para(model, words["case_records"], True)
        for case in snap["cases"]:
            number = case_numbers[case["id"]]
            title = case.get("title_fr") if language == "fr" else case.get("title")
            _heading(model, words["case"] + " " + str(number) + "  " + (title or case["id"]), 2, _anchor(case["id"]))
            for key, label in (("review_questions", "questions"), ("expected_observations", "observations")):
                values = case.get(key + "_fr" if language == "fr" else key, [])
                if values:
                    _heading(model, words[label], 3)
                    _list(model, values)
            source_rows = case.get("source_rows", [])
            selected = [row for row in source_rows if isinstance(row, dict) and row.get("language") in (None, "", language)]
            locators = []
            for row in selected:
                locator = row.get("csv_record", row.get("record_number", row.get("row", "")))
                locators.append([row.get("source_id", ""), locator, row.get("row_sha256", "")])
            _table(model, [words["sources"], words["record"], "SHA256"], locators, [2, 1, 3])
            if case.get("source_ids"):
                _para(model, words["sources"] + ": " + _source_citations(case["source_ids"], source_numbers), True)
    _heading(model, words["source_register"])
    for source in snap["sources"]:
        _heading(model, "[" + str(source_numbers[source["id"]]) + "]  " + source["title"], 2)
        details = [source.get(k) for k in ("publisher", "version", "language", "representation") if source.get(k)]
        _para(model, " · ".join(details))
        _link(model, source.get("url", ""), source.get("url"))
        if source.get("license"):
            _para(model, source["license"], True)
        if source.get("sha256"):
            _para(model, "SHA256: " + source["sha256"], True)
    _para(model, words["edition_id"] + ": " + edition["id"] + "\n" +
          words["digest"] + ": " + edition["sha256"], True)
    if edition.get("source_edition_sha256") and edition["source_edition_sha256"] != edition["sha256"]:
        _para(model, ("Original frozen edition SHA256: " if language == "en" else
                     "SHA256 de l’édition originale consignée : ") + edition["source_edition_sha256"], True)
    return model



def _case_rows_model(edition, language):
    words = WORDS[language]
    title = "Original source case records" if language == "en" else "Enregistrements sources des cas pratiques"
    model = _model(title, edition["label"] + " · " + words["language"], language, edition["frozen_at"])
    _para(model, "The following cells retain their original source names, record numbers and string values."
          if language == "en" else
          "Les cellules ci-dessous conservent leurs noms de champs, numéros d’enregistrement et valeurs textuelles d’origine.")
    for case in edition["snapshot"]["cases"]:
        label = case.get("title_fr") if language == "fr" else case.get("title")
        _heading(model, label or case["id"], 1, _anchor(case["id"]))
        _para(model, case["id"], True)
        groups = {}
        for row in case.get("source_rows", []):
            if isinstance(row, dict) and row.get("language") in (None, "", language):
                groups.setdefault(row.get("source_id", ""), []).append(row)
        for source_id, rows in groups.items():
            _heading(model, source_id or words["sources"], 2)
            for offset in range(0, len(rows), 3):
                chunk = rows[offset:offset + 3]
                if all(isinstance(row.get("cells"), list) for row in chunk):
                    fields = {}
                    values = []
                    for row in chunk:
                        record = {}
                        for cell in row["cells"]:
                            key = (_text(cell.get("column", "")), _text(cell.get("name", "")))
                            fields.setdefault(key, None)
                            if key in record:
                                raise PublicationError("A retained case record repeats the same source column.")
                            record[key] = _text(cell.get("value"))
                        values.append(record)
                    headers = [("Original field" if language == "en" else "Champ original")] + [
                        words["record"] + " " + _text(row.get("csv_record", row.get("record_number", ""))) for row in chunk]
                    data = [[number + "  " + name] + [record.get((number, name), "") for record in values]
                            for number, name in fields]
                    _table(model, headers, data, [2.3] + [1] * len(chunk))
                else:
                    for row in chunk:
                        _table(model, ["Field" if language == "en" else "Champ", "Value" if language == "en" else "Valeur"],
                               [[key, _display(value)] for key, value in row.items()], [1, 3])
                for row in chunk:
                    _para(model, words["record"] + " " + _text(row.get("csv_record", row.get("record_number", ""))) +
                          " · SHA256 " + _text(row.get("row_sha256")), True)
    return model


def _comparison(previous, edition):
    before = _public_snapshot(previous["snapshot"])
    after = _public_snapshot(edition["snapshot"])
    detail = engine.compare_editions(before, after)
    for key in ("cases", "language_reviews"):
        detail[key + "_changed"] = engine.canonical(before[key]) != engine.canonical(after[key])
    return {"before": _metadata(previous), "after": _metadata(edition), "changes": detail,
            "before_snapshot": before, "after_snapshot": after,
            "non_public_fields_changed": engine.canonical(previous["snapshot"]) != engine.canonical(edition["snapshot"])
                                         and engine.canonical(before) == engine.canonical(after)}


def _diff_words(before, after):
    left, right = re.findall(r"\s+|\S+", _text(before)), re.findall(r"\s+|\S+", _text(after))
    parts = []
    for tag, a, b, c, d in difflib.SequenceMatcher(None, left, right, autojunk=False).get_opcodes():
        if tag in ("equal", "delete", "replace") and a != b:
            parts.append({"text": "".join(left[a:b]), "change": "same" if tag == "equal" else "delete"})
        if tag in ("insert", "replace") and c != d:
            parts.append({"text": "".join(right[c:d]), "change": "insert"})
    return parts


def _comparison_model(comparison):
    a, b = comparison["before"], comparison["after"]
    model = _model("Edition comparison", a["label"] + " to " + b["label"], "en", b["frozen_at"])
    _para(model, "Deleted text is struck through. Added text is underlined. English and French changes remain separate.")
    _table(model, ["Edition", "Identifier", "SHA256"],
           [[a["label"], a["edition_id"], a["edition_sha256"]],
            [b["label"], b["edition_id"], b["edition_sha256"]]], [1, 2, 3])
    changes = comparison["changes"]
    for change in changes["clauses"]:
        old, new = change.get("before", {}), change.get("after", {})
        row = new or old
        _heading(model, row["number"] + "  " + row["title_en"])
        _para(model, change["kind"].capitalize() + " · " + row["id"], True)
        for field, label in (("title_en", "English title"), ("body_en", "English text"),
                             ("title_fr", "Titre français"), ("body_fr", "Texte français")):
            if old.get(field, "") != new.get(field, ""):
                _heading(model, label, 2)
                model["blocks"].append({"kind": "diff", "parts": _diff_words(old.get(field, ""), new.get(field, ""))})
        for field, label in (("table_en", "English table"), ("table_fr", "Tableau français")):
            if old.get(field, []) != new.get(field, []):
                _heading(model, label, 2)
                for basis, value in (("Before", old.get(field, [])), ("After", new.get(field, []))):
                    _para(model, basis, True)
                    if value:
                        _table(model, value[0], value[1:] or [[""] * len(value[0])])
        metadata = []
        for field in ("number", "parent_id", "kind", "source_ids", "case_ids"):
            if engine.canonical(old.get(field)) != engine.canonical(new.get(field)):
                metadata.append([field, _display(old.get(field)), _display(new.get(field))])
        _table(model, ["Field", "Before", "After"], metadata, [1, 2, 2])
    for key, label in (("scope", "Scope"), ("references", "Related references"), ("sources", "Source versions"),
                       ("cases", "Worked cases"), ("language_reviews", "Language review records")):
        left, right = comparison["before_snapshot"][key], comparison["after_snapshot"][key]
        if engine.canonical(left) == engine.canonical(right):
            continue
        _heading(model, label)
        if isinstance(left, dict):
            for field in sorted(set(left) | set(right)):
                if engine.canonical(left.get(field)) != engine.canonical(right.get(field)):
                    _heading(model, field.replace("_", " "), 2)
                    model["blocks"].append({"kind": "diff", "parts": _diff_words(_display(left.get(field)), _display(right.get(field)))})
        else:
            identity = "clause_id" if key == "language_reviews" else "id"
            old = {r[identity]: r for r in left}
            new = {r[identity]: r for r in right}
            for item in sorted(set(old) | set(new)):
                if engine.canonical(old.get(item)) == engine.canonical(new.get(item)):
                    continue
                _heading(model, item, 2)
                x, y = old.get(item, {}), new.get(item, {})
                for field in sorted(set(x) | set(y)):
                    if engine.canonical(x.get(field)) == engine.canonical(y.get(field)):
                        continue
                    if field == "source_rows":
                        _table(model, ["Original case records", "Before", "After"],
                               [["Count", len(x.get(field, [])), len(y.get(field, []))],
                                ["SHA256", engine.digest(x.get(field, [])), engine.digest(y.get(field, []))]], [1, 2, 2])
                    else:
                        _heading(model, field.replace("_", " "), 3)
                        model["blocks"].append({"kind": "diff", "parts": _diff_words(_display(x.get(field)), _display(y.get(field)))})
    if not changes["clauses"] and not any(changes.get(k) for k in
        ("scope_changed", "references_changed", "source_versions_changed", "cases_changed", "language_reviews_changed")):
        _para(model, "The selected public text, source references, cases and language review records are unchanged.")
    if comparison["non_public_fields_changed"]:
        _para(model, "The edition identity differs because non-public editorial fields changed. Those fields remain in the private workspace.")
    return model


def _display(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list) and all(isinstance(x, (str, int, float)) for x in value):
        return "\n".join(str(x) for x in value)
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2)


def _latest(rows, *fields):
    return sorted(rows, key=lambda r: tuple(_text(r.get(k)) for k in fields) + (_text(r.get("id")),))[-1] if rows else None



def _field_label(name):
    labels = {"id": "Record ID", "source_ids": "Source IDs", "input_sha256": "Retained input SHA256",
              "edition_id": "Edition ID", "ballot_id": "Ballot ID", "ola_events": "Official Languages Act",
              "scope": "Scope", "selected_procedure": "Selected development procedure",
              "technical_decisions": "Technical decisions", "ballot_decisions": "Ballot decisions",
              "consultations": "Consultations", "communications": "Communication tasks",
              "delivery_register": "Scope and delivery register", "context_scope": "Context coverage",
              "scc_response": "Recorded SCC response", "scope_binding_current": "Scope binding current",
              "minimum_business_days": "Minimum business days", "next_due": "Next report due",
              "last_report_submitted_on": "Last report submitted on",
              "recorded_closeout_conditions_met": "Recorded closeout conditions met",
              "accepted_reports_with_received_deliverables": "Accepted reports with received deliverables",
              "conditions_recorded": "Cash release conditions recorded",
              "provisional_review_ready_on": "Provisional review ready date",
              "mandatory_notice_exception_supported": "Mandatory-notice exception supported",
              "planning_due_before_first_phase_report": "Planning due date before first phase report",
              "final_invoice_register_confirmed": "Final invoice register confirmed"}
    return labels.get(name, name.replace("_", " ").capitalize())


def _fact_value(value):
    if value is None or value == "":
        return "Not recorded"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        if not value:
            return "No entries recorded"
        if all(not isinstance(item, (list, dict)) for item in value):
            return ", ".join(_text(item) for item in value)
    return _display(value)


def _facts(model, value, *, omit=()):
    if not isinstance(value, dict):
        _para(model, _fact_value(value))
        return
    rows = [[_field_label(key), _fact_value(item)] for key, item in value.items() if key not in omit]
    if rows:
        _table(model, ["Recorded field", "Value"], rows, [2, 4])
    else:
        _para(model, "No fields recorded.")


def _report_context(model, context, title):
    _heading(model, title, 2)
    if context is None:
        _para(model, "Structured development-process, OLA and communication context was not captured in this input. "
              "The original narrative and records remain unchanged; current project records are shown separately.")
        return
    if not isinstance(context, dict):
        raise PublicationError("A retained phase report context must be an object.")
    _para(model, "Context SHA256 " + engine.digest(context), True)
    for name, value in context.items():
        _heading(model, _field_label(name), 3)
        if name == "scope" and isinstance(value, dict):
            # The complete original scope remains in the same retained snapshot.
            fields = ("id", "title_en", "title_fr", "deliverable_type", "roadmap_issues",
                      "procedure_id", "source_ids")
            _facts(model, {key: value[key] for key in fields if key in value})
        elif isinstance(value, list):
            if not value:
                _para(model, "No entries captured.")
            for index, item in enumerate(value):
                if isinstance(item, dict):
                    label = item.get("title") or item.get("title_en") or item.get("id") or str(index + 1)
                    _para(model, label)
                    _facts(model, item)
                else:
                    _para(model, _fact_value(item))
        else:
            _facts(model, value)


def _communication_workpaper(model, analysis, phase):
    rows = [row for row in analysis.get("communication_reviews", []) if row.get("phase") == phase]
    if not rows:
        return
    _heading(model, "Current SCC communications review", 2)
    _para(model, "Calculated dates describe the selected recorded calendar and submissions. "
          "A provisional date does not stand in for an SCC response.")
    for row in rows:
        _heading(model, row.get("title") or row["id"], 3)
        _facts(model, row)


def _phase_model(state, phase, phase_analysis):
    title = engine.PHASE_NAMES[phase]
    model = _model("Phase " + phase + "  " + title, state["title"])
    current = engine.phase_snapshot(state, phase)
    funding = current["funding"]
    _para(model, "Current phase workpaper. All money is CAD. Net delivery amounts and HST remain separate.")
    if funding:
        _para(model, "Funding record: " + funding["title"] + " · " + funding["status"] +
              ". Phase base amount: CAD " + _money(funding["phase_amounts"][phase]) + ".")
    else:
        _para(model, "A funding agreement is not recorded in the selected workspace.")
    _report_context(model, current.get("report_context"), "Current development, OLA and communications basis")
    _communication_workpaper(model, phase_analysis, phase)
    cash = phase_analysis.get("cash_release_conditions")
    if cash is not None:
        _heading(model, "Current SCC phase cash conditions", 2)
        _facts(model, cash)
    costs = current["costs"]
    rows = []
    for basis in ("estimate", "actual"):
        rows.append([basis.capitalize(), _money(_amount(costs, "net_cad", basis=basis, category="base")),
                     _money(_amount(costs, "net_cad", basis=basis, category="ola")),
                     _money(_amount(costs, "hst_cad", basis=basis))])
    _table(model, ["Cost basis", "Base net CAD", "OLA net CAD", "HST CAD"], rows, [1.4, 1, 1, 1], (1, 2, 3))
    _heading(model, "Work and deliverables")
    _table(model, ["Task", "Owner", "Due", "Status"],
           [[r["title"], r["owner"], r["due"], r["status"]] for r in current["tasks"]], [3, 2, 1.2, 1.5])
    _table(model, ["Deliverable", "Due", "Status", "Acceptance recorded"],
           [[r["title"], r["due"], r["status"], r["accepted_at"]] for r in current["deliverables"]], [3, 1.2, 1.4, 2])
    _heading(model, "Source costs")
    for cost in costs:
        _para(model, cost["title"] + " · " + cost["id"])
        _table(model, ["Basis and category", "Net CAD", "HST CAD", "Date"],
               [[cost["basis"] + " / " + cost["category"], _money(cost["net_cad"]), _money(cost["hst_cad"]), cost["incurred"]]],
               [2, 1, 1, 1.4], (1, 2))
        _para(model, "Substantiation: " + (", ".join(cost["substantiation_ids"]) or "Not recorded") +
              (". External reference: " + cost["external_reference"] if cost["external_reference"] else ""), True)
    reports = sorted((r for r in state["phase_reports"] if r["phase"] == phase),
                     key=lambda r: (r["recorded_at"], r["id"]))
    _heading(model, "Retained phase reports")
    if not reports:
        _para(model, "No phase report has been captured for this phase.")
    for report in reports:
        if engine.digest(report["input_snapshot"]) != report["input_sha256"]:
            raise PublicationError("A retained phase report differs from its snapshot hash: " + report["id"])
        _heading(model, report["title"], 2)
        _para(model, report["summary"])
        _para(model, "Report " + report["id"] + " · through " + (report["through"] or "date not recorded") +
              " · captured " + report["recorded_at"], True)
        _para(model, "Submitted: " + (report["submitted_at"] or "Not recorded") +
              ". Current phase " + ("differs from" if engine.digest(current) != report["input_sha256"] else "matches") +
              " this retained report.")
        snap = report["input_snapshot"]
        retained_context = snap.get("report_context")
        if retained_context is not None and retained_context == current.get("report_context"):
            _para(model, "The retained development, OLA and communication context matches the current context shown above. "
                  "Retained context SHA256 " + engine.digest(retained_context) + ".", True)
        else:
            _report_context(model, retained_context, "Retained development, OLA and communications basis")
        receipt_conditions = next((row.get("delivery_receipt_conditions") for row in
                                   phase_analysis.get("reports", []) if row.get("report_id", row.get("id")) == report["id"]), None)
        if receipt_conditions is not None:
            _heading(model, "Recorded deliverable receipt conditions", 3)
            _facts(model, receipt_conditions)
        retained_funding = snap.get("funding")
        if retained_funding:
            _para(model, "Retained funding: " + retained_funding["title"] + " · " + retained_funding["status"] +
                  ". Phase base CAD " + _money(retained_funding["phase_amounts"][phase]) +
                  "; OLA contingency CAD " + _money(retained_funding["ola_cad"]) + ".")
        _table(model, ["Retained deliverable", "Due", "Status", "Acceptance recorded"],
               [[r["title"], r["due"], r["status"], r["accepted_at"]] for r in snap["deliverables"]], [3, 1.2, 1.4, 2])
        _table(model, ["Retained cost", "Basis", "Net CAD", "HST CAD"],
               [[c["title"], c["basis"] + "/" + c["category"], _money(c["net_cad"]), _money(c["hst_cad"])]
                for c in snap["costs"]], [3, 1.5, 1, 1], (2, 3))
        responses = sorted((r for r in state["report_responses"] if r["report_id"] == report["id"]),
                           key=lambda r: (r["date"] or "", r["id"]))
        if not responses:
            _para(model, "Report response: Not recorded.")
        for response in responses:
            _para(model, "Recorded response: " + response["decision"] + " · " + (response["date"] or "") +
                  " · " + response["authority"] + ". " + response["rationale"])
            _para(model, "Evidence: " + ", ".join(response["evidence_ids"]) + " · report SHA256 " +
                  response["report_sha256"], True)
        _para(model, "Retained input SHA256 " + report["input_sha256"], True)
    invoices = [r for r in state["invoices"] if r["phase"] == phase]
    _heading(model, "Invoices and recorded cash")
    balances = {r["invoice_id"]: r for r in engine.payment_balances(state)}
    if not invoices:
        _para(model, "No invoice is recorded for this phase.")
    for invoice in invoices:
        if engine.digest(invoice["input_snapshot"]) != invoice["input_sha256"]:
            raise PublicationError("A retained invoice differs from its snapshot hash: " + invoice["id"])
        balance = balances[invoice["id"]]
        _heading(model, "Invoice " + invoice["number"], 2)
        _para(model, "Status: " + balance["status"] + " · issued " + (invoice["issued"] or "Not recorded") +
              " · report " + invoice["report_id"])
        _table(model, ["Net CAD", "HST CAD", "Total CAD", "Collected CAD", "Remaining CAD"],
               [[_money(invoice["net_cad"]), _money(invoice["hst_cad"]), _money(invoice["total_cad"]),
                 _money(balance["collected_cad"]), _money(balance["remaining_cad"])]], numeric=range(5))
        _table(model, ["Source cost", "Category", "Allocated net CAD"],
               [[r["cost_id"], r["category"], _money(r["net_cad"])] for r in invoice["allocations"]],
               [3, 1.5, 1.5], (2,))
        _para(model, invoice.get("terms", ""))
        payments = [r for r in state["payments"] if r["invoice_id"] == invoice["id"]]
        _table(model, ["Cash record", "Date", "CAD", "External reference"],
               [[r["kind"], r["date"], _money(r["amount_cad"]), r["external_reference"]] for r in payments],
               [1.2, 1.2, 1, 3], (2,))
        _para(model, "Invoice " + invoice["id"] + " · retained input SHA256 " + invoice["input_sha256"], True)
    return model


def _management_model(state, analysis):
    model = _model("SCC delivery and financial workpapers", state["title"])
    funding = engine.single(state, "funding")
    _para(model, "Four-phase workpapers for the selected project state. The following records identify scope delivery, "
          "retained phase reports, report responses, invoice allocations and recorded cash. Amounts are CAD with HST shown separately.")
    _para(model, "Project mode: " + state["mode"] + " · funding status: " + analysis["funding_status"] +
          " · next progress report due: " + (analysis.get("next_progress_report_due") or "Not recorded") + ".")
    if funding:
        _heading(model, "Recorded funding")
        _table(model, ["Phase", "Base amount CAD"],
               [[p + "  " + engine.PHASE_NAMES[p], _money(funding["phase_amounts"][p])] for p in engine.PHASES],
               [4, 1], (1,))
        _table(model, ["Funding component", "CAD"],
               [["Base delivery", _money(funding["base_cad"])], ["OLA contingency", _money(funding["ola_cad"])],
                ["Maximum recorded net amount", _money(Decimal(funding["base_cad"]) + Decimal(funding["ola_cad"]))]],
               [4, 1], (1,))
        _para(model, funding["title"] + " · " + funding["status"] + " · agreement sources: " +
              (", ".join(funding["agreement_source_ids"]) or "Not recorded"), True)
    if analysis.get("reporting_schedule") is not None:
        _heading(model, "Current reporting calendar")
        _facts(model, analysis["reporting_schedule"])
    phases = {r["phase"]: r for r in analysis["phases"]}
    for phase in engine.PHASES:
        phase_analysis = dict(phases.get(phase, {}))
        phase_analysis["communication_reviews"] = analysis.get("communication_reviews", [])
        child = _phase_model(state, phase, phase_analysis)
        _heading(model, child["title"])
        model["blocks"].extend(child["blocks"])
    if analysis.get("review_link_dispatches"):
        _heading(model, "Current public-review launch handoffs")
        for dispatch in analysis["review_link_dispatches"]:
            _facts(model, dispatch)
    if analysis.get("completion") is not None:
        _heading(model, "Current scope and final-invoice closeout")
        _facts(model, analysis["completion"])
    _heading(model, "Comment disposition and receiving work")
    comments = sorted(state["comments"], key=lambda r: (r["edition_id"], r["clause_id"], r["id"]))
    if not comments:
        _para(model, "No consultation comments are recorded.")
    for comment in comments:
        _heading(model, comment["id"] + "  " + comment["kind"], 2)
        _para(model, "Edition " + comment["edition_id"] + " · clause " + (comment["clause_id"] or "General") +
              " · disposition " + comment["disposition"], True)
        _para(model, comment["body"])
        _para(model, "Rationale: " + (comment["rationale"] or "Not recorded"))
        _para(model, "Resolution: " + (comment["resolution_edition_id"] or "Not recorded") + " / " +
              (comment["resolution_clause_id"] or "Not recorded") + ". Notification: " +
              (comment["notification_at"] or "Not recorded") + ".", True)
    _heading(model, "Recorded technical and procedural decisions")
    if not state["ballot_closures"]:
        _para(model, "No ballot closure is recorded.")
    for closure in state["ballot_closures"]:
        _heading(model, closure["ballot_id"] + "  " + closure["recorded_decision"], 2)
        _para(model, closure["authority"] + " · " + closure["recorded_at"] + ". " + closure["rationale"])
        _para(model, "Retained ballot input SHA256 " + closure["input_sha256"], True)
        _para(model, "Recorded calculation\n" + _display(closure["calculation"]), True)
        for review in [r for r in state["procedural_reviews"] if r["closure_id"] == closure["id"]]:
            _para(model, "Procedural review " + review["id"] + ": " + review["decision"] +
                  ". " + review["rationale"])
            _para(model, "Reviewer " + review["reviewer_id"] + " · reviewed input SHA256 " +
                  review["reviewed_input_sha256"], True)
    _heading(model, "Project changes")
    for change in state["changes"]:
        _heading(model, change["title"], 2)
        _para(model, change["kind"] + " · " + change["decision"] + ". " + change["rationale"])
        _table(model, ["Original due", "Proposed due", "Base change CAD", "OLA change CAD"],
               [[change["original_due"], change["new_due"], _money(change["base_change_cad"]),
                 _money(change["ola_change_cad"])]], [1, 1, 1, 1], (2, 3))
        _para(model, "Authority: " + (change["authority"] or "Not recorded") +
              " · decision date: " + (change["decision_date"] or "Not recorded"), True)
    _heading(model, "Receiving follow-up")
    _table(model, ["Area", "Record", "Recorded finding"],
           [[r.get("collection", ""), r.get("record_id", ""), r["message"]] for r in analysis["findings"]],
           [1.2, 1.8, 4])
    _para(model, "Selected project state SHA256 " + engine.digest(state), True)
    return model


def _html(model):
    parts = ['<!doctype html><html lang="' + ("fr-CA" if model["language"] == "fr" else "en-CA") + '"><head>',
             '<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
             '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'">',
             "<title>" + escape(model["title"]) + "</title><style>" + CSS + "</style></head><body>",
             "<h1>" + escape(model["title"]) + "</h1>"]
    if model["subtitle"]:
        parts.append('<p class="subtitle">' + escape(model["subtitle"]) + "</p>")
    for block in model["blocks"]:
        kind = block["kind"]
        if kind == "heading":
            level = block["level"] + 1
            identifier = ' id="' + block["anchor"] + '"' if block.get("anchor") else ""
            parts.append(f"<h{level}{identifier}>" + escape(block["text"]) + f"</h{level}>")
        elif kind == "paragraph":
            parts.append('<p' + (' class="small"' if block["small"] else "") + ">" +
                         escape(block["text"]).replace("\n", "<br>") + "</p>")
        elif kind == "list":
            parts.append("<ul>" + "".join("<li>" + escape(x) + "</li>" for x in block["items"]) + "</ul>")
        elif kind == "link":
            if block["url"]:
                parts.append('<p class="small"><a rel="noopener noreferrer" href="' + escape(block["url"], quote=True) +
                             '">' + escape(block["text"]) + "</a></p>")
            elif block["text"]:
                parts.append('<p class="small">' + escape(block["text"]) + "</p>")
        elif kind == "toc":
            parts.append('<ol class="toc">' + "".join('<li><a href="#' + x["anchor"] + '">' +
                         escape(x["text"]) + "</a></li>" for x in block["items"]) + "</ol>")
        elif kind == "diff":
            pieces = []
            for piece in block["parts"]:
                text = escape(piece["text"]).replace("\n", "<br>")
                tag = {"delete": "del", "insert": "ins"}.get(piece["change"])
                pieces.append(("<" + tag + ">" + text + "</" + tag + ">") if tag else text)
            parts.append("<p>" + "".join(pieces) + "</p>")
        elif kind == "table":
            total = sum(block["weights"])
            parts.append("<table><colgroup>" + "".join('<col style="width:' + str(w / total * 100) + '%">'
                         for w in block["weights"]) + "</colgroup><thead><tr>")
            for i, value in enumerate(block["headings"]):
                parts.append('<th scope="col"' + (' class="number"' if i in block["numeric"] else "") + ">" + escape(value) + "</th>")
            parts.append("</tr></thead><tbody>")
            for row in block["rows"]:
                parts.append("<tr>" + "".join('<td' + (' class="number"' if i in block["numeric"] else "") + ">" +
                             escape(value).replace("\n", "<br>") + "</td>" for i, value in enumerate(row)) + "</tr>")
            parts.append("</tbody></table>")
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")


def _docx(model):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        raise PublicationError("DOCX output requires python-docx 1.2.0 or a compatible installed version.") from error
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.72)
    section.footer_distance = Inches(.32)
    styles = document.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.1
        rp = style.element.get_or_add_rPr()
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "fr-CA" if model["language"] == "fr" else "en-CA")
        rp.append(language)
    styles["Title"].font.size, styles["Title"].font.bold = Pt(21), True
    styles["Subtitle"].font.size = Pt(11)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        styles[name].font.size, styles[name].font.bold = Pt(size), True
        styles[name].paragraph_format.space_before = Pt(16 if name == "Heading 1" else 10)
        styles[name].paragraph_format.keep_with_next = True
    cp = document.core_properties
    cp.title, cp.subject = model["title"], model["subtitle"]
    cp.author, cp.last_modified_by = "Standards drafting workspace", ""
    cp.language = "fr-CA" if model["language"] == "fr" else "en-CA"
    try:
        instant = datetime.fromisoformat(model["timestamp"].replace("Z", "+00:00")).astimezone(timezone.utc)
    except (ValueError, TypeError):
        instant = datetime(2000, 1, 1, tzinfo=timezone.utc)
    cp.created = cp.modified = instant.replace(tzinfo=None)
    document.add_paragraph(model["title"], "Title")
    if model["subtitle"]:
        document.add_paragraph(model["subtitle"], "Subtitle")
    bookmark_id = 0

    def link(paragraph, label, target, internal=False):
        h = OxmlElement("w:hyperlink")
        if internal:
            h.set(qn("w:anchor"), target)
        else:
            relation = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
            h.set(qn("r:id"), paragraph.part.relate_to(target, relation, is_external=True))
        r = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "000000")
        underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
        props.extend([color, underline])
        if not internal:
            size = OxmlElement("w:sz"); size.set(qn("w:val"), "18"); props.append(size)
        r.append(props)
        text = OxmlElement("w:t"); text.text = label; r.append(text)
        h.append(r); paragraph._p.append(h)

    for block in model["blocks"]:
        kind = block["kind"]
        if kind == "heading":
            p = document.add_paragraph(block["text"], "Heading " + str(block["level"]))
            if block.get("anchor"):
                start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), block["anchor"])
                end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_id))
                p._p.insert(1 if p._p.pPr is not None else 0, start); p._p.append(end); bookmark_id += 1
        elif kind in ("paragraph", "link"):
            p = document.add_paragraph()
            if kind == "link" and block["url"]:
                link(p, block["text"], block["url"])
            else:
                p.add_run(block["text"])
            if block.get("small") or kind == "link":
                for run in p.runs:
                    run.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(5)
        elif kind == "list":
            for item in block["items"]:
                document.add_paragraph(item, "List Bullet")
        elif kind == "toc":
            for item in block["items"]:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Inches(.18 * (item["level"] - 1))
                link(p, item["text"], item["anchor"], True)
        elif kind == "diff":
            p = document.add_paragraph()
            for piece in block["parts"]:
                run = p.add_run(piece["text"])
                if piece["change"] == "delete":
                    run.font.strike = True
                elif piece["change"] == "insert":
                    run.font.underline = True
                    run.font.bold = True
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["headings"]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [Inches(7.06 * w / sum(block["weights"])) for w in block["weights"]]
            for column, width in zip(table.columns, widths):
                column.width = width
            for i, value in enumerate(block["headings"]):
                table.rows[0].cells[i].text = value
            for row in block["rows"]:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = value
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row_index, row in enumerate(table.rows):
                for i, cell in enumerate(row.cells):
                    cell.width = widths[i]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tc = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        line = OxmlElement("w:" + edge)
                        for k, v in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            line.set(qn("w:" + k), v)
                        borders.append(line)
                    tc.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "bottom", "left", "right"):
                        margin = OxmlElement("w:" + edge)
                        margin.set(qn("w:w"), "85" if edge in ("top", "bottom") else "100")
                        margin.set(qn("w:type"), "dxa"); margins.append(margin)
                    tc.append(margins)
                    if row_index == 0:
                        shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E9E9E9"); tc.append(shading)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1.05
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in block["numeric"] else WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.bold = row_index == 0
            after = document.add_paragraph()
            after.paragraph_format.space_after = Pt(2)
            after.paragraph_format.line_spacing = 1
            after.add_run().font.size = Pt(1)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(("Page " if model["language"] == "en" else "Page ") + "")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    output = BytesIO(); document.save(output)
    # Source-bound content produces stable bytes across repeated downloads.
    normalized = BytesIO()
    with zipfile.ZipFile(output) as original, zipfile.ZipFile(normalized, "w", zipfile.ZIP_DEFLATED) as final:
        for name in sorted(original.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            final.writestr(info, original.read(name), compress_type=zipfile.ZIP_DEFLATED)
    return normalized.getvalue()


def _pdf(model):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError as error:
        raise PublicationError("PDF output requires ReportLab 4.4.9 or a compatible installed version.") from error
    with FONT_LOCK:
        for name, file in zip(FONT_NAMES, ("Vera.ttf", "VeraBd.ttf", "VeraIt.ttf", "VeraBI.ttf")):
            if name not in pdfmetrics.getRegisteredFontNames():
                path = FONT_DIR / file
                if not path.is_file():
                    import reportlab
                    path = Path(reportlab.__file__).resolve().parent / "fonts" / file
                if not path.is_file():
                    raise PublicationError("The installed ReportLab package is missing its original embedded font: " + file)
                pdfmetrics.registerFont(TTFont(name, str(path)))
        pdfmetrics.registerFontFamily(FONT_NAMES[0], normal=FONT_NAMES[0], bold=FONT_NAMES[1],
                                      italic=FONT_NAMES[2], boldItalic=FONT_NAMES[3])
    base = ParagraphStyle("Body", fontName=FONT_NAMES[0], fontSize=10.5, leading=14.5,
                          spaceAfter=8, splitLongWords=True, allowWidows=0, allowOrphans=0)
    styles = {
        "body": base,
        "small": ParagraphStyle("Small", parent=base, fontSize=8, leading=10.8, spaceAfter=5),
        "title": ParagraphStyle("Title", parent=base, fontName=FONT_NAMES[1], fontSize=20, leading=25, spaceAfter=14),
        "subtitle": ParagraphStyle("Subtitle", parent=base, fontSize=11, leading=15, spaceAfter=16),
        "heading1": ParagraphStyle("Heading1", parent=base, fontName=FONT_NAMES[1], fontSize=14, leading=18,
                                   spaceBefore=16, spaceAfter=9, keepWithNext=True),
        "heading2": ParagraphStyle("Heading2", parent=base, fontName=FONT_NAMES[1], fontSize=12, leading=16,
                                   spaceBefore=11, spaceAfter=7, keepWithNext=True),
        "heading3": ParagraphStyle("Heading3", parent=base, fontName=FONT_NAMES[1], fontSize=10.5, leading=14.5,
                                   spaceBefore=8, spaceAfter=5, keepWithNext=True),
    }
    def safe(text):
        return escape(_text(text), quote=False).replace("\n", "<br/>")
    story = [Paragraph(safe(model["title"]), styles["title"])]
    if model["subtitle"]:
        story.append(Paragraph(safe(model["subtitle"]), styles["subtitle"]))
    for block in model["blocks"]:
        kind = block["kind"]
        if kind == "heading":
            prefix = '<a name="' + block["anchor"] + '"/>' if block.get("anchor") else ""
            story.append(Paragraph(prefix + safe(block["text"]), styles["heading" + str(block["level"])]))
        elif kind == "paragraph":
            story.append(Paragraph(safe(block["text"]), styles["small" if block["small"] else "body"]))
        elif kind == "list":
            for item in block["items"]:
                story.append(Paragraph("• " + safe(item), base))
        elif kind == "link":
            text = safe(block["text"])
            if block["url"]:
                text = '<a color="#000000" href="' + escape(block["url"], quote=True) + '">' + text + "</a>"
            if text:
                story.append(Paragraph(text, styles["small"]))
        elif kind == "toc":
            for item in block["items"]:
                story.append(Paragraph('<a color="#000000" href="#' + item["anchor"] + '">' +
                                       safe(item["text"]) + "</a>", base))
        elif kind == "diff":
            result = []
            for piece in block["parts"]:
                tag = {"delete": "strike", "insert": "u"}.get(piece["change"])
                result.append(("<" + tag + ">" + safe(piece["text"]) + "</" + tag + ">") if tag else safe(piece["text"]))
            story.append(Paragraph("".join(result) or " ", base))
        elif kind == "table":
            content = []
            for row_index, row in enumerate([block["headings"]] + block["rows"]):
                cells = []
                for i, text in enumerate(row):
                    style = ParagraphStyle("Cell", parent=base, fontName=FONT_NAMES[1] if row_index == 0 else FONT_NAMES[0],
                                           fontSize=8.1, leading=10.5, spaceAfter=0,
                                           alignment=TA_RIGHT if i in block["numeric"] else TA_LEFT)
                    cells.append(Paragraph(safe(text) or " ", style))
                content.append(cells)
            weights = block["weights"]
            table = Table(content, colWidths=[508 * x / sum(weights) for x in weights], repeatRows=1,
                          hAlign="LEFT", splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9E9E9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 10)])
    class StableCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)
    def page(canvas, document):
        canvas.saveState()
        canvas.setFont(FONT_NAMES[0], 8)
        canvas.drawRightString(letter[0] - 52, 25, "Page " + str(document.page))
        canvas.restoreState()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter, leftMargin=52, rightMargin=52,
                                 topMargin=52, bottomMargin=46, title=model["title"],
                                 author="Standards drafting workspace", subject=model["subtitle"])
    document.build(story, onFirstPage=page, onLaterPages=page, canvasmaker=StableCanvas)
    return output.getvalue()


def _document_files(prefix, model):
    return {prefix + ".html": _html(model), prefix + ".docx": _docx(model), prefix + ".pdf": _pdf(model)}


def _manifest(files, context):
    return _json({**context, "files": [
        {"path": path, "bytes": len(data), "sha256": sha256(data).hexdigest()} for path, data in sorted(files.items())
    ]})


def edition_files(edition, *, publication=None, comparison_edition=None):
    """Return the selected immutable edition in both languages and its comparison."""
    edition = _frozen(edition)
    def public_rows(selected):
        snapshot = selected["snapshot"]
        sources = snapshot["sources"]
        if any(source.get("citation_visibility") != "public" for source in sources):
            raise PublicationError("Project the selected edition to explicitly public source references before publication.")
        public_ids = {source["id"] for source in sources}
        for record in [snapshot["scope"]] + snapshot["clauses"] + snapshot["references"] + snapshot["cases"]:
            if any(source_id not in public_ids for source_id in record.get("source_ids", [])):
                raise PublicationError("Every published citation must reference an explicitly public source record.")
        for case in snapshot["cases"]:
            for row in case.get("source_rows", []):
                if not isinstance(row, dict) or row.get("source_id") not in public_ids:
                    raise PublicationError("Every published case row must be bound to an explicitly public source record.")
    public_rows(edition)
    if comparison_edition is not None:
        comparison_edition = _frozen(comparison_edition)
        public_rows(comparison_edition)
    if publication is not None:
        publication = deepcopy(publication)
        if publication.get("edition_id") != edition["id"]:
            raise PublicationError("The publication record belongs to another edition.")
    context = _metadata(edition, publication)
    files = {}
    for language in ("en", "fr"):
        files.update(_document_files("publication/standard-" + language, _standard_model(edition, publication, language)))
        files["publication/case-records-" + language + ".html"] = _html(_case_rows_model(edition, language))
    public = _public_snapshot(edition["snapshot"])
    files["publication/edition.json"] = _json({**context, "export_snapshot_sha256": engine.digest(public),
                                              "snapshot": public})
    files["publication/source-register.json"] = _json(public["sources"])
    files["publication/case-records.json"] = _json(public["cases"])
    files["publication/language-review.json"] = _json(public["language_reviews"])
    if comparison_edition is not None:
        comparison = _comparison(_frozen(comparison_edition), edition)
        files.update(_document_files("publication/edition-comparison", _comparison_model(comparison)))
        files["publication/edition-comparison.json"] = _json(comparison)
    index = _model("English and French publication", edition["label"], "en", edition["frozen_at"])
    _para(index, "Open the language edition below. The Word documents contain editable numbered headings, paragraphs and tables. "
          "The PDF editions are ready to read or print. Source references and retained original case cells accompany the same edition.")
    # Only fixed filenames become local index links.
    links = [("English Word document", "standard-en.docx"), ("Document Word français", "standard-fr.docx"),
             ("English PDF", "standard-en.pdf"), ("PDF français", "standard-fr.pdf"),
             ("English web edition", "standard-en.html"), ("Édition Web française", "standard-fr.html"),
             ("Original source case cells", "case-records-en.html"), ("Cellules sources originales", "case-records-fr.html")]
    if comparison_edition is not None:
        links.extend([("Edition comparison", "edition-comparison.html"), ("Editable edition comparison", "edition-comparison.docx")])
    index_html = _html(index).decode().replace("</body>", "<ul>" + "".join(
        '<li><a href="' + path + '">' + escape(label) + "</a></li>" for label, path in links) +
        '</ul><p class="small">Edition ' + escape(edition["id"]) + " · SHA256 " + edition["sha256"] + "</p></body>")
    files["publication/index.html"] = index_html.encode()
    files["publication/manifest.json"] = _manifest(files, context)
    return files


def delivery_files(state, analysis):
    """Return private four-phase, approval, financial and receiving workpapers."""
    state, analysis = deepcopy(state), deepcopy(analysis)
    engine.validate_state(state)
    if engine.canonical(analysis) != engine.canonical(engine.analyze(state)):
        raise PublicationError("Delivery analysis does not describe the selected captured state; recapture state and analysis together.")
    model = _management_model(state, analysis)
    files = _document_files("management/delivery-workpapers", model)
    phases = {row["phase"]: row for row in analysis["phases"]}
    for phase in engine.PHASES:
        phase_analysis = dict(phases.get(phase, {}))
        phase_analysis["communication_reviews"] = analysis.get("communication_reviews", [])
        files["management/phase-" + phase + ".html"] = _html(_phase_model(state, phase, phase_analysis))
    selected = ("funding", "tasks", "deliverables", "costs", "phase_reports", "report_responses", "invoices",
                "invoice_responses", "payments", "changes", "comments", "ballots", "votes",
                "ballot_closures", "procedural_reviews", "ola_events", "maintenance", "uptake")
    records = {key: state[key] for key in selected}
    records.update({key: state.get(key, []) for key in ("communications", "delivery_registers")})
    records["current_phase_snapshots"] = {phase: engine.phase_snapshot(state, phase) for phase in engine.PHASES}
    files["management/retained-records.json"] = _json({"project_id": state["id"], "mode": state["mode"],
                                                     "state_sha256": engine.digest(state), **records})
    files["management/delivery-analysis.json"] = _json(analysis)
    files["management/manifest.json"] = _manifest(files, {"schema": SCHEMA, "project_id": state["id"],
                                                        "state_sha256": engine.digest(state), "mode": state["mode"]})
    return files

