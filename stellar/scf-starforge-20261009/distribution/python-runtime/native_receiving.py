#!/usr/bin/env python3
"""Operate one source-pinned Python workbook receiving journey.

The maintained runtime installer, Python wheel engine, AnalysisWorkspace and
WorkbookDelivery remain the implementations. Run this driver only within the
original admitted native operation. All customer state and output stay private.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import date, datetime, time as time_value, timezone
import fcntl
import hashlib
import importlib
import importlib.metadata
import json
import os
from pathlib import Path, PurePosixPath
import re
import selectors
import signal
import stat
import subprocess
import sys
import time
import uuid

SCHEMA = "stellar-forge.python-native-receiving.v1"
PROJECT = "stellar/scf-starforge-20261009/sf88-operator-buyer-cli/python"
MAX_JSON = 16 * 1024 * 1024
HEX = re.compile(r"^[a-f0-9]{64}$")
COMMIT = re.compile(r"^[a-f0-9]{40}$")
_ACTIVE = None


class ReceivingError(RuntimeError):
    def __init__(self, code, detail=None):
        self.code, self.detail = code, detail
        super().__init__(code)


def need(value, code, detail=None):
    if not value:
        raise ReceivingError(code, detail)


def encoded(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2,
                       allow_nan=False) + "\n").encode("utf-8")


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def stamp():
    return datetime.now(timezone.utc).isoformat()


def absolute(value, *, file=False):
    need(isinstance(value, str) and value and "\0" not in value,
         "ABSOLUTE_PATH_REQUIRED")
    p = Path(value)
    need(p.is_absolute(), "ABSOLUTE_PATH_REQUIRED")
    if file:
        info = p.lstat()
        need(stat.S_ISREG(info.st_mode) and not p.is_symlink(), "REGULAR_FILE_REQUIRED")
    return p


def private(path):
    p = Path(path).absolute()
    p.mkdir(mode=0o700, parents=True, exist_ok=True)
    info = p.lstat()
    need(stat.S_ISDIR(info.st_mode) and p.resolve() == p
         and info.st_uid == os.getuid() and not info.st_mode & 0o077,
         "PRIVATE_DIRECTORY_REQUIRED", str(p))
    return p


def sync(path):
    fd = os.open(path, os.O_RDONLY | os.O_DIRECTORY)
    try:
        os.fsync(fd)
    finally:
        os.close(fd)


def write(path, value, *, create=False):
    path = Path(path)
    raw = value if isinstance(value, bytes) else encoded(value)
    temporary = path.parent / (".pending-" + uuid.uuid4().hex)
    fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(fd, "wb") as output:
            output.write(raw)
            output.flush()
            os.fsync(output.fileno())
        if create:
            os.link(temporary, path)
        else:
            os.replace(temporary, path)
        sync(path.parent)
    finally:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass


def read(path, maximum=MAX_JSON):
    p = absolute(str(path), file=True)
    need(p.stat().st_size <= maximum, "JSON_SIZE_LIMIT")
    with p.open("rb") as stream:
        raw = stream.read(maximum + 1)
    need(len(raw) <= maximum, "JSON_SIZE_LIMIT")
    return json.loads(raw)


def environment():
    result = dict(os.environ)
    for key in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV"):
        result.pop(key, None)
    result.update(PYTHONNOUSERSITE="1", PYTHONDONTWRITEBYTECODE="1",
                  PIP_DISABLE_PIP_VERSION_CHECK="1")
    return result


def interrupted(signum, frame):
    raise KeyboardInterrupt


def stop(process):
    if process.poll() is not None:
        return
    # The driver starts and owns this group; no peer's process is selected.
    for sig, seconds in ((signal.SIGINT, 15), (signal.SIGTERM, 10), (signal.SIGKILL, 5)):
        if process.poll() is not None:
            break
        try:
            os.killpg(process.pid, sig)
        except ProcessLookupError:
            break
        try:
            process.wait(timeout=seconds)
        except subprocess.TimeoutExpired:
            pass


@contextmanager
def operation_lease(directory):
    p = directory / "operation.lock"
    fd = os.open(p, os.O_CREAT | os.O_RDWR | os.O_NOFOLLOW, 0o600)
    try:
        info, named = os.fstat(fd), p.lstat()
        need(stat.S_ISREG(info.st_mode) and info.st_nlink == 1
             and not info.st_mode & 0o077 and info.st_uid == os.getuid()
             and (info.st_dev, info.st_ino) == (named.st_dev, named.st_ino),
             "PRIVATE_OPERATION_LOCK_REQUIRED")
        try:
            fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise ReceivingError("ORIGINAL_RECEIVING_OPERATION_RUNNING") from None
        yield fd
    finally:
        os.close(fd)


class Journey:
    def __init__(self, plan, root, fd):
        self.plan, self.root, self.fd = plan, root, fd
        self.plan_sha256 = hashlib.sha256(encoded(plan)).hexdigest()
        self.path = root / "operation.json"
        if self.path.exists():
            self.state = read(self.path)
            need(self.state["plan_sha256"] == self.plan_sha256
                 and self.state["operation_id"] == plan["operation_id"],
                 "RECEIVING_OPERATION_INPUT_CHANGED")
        else:
            self.state = {"schema": SCHEMA, "operation_id": plan["operation_id"],
                          "plan_sha256": self.plan_sha256, "created_at": stamp(),
                          "state": "PREPARED", "steps": {}}
            write(root / "plan.json", plan, create=True)
            write(self.path, self.state, create=True)
        private(root / "steps")

    def save(self):
        self.state["updated_at"] = stamp()
        write(self.path, self.state)

    def step(self, name, intent, apply, *, recovered=None, refresh=False):
        need(re.fullmatch(r"[a-z][a-z0-9-]{0,79}", name), "STEP_NAME_INVALID")
        expected = hashlib.sha256(encoded(intent)).hexdigest()
        row = self.state["steps"].get(name)
        directory = private(self.root / "steps" / name)
        if row:
            need(row["intent_sha256"] == expected, "RECEIVING_STEP_INPUT_CHANGED", name)
            if row["state"] == "COMPLETED" and not refresh:
                need(digest(directory / "result.json") == row["result_sha256"],
                     "RECEIVING_STEP_RESULT_CHANGED", name)
                return read(directory / "result.json")
        row = row or {"intent_sha256": expected, "attempts": 0}
        row.update(state="RUNNING", attempts=row["attempts"] + 1, started_at=stamp())
        self.state["steps"][name] = row
        self.state["state"] = "RUNNING"
        self.save()
        try:
            result = recovered(directory) if recovered else None
            if result is None:
                result = apply(directory, row["attempts"])
            write(directory / "result.json", result)
            row.update(state="COMPLETED", result_sha256=digest(directory / "result.json"),
                       completed_at=stamp())
            self.save()
            return result
        except BaseException as error:
            row.update(state="INTERRUPTED" if isinstance(error, KeyboardInterrupt) else "FAILED",
                       error=getattr(error, "code", type(error).__name__), finished_at=stamp())
            self.state["state"] = row["state"]
            self.save()
            raise

    def command(self, directory, attempt, argv, *, timeout=1800):
        global _ACTIVE
        log = directory / ("command-" + str(attempt))
        stdout, stderr = Path(str(log) + ".stdout"), Path(str(log) + ".stderr")
        with stdout.open("xb") as output, stderr.open("xb") as errors:
            os.fchmod(output.fileno(), 0o600)
            os.fchmod(errors.fileno(), 0o600)
            process = subprocess.Popen([str(a) for a in argv], stdin=subprocess.DEVNULL,
                stdout=output, stderr=errors, env=environment(), start_new_session=True,
                pass_fds=(self.fd,))
            _ACTIVE = process
            try:
                code = process.wait(timeout=timeout)
            except subprocess.TimeoutExpired:
                raise ReceivingError("NATIVE_COMMAND_TIMEOUT", {"stdout": str(stdout)})
            finally:
                stop(process)
                _ACTIVE = None
                output.flush()
                errors.flush()
                os.fsync(output.fileno())
                os.fsync(errors.fileno())
        need(code == 0, "NATIVE_COMMAND_FAILED",
             {"exit_code": code, "stdout": str(stdout), "stderr": str(stderr)})
        need(stdout.stat().st_size <= MAX_JSON, "NATIVE_RESULT_SIZE_LIMIT")
        try:
            return json.loads(stdout.read_bytes())
        except (ValueError, UnicodeError):
            raise ReceivingError("NATIVE_JSON_RESULT_REQUIRED", {"stdout": str(stdout)}) from None

    def engine(self, name, runtime, command, arguments=()):
        intent = {"engine_sha256": self.plan["engine_sha256"],
                  "runtime": str(runtime), "command": command, "arguments": list(arguments)}

        def recover(directory):
            path = directory / "engine-progress.json"
            if not path.exists():
                return None
            retained = read(path)
            if (retained.get("stage") == "complete"
                    and retained.get("state") == "COMPLETED"
                    and retained.get("command") == command
                    and isinstance(retained.get("result"), dict)):
                return retained["result"]
            return None

        def apply(directory, attempt):
            argv = [self.plan["python"], self.plan["engine"], "--root", str(runtime),
                    "--lease-fd", str(self.fd), "--progress-file",
                    str(directory / "engine-progress.json"), command, *arguments]
            return self.command(directory, attempt, argv)
        return self.step(name, intent, apply,
                         recovered=None if command == "inspect" else recover,
                         refresh=(command == "inspect"))


# This calls the maintained source installer directly. It does not unpack source,
# resolve Node dependencies, or infer a source manifest on its own.
INSTALLER = r"""
import { readFile, access, open, rename, unlink } from 'node:fs/promises';
import { randomUUID } from 'node:crypto';
import { join } from 'node:path';
import { pathToFileURL } from 'node:url';
const request = JSON.parse(await readFile(process.argv[1], 'utf8'));
const installer = await import(pathToFileURL(request.installer_module).href);
const exists = async p => { try { await access(p); return true; } catch { return false; } };
const selectionPath = request.selection_file;
let selection = {
  schema:'stellar-forge.python-source-selection.v1',
  originalDirectory:request.installation_directory,
  directory:request.installation_directory,from:null,sequence:0
};
async function retainSelection() {
  const temporary=selectionPath+'.pending-'+randomUUID();
  const handle=await open(temporary,'wx',0o600);
  try { await handle.writeFile(JSON.stringify(selection)); await handle.sync(); }
  finally { await handle.close(); }
  try { await rename(temporary,selectionPath); }
  finally { await unlink(temporary).catch(()=>{}); }
  const parent=await open(request.step_directory,'r');
  try { await parent.sync(); } finally { await parent.close(); }
}
if (await exists(selectionPath)) {
  selection=JSON.parse(await readFile(selectionPath,'utf8'));
  if (selection.schema!=='stellar-forge.python-source-selection.v1'
      || selection.originalDirectory!==request.installation_directory
      || !Number.isSafeInteger(selection.sequence) || selection.sequence<0)
    throw Object.assign(new Error('ORIGINAL_SOURCE_SELECTION_CHANGED'),{code:'ORIGINAL_SOURCE_SELECTION_CHANGED'});
  const expected=selection.sequence===0 ? request.installation_directory
    : join(request.step_directory,'source-recovered-'+selection.sequence);
  const previous=selection.sequence===0 ? null : selection.sequence===1
    ? request.installation_directory : join(request.step_directory,'source-recovered-'+(selection.sequence-1));
  if(selection.directory!==expected || selection.from!==previous)
    throw Object.assign(new Error('ORIGINAL_SOURCE_SELECTION_CHANGED'),{code:'ORIGINAL_SOURCE_SELECTION_CHANGED'});
} else {
  await retainSelection();
}
if (!await exists(join(selection.directory,'runtime-install.json'))) {
  if (await exists(selection.directory)) {
    // The existing installer owns recovery of its own partial archive/intent.
    selection={
      ...selection,from:selection.directory,sequence:selection.sequence+1,
      directory:join(request.step_directory,'source-recovered-'+(selection.sequence+1))
    };
    await retainSelection();
  }
  if(selection.sequence>0) {
    await installer.recoverRuntime({
      from:selection.from,directory:selection.directory,python:request.python,install:false
    });
  } else {
    if(!request.prepare)
      throw Object.assign(new Error('ORIGINAL_SOURCE_INSTALLATION_REQUIRED'),{code:'ORIGINAL_SOURCE_INSTALLATION_REQUIRED'});
    await installer.prepareRuntimeInstall({
      ...request.prepare,directory:selection.directory,python:request.python
    });
  }
}
if (request.install_zero_dependencies) {
  const profile=JSON.parse(await readFile(join(selection.directory,'runtime-profile.json'),'utf8'));
  if(Object.keys(profile.packageJson.dependencies ?? {}).length!==0)
    throw Object.assign(new Error('ORIGINAL_NODE_COHORT_REQUIRED'),{code:'ORIGINAL_NODE_COHORT_REQUIRED'});
  await installer.installRuntime({directory:selection.directory,offline:true,npm:request.npm ?? 'npm'});
}
const observed=await installer.inspectRuntimeInstall({
  directory:selection.directory,load:true,source:true,python:request.python
});
const plan=JSON.parse(await readFile(join(selection.directory,'runtime-install.json'),'utf8'));
console.log(JSON.stringify({
  status:observed.status,ready:observed.ready,sourceVerified:observed.sourceVerified,
  directory:observed.directory,installationId:observed.installationId,profileId:observed.profileId,
  sourceRoot:observed.sourceRoot,sourceCommit:observed.sourceCommit,
  sourceManifestSha256:plan.source.manifestSha256,archiveSha256:plan.archiveSha256,
  dependencyRoot:observed.dependencyRoot,packageVersion:observed.packageVersion,
  recovery:selection.sequence>0 ? selection : null
}));
"""


def select_source(journey, name, specification):
    need(isinstance(specification, dict), "ORIGINAL_SOURCE_SELECTION_REQUIRED")
    module = absolute(specification["installer_module"], file=True)
    need(HEX.fullmatch(specification["installer_sha256"])
         and digest(module) == specification["installer_sha256"], "INSTALLER_SOURCE_CHANGED")
    need(COMMIT.fullmatch(specification["commit"])
         and HEX.fullmatch(specification["manifest_sha256"]), "SOURCE_IDENTITY_REQUIRED")
    request = dict(specification)
    request["python"] = journey.plan["python"]
    absolute(request["installation_directory"])

    def apply(directory, attempt):
        path = directory / "installer-request.json"
        selected_request = {**request, "step_directory": str(directory),
                            "selection_file": str(directory / "selected-source.json")}
        if not path.exists():
            write(path, selected_request, create=True)
        else:
            need(read(path) == selected_request, "INSTALLER_REQUEST_CHANGED")
        return journey.command(directory, attempt, [
            journey.plan["node"], "--input-type=module", "-e", INSTALLER, str(path)])

    result = journey.step(name, request, apply, refresh=True)
    need(result.get("sourceVerified") is True
         and result["sourceCommit"] == specification["commit"]
         and result["sourceManifestSha256"] == specification["manifest_sha256"],
         "ORIGINAL_SOURCE_SELECTION_CHANGED")
    if specification.get("root") and not result.get("recovery"):
        need(result["sourceRoot"] == specification["root"], "ORIGINAL_SOURCE_ROOT_CHANGED")
    if result.get("recovery"):
        need(result["recovery"]["originalDirectory"] == specification["installation_directory"],
             "ORIGINAL_RECOVERY_SOURCE_CHANGED")
    need((Path(result["sourceRoot"]) / PROJECT / "pyproject.toml").is_file(),
         "ORIGINAL_PYTHON_PROJECT_ABSENT")
    return result


def source_matches(source, selection):
    release = source.get("releaseId", source.get("sourceCommit"))
    manifest = source.get("releaseDigest", source.get("sourceManifestSha256"))
    need(release == selection["sourceCommit"]
         and manifest == selection["sourceManifestSha256"], "PYTHON_SOURCE_IDENTITY_CHANGED")


def candidate_id(result):
    value = result.get("candidate", result.get("id"))
    need(isinstance(value, str) and HEX.fullmatch(value), "CANDIDATE_ID_REQUIRED")
    return value


def installation_id(result):
    value = result.get("installation", result.get("id"))
    need(isinstance(value, str) and HEX.fullmatch(value), "INSTALLATION_ID_REQUIRED")
    return value


def cohort(journey, prefix, root, selection):
    project = Path(selection["sourceRoot"]) / PROJECT
    extras = sorted(journey.plan.get("extras", ["workbooks"]))
    result = journey.engine(prefix + "-resolve", root, "resolve", [
        "--python", journey.plan["python"], "--source", str(project), "--extras", ",".join(extras),
        "--release-id", selection["sourceCommit"],
        "--release-digest", selection["sourceManifestSha256"]])
    source_matches(result["source"], selection)
    need(result["extras"] == extras, "EXACT_DECLARED_EXTRAS_REQUIRED")
    identifier = candidate_id(result)
    checked = journey.engine(prefix + "-candidate", root, "inspect", ["--candidate", identifier])
    need(candidate_id(checked) == identifier and checked["source"] == result["source"],
         "RETAINED_CANDIDATE_CHANGED")
    return result


def install_cohort(journey, prefix, root, candidate, selection):
    identifier = candidate_id(candidate)
    installed = journey.engine(prefix + "-install", root, "install", [
        "--candidate", identifier, "--python", journey.plan["python"]])
    need(installed["candidate"] == identifier and installed["source"] == candidate["source"],
         "INSTALLATION_SOURCE_CHANGED")
    source_matches(installed["source"], selection)
    need(installed["capabilities"]["analysis"] is True
         and installed["capabilities"]["workbooks"] is True
         and installed["extras"] == candidate["extras"]
         and all(installed["capabilities"].get(name) is True for name in candidate["extras"]),
         "SELECTED_PYTHON_RUNTIME_NOT_READY")
    checked = journey.engine(prefix + "-installation", root, "inspect", [
        "--installation", installation_id(installed)])
    need(installation_id(checked) == installation_id(installed)
         and checked["source"] == installed["source"], "RETAINED_INSTALLATION_CHANGED")
    return installed


def cli(journey, name, installed, arguments):
    return journey.step(name, {"installation": installation_id(installed), "arguments": arguments},
        lambda directory, attempt: journey.command(directory, attempt,
            [installed["python"], *arguments], timeout=300))



def file_identity(info):
    return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns)


def customer_tree(directory):
    """Inventory the original owner's complete closed directory without opening SQL."""
    root = absolute(str(directory))
    info = root.lstat()
    need(stat.S_ISDIR(info.st_mode) and root.resolve() == root,
         "CUSTOMER_DIRECTORY_REQUIRED")
    directories, files, pending = [], [], [root]
    while pending:
        current = pending.pop()
        for path in sorted(current.iterdir()):
            entry = path.lstat()
            relative = path.relative_to(root).as_posix()
            if stat.S_ISDIR(entry.st_mode):
                directories.append(relative)
                pending.append(path)
                continue
            need(stat.S_ISREG(entry.st_mode), "CUSTOMER_COPY_REGULAR_FILES_REQUIRED", relative)
            fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
            with os.fdopen(fd, "rb") as stream:
                opened = os.fstat(stream.fileno())
                need(file_identity(opened) == file_identity(entry), "CUSTOMER_SOURCE_CHANGED", relative)
                sha = hashlib.file_digest(stream, "sha256").hexdigest()
                need(file_identity(os.fstat(stream.fileno())) == file_identity(opened)
                     and file_identity(path.lstat()) == file_identity(opened),
                     "CUSTOMER_SOURCE_CHANGED", relative)
            files.append({"path": relative, "bytes": entry.st_size, "sha256": sha})
    return {"directories": sorted(directories), "files": sorted(files, key=lambda x: x["path"])}


def copy_closed_customer(request):
    """Receive the owner's full directory under its existing analysis.lock."""
    from stellar_forge_buyer.analysis_workbooks import _publish_received_directory

    os.fstat(request["lease_fd"])
    source = absolute(request["customer"]["database"])
    target = absolute(request["target"])
    control = private(request["directory"])
    private(target.parent)
    snapshot_path = control / "copy-snapshot.json"
    snapshot = read(snapshot_path) if snapshot_path.exists() else None
    expected = {"schema": "stellar-forge.python-customer-copy.v1",
                "source": str(source), "database": str(target),
                "workspace_id": request["customer"]["workspace_id"]}
    if snapshot is not None:
        need(all(snapshot.get(k) == v for k, v in expected.items()),
             "CUSTOMER_COPY_SELECTION_CHANGED")
        need(hashlib.sha256(encoded(snapshot["inventory"])).hexdigest()
             == snapshot["inventory_sha256"], "CUSTOMER_COPY_SNAPSHOT_CHANGED")

    def completed():
        inventory = snapshot["inventory"]
        need(customer_tree(target) == inventory, "RECEIVED_CUSTOMER_COPY_CHANGED")
        result = {**expected, "state": "COMPLETED",
                  "inventory_sha256": snapshot["inventory_sha256"],
                  "snapshot_sha256": digest(snapshot_path),
                  "files": len(inventory["files"]),
                  "bytes": sum(x["bytes"] for x in inventory["files"]),
                  "source_preserved_during_copy": True}
        write(control / "copy-result.json", result)
        return result

    if target.exists():
        need(snapshot is not None, "CUSTOMER_COPY_DESTINATION_EXISTS")
        return completed()

    # This is the original AnalysisWorkspace lock, opened without creating or
    # changing any source file. The full database/artifact directory is the
    # owner's documented copy unit; retained producer paths stay unchanged.
    lock_path = source / "analysis.lock"
    lock_fd = os.open(lock_path, os.O_RDONLY | os.O_NOFOLLOW)
    try:
        info, named = os.fstat(lock_fd), lock_path.lstat()
        need(stat.S_ISREG(info.st_mode) and info.st_nlink == 1
             and not info.st_mode & 0o077 and info.st_uid == os.getuid()
             and (info.st_dev, info.st_ino) == (named.st_dev, named.st_ino),
             "ORIGINAL_CUSTOMER_LOCK_REQUIRED")
        try:
            fcntl.flock(lock_fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise ReceivingError("ORIGINAL_CUSTOMER_STILL_OPEN") from None
        need(file_identity(os.fstat(lock_fd)) == file_identity(lock_path.lstat()),
             "ORIGINAL_CUSTOMER_LOCK_CHANGED")
        inventory = customer_tree(source)
        need(any(x["path"] == "datasets.sqlite3" for x in inventory["files"]),
             "ORIGINAL_CUSTOMER_DATABASE_REQUIRED")
        if snapshot is None:
            snapshot = {**expected, "captured_at": stamp(), "inventory": inventory,
                        "inventory_sha256": hashlib.sha256(encoded(inventory)).hexdigest()}
            write(snapshot_path, snapshot, create=True)
        else:
            need(snapshot["inventory"] == inventory, "CUSTOMER_COPY_SOURCE_CHANGED")
        staging = private(control / "customer-stage")
        for relative in inventory["directories"]:
            private(staging / relative)
        temporary = control / "pending-file"
        for row in inventory["files"]:
            original, destination = source / row["path"], staging / row["path"]
            if destination.exists():
                need(stat.S_ISREG(destination.lstat().st_mode)
                     and not destination.is_symlink()
                     and destination.stat().st_size == row["bytes"]
                     and digest(destination) == row["sha256"], "CUSTOMER_COPY_STAGE_CHANGED", row["path"])
                continue
            if temporary.exists():
                need(stat.S_ISREG(temporary.lstat().st_mode) and not temporary.is_symlink(),
                     "CUSTOMER_COPY_STAGE_CHANGED")
                temporary.unlink()
            input_fd = os.open(original, os.O_RDONLY | os.O_NOFOLLOW)
            output_fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
            sha, count = hashlib.sha256(), 0
            with os.fdopen(input_fd, "rb") as input_stream, os.fdopen(output_fd, "wb") as output:
                before = os.fstat(input_stream.fileno())
                for chunk in iter(lambda: input_stream.read(1024 * 1024), b""):
                    output.write(chunk)
                    sha.update(chunk)
                    count += len(chunk)
                output.flush()
                os.fsync(output.fileno())
                need(file_identity(before) == file_identity(os.fstat(input_stream.fileno()))
                     and file_identity(before) == file_identity(original.lstat())
                     and count == row["bytes"] and sha.hexdigest() == row["sha256"],
                     "CUSTOMER_SOURCE_CHANGED", row["path"])
            try:
                os.link(temporary, destination)
                sync(destination.parent)
            finally:
                try:
                    temporary.unlink(missing_ok=True)
                except OSError:
                    pass
        need(customer_tree(source) == inventory and customer_tree(staging) == inventory,
             "CUSTOMER_COPY_SOURCE_CHANGED")
        for relative in sorted(inventory["directories"], key=lambda x: x.count("/"), reverse=True):
            sync(staging / relative)
        sync(staging)
        _publish_received_directory(staging, target)
        return completed()
    finally:
        os.close(lock_fd)


def customer_transfer_inventory(customer, archive):
    transfer = customer['directory_transfer']
    destination = absolute(str(Path(customer['database']).parent))
    def relative(value):
        need(isinstance(value, str) and value and "\\" not in value and "\0" not in value
             and not value.startswith("/") and all(p not in ("", ".", "..") for p in value.split("/")),
             "CUSTOMER_TRANSFER_PATH_INVALID")
        return value

    members = archive.infolist()
    names = [x.filename for x in members]
    need(len(names) <= 100000 and len(names) == len(set(names))
         and "transfer-manifest.json" in names, "CUSTOMER_TRANSFER_INVENTORY_INVALID")
    manifest_info = archive.getinfo("transfer-manifest.json")
    need(manifest_info.file_size <= MAX_JSON, "CUSTOMER_TRANSFER_MANIFEST_TOO_LARGE")
    manifest_bytes = archive.read(manifest_info)
    need(hashlib.sha256(manifest_bytes).hexdigest() == transfer["manifest_sha256"],
         "CUSTOMER_TRANSFER_MANIFEST_CHANGED")
    manifest = json.loads(manifest_bytes)
    need(manifest.get("schema") == "stellar-forge.customer-directory-transfer.v1"
         and manifest.get("database_root") == "customer-database"
         and manifest.get("workspace_id") == customer["workspace_id"]
         and manifest.get("source_preserved_during_capture") is True,
         "CUSTOMER_TRANSFER_IDENTITY_CHANGED")
    workbook = manifest.get("workbook", {})
    need(workbook.get("name") == customer["workbook_name"]
         and workbook.get("revision") == customer["revision"]
         and workbook.get("build_id") == customer["delivery_id"]
         and workbook.get("row_count") == customer["row_count"],
         "CUSTOMER_TRANSFER_WORKBOOK_CHANGED")
    need(isinstance(manifest.get("files"), list)
         and isinstance(manifest.get("directories"), list),
         "CUSTOMER_TRANSFER_INVENTORY_INVALID")
    files = {}
    for row in manifest["files"]:
        name = relative(row.get("path"))
        need(name not in files and name != "transfer-manifest.json"
             and type(row.get("bytes")) is int and row["bytes"] >= 0
             and isinstance(row.get("sha256"), str) and HEX.fullmatch(row["sha256"]),
             "CUSTOMER_TRANSFER_INVENTORY_INVALID")
        files[name] = {"path": name, "bytes": row["bytes"], "sha256": row["sha256"]}
    directories = sorted(relative(x) for x in manifest["directories"])
    need(len(directories) == len(set(directories))
         and not set(directories) & set(files)
         and sum(x["bytes"] for x in files.values()) <= 8 * 1024**3,
         "CUSTOMER_TRANSFER_INVENTORY_INVALID")
    files["transfer-manifest.json"] = {
        "path": "transfer-manifest.json", "bytes": len(manifest_bytes),
        "sha256": transfer["manifest_sha256"]}
    expected_names = set(files) | {x + "/" for x in directories}
    need(set(names) == expected_names, "CUSTOMER_TRANSFER_MEMBERS_CHANGED")
    for member in members:
        mode = (member.external_attr >> 16) & 0o170000
        need(not member.flag_bits & 1 and not stat.S_ISLNK(mode),
             "CUSTOMER_TRANSFER_MEMBER_INVALID")
        if member.is_dir():
            need(mode in (0, stat.S_IFDIR) and member.file_size == 0,
                 "CUSTOMER_TRANSFER_MEMBER_INVALID")
        else:
            need(mode in (0, stat.S_IFREG)
                 and member.file_size == files[member.filename]["bytes"],
                 "CUSTOMER_TRANSFER_MEMBER_INVALID")
    original_inventory = {
        "directories": [x[len("customer-database/"):] for x in directories
                        if x.startswith("customer-database/")],
        "files": [{**row, "path": name[len("customer-database/"):]}
                  for name, row in sorted(files.items())
                  if name.startswith("customer-database/")]}
    need(hashlib.sha256(encoded(original_inventory)).hexdigest()
         == manifest.get("database_inventory_sha256")
         and len(original_inventory["files"]) == manifest.get("database_files")
         and sum(x["bytes"] for x in original_inventory["files"])
         == manifest.get("database_bytes"),
         "CUSTOMER_TRANSFER_DATABASE_INVENTORY_CHANGED")
    need("customer-database/analysis.lock" in files
         and "customer-database/datasets.sqlite3" in files
         and str(destination / "customer-database") == customer["database"]
         and str(destination / "edition-operation/handoff.zip") == customer["archive"]
         and files.get("edition-operation/handoff.zip", {}).get("sha256")
         == customer["archive_sha256"]
         and files.get("edition-operation/manifest.json", {}).get("sha256")
         == customer["manifest_sha256"], "CUSTOMER_TRANSFER_PATH_OR_HANDOFF_CHANGED")
    expected_tree = {"directories": directories,
                     "files": [row for _, row in sorted(files.items())]}
    return manifest, files, directories, expected_tree


def receive_customer_archive(request):
    """Hydrate the original locked directory transfer with the installed publisher."""
    from stellar_forge_buyer.analysis_workbooks import _publish_received_directory
    import zipfile

    os.fstat(request["lease_fd"])
    customer = request["customer"]
    transfer = customer["directory_transfer"]
    archive_path = absolute(transfer["archive"], file=True)
    destination = absolute(str(Path(customer["database"]).parent))
    control = private(request["directory"])
    private(destination.parent)
    need(destination != archive_path.parent and not archive_path.is_relative_to(destination),
         "CUSTOMER_TRANSFER_DESTINATION_OVERLAP")
    expected = {"schema": "stellar-forge.python-customer-receive.v1",
                "operation_id": request["operation_id"], "destination": str(destination),
                "archive_sha256": transfer["sha256"], "manifest_sha256": transfer["manifest_sha256"]}
    intent_path = control / "receive-intent.json"
    if intent_path.exists():
        need(read(intent_path) == expected, "CUSTOMER_TRANSFER_INTENT_CHANGED")
    else:
        write(intent_path, expected, create=True)
    stage = destination.parent / ("." + destination.name + "-receiving-" +
                                  hashlib.sha256(encoded(expected)).hexdigest()[:20])

    fd = os.open(archive_path, os.O_RDONLY | os.O_NOFOLLOW)
    with os.fdopen(fd, "rb") as archive_stream:
        opened = os.fstat(archive_stream.fileno())
        need(stat.S_ISREG(opened.st_mode)
             and file_identity(opened) == file_identity(archive_path.lstat()),
             "CUSTOMER_TRANSFER_ARCHIVE_CHANGED")
        need(hashlib.file_digest(archive_stream, "sha256").hexdigest() == transfer["sha256"],
             "CUSTOMER_TRANSFER_ARCHIVE_CHANGED")
        archive_stream.seek(0)
        with zipfile.ZipFile(archive_stream) as archive:
            manifest, files, directories, expected_tree = customer_transfer_inventory(customer, archive)
            if destination.exists():
                need(customer_tree(destination) == expected_tree, "CUSTOMER_TRANSFER_DESTINATION_EXISTS")
            else:
                private(stage)
                for directory in sorted(directories, key=lambda x: (x.count("/"), x)):
                    private(stage / directory)
                for name, row in sorted(files.items()):
                    output = stage / name
                    private(output.parent)
                    temporary = output.parent / (".transfer-part-" +
                                                  hashlib.sha256(name.encode()).hexdigest())
                    need(temporary.relative_to(stage).as_posix() not in
                         (set(files) | {x + "/" for x in directories}),
                         "CUSTOMER_TRANSFER_TEMPORARY_CONFLICT")
                    if output.exists():
                        need(output.is_file() and not output.is_symlink()
                             and output.stat().st_size == row["bytes"] and digest(output) == row["sha256"],
                             "CUSTOMER_TRANSFER_STAGED_FILE_CHANGED", name)
                        if temporary.exists():
                            info = temporary.lstat()
                            need(stat.S_ISREG(info.st_mode) and info.st_uid == os.getuid(),
                                 "CUSTOMER_TRANSFER_TEMPORARY_CHANGED")
                            temporary.unlink(); sync(output.parent)
                        continue
                    if temporary.exists():
                        info = temporary.lstat()
                        need(stat.S_ISREG(info.st_mode) and info.st_uid == os.getuid()
                             and info.st_nlink == 1, "CUSTOMER_TRANSFER_TEMPORARY_CHANGED")
                    count, sha = 0, hashlib.sha256()
                    target_fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_TRUNC |
                                        os.O_NOFOLLOW, 0o600)
                    with os.fdopen(target_fd, "wb") as sink, archive.open(name) as source:
                        while data := source.read(1024 * 1024):
                            count += len(data)
                            need(count <= row["bytes"], "CUSTOMER_TRANSFER_FILE_SIZE_CHANGED", name)
                            sink.write(data); sha.update(data)
                        sink.flush(); os.fsync(sink.fileno())
                    need(count == row["bytes"] and sha.hexdigest() == row["sha256"],
                         "CUSTOMER_TRANSFER_FILE_CHANGED", name)
                    os.link(temporary, output); sync(output.parent); temporary.unlink()
                need(customer_tree(stage) == expected_tree, "CUSTOMER_TRANSFER_STAGED_TREE_CHANGED")
                for directory in sorted(directories, key=lambda x: x.count("/"), reverse=True):
                    sync(stage / directory)
                sync(stage)
                _publish_received_directory(stage, destination)
            need(file_identity(os.fstat(archive_stream.fileno())) == file_identity(opened)
                 and file_identity(archive_path.lstat()) == file_identity(opened),
                 "CUSTOMER_TRANSFER_ARCHIVE_CHANGED")
    result = {**expected, "state": "COMPLETED", "database": customer["database"],
              "database_files": manifest["database_files"],
              "database_bytes": manifest["database_bytes"],
              "database_inventory_sha256": manifest["database_inventory_sha256"],
              "source_request": manifest.get("source_request")}
    write(control / "receive-result.json", result)
    return result


def hydrate_customer_transfer(journey, installed):
    customer = journey.plan["customer"]
    if not customer.get("directory_transfer"):
        return None
    intent = {"installation": installation_id(installed), "customer": customer}

    def apply(directory, attempt):
        request = {**intent, "operation_id": journey.plan["operation_id"],
                   "directory": str(directory), "lease_fd": journey.fd}
        request_path = directory / ("customer-receive-request-" + str(attempt) + ".json")
        write(request_path, request, create=True)
        return journey.command(directory, attempt,
            [installed["python"], str(Path(__file__).absolute()), "_customer-receive",
             "--request", str(request_path)], timeout=600)
    return journey.step("customer-directory-receive", intent, apply)


def receive_customer_copy(journey, installed):
    target = private(journey.root / "customer") / "database"
    intent = {"customer": journey.plan["customer"], "target": str(target),
              "installation": installation_id(installed)}

    def apply(directory, attempt):
        request = {**intent, "directory": str(directory), "lease_fd": journey.fd}
        path = directory / ("copy-request-" + str(attempt) + ".json")
        write(path, request, create=True)
        return journey.command(directory, attempt, [
            installed["python"], str(Path(__file__).absolute()), "_customer-copy",
            "--request", str(path)], timeout=600)

    return journey.step("customer-copy", intent, apply)


def use_customer(journey, name, installed, selection, customer, *, build=False, expected_build_id=None):
    directory = private(journey.root / "customer" / name)
    request = {
        "schema": SCHEMA, "operation_id": journey.plan["operation_id"],
        "customer": customer, "installation": installed,
        "source": selection, "node": journey.plan["node"], "directory": str(directory),
        "build": build, "build_id": hashlib.sha256(
            (journey.plan["operation_id"] + "/reviewed-workbook-build").encode()).hexdigest()[:32],
        "expected_build_id": expected_build_id, "lease_fd": journey.fd
    }
    intent = {key: value for key, value in request.items() if key != "lease_fd"}

    def apply(step_directory, attempt):
        path = step_directory / ("customer-request-" + str(attempt) + ".json")
        write(path, request, create=True)
        return journey.command(step_directory, attempt, [
            installed["python"], str(Path(__file__).absolute()), "_customer", "--request", str(path)],
            timeout=600)

    return journey.step(name, intent, apply)


def allocated_tree(directory):
    root = absolute(directory)
    need(root.is_dir() and not root.is_symlink(), "RECEIVING_ALLOCATION_ROOT_REQUIRED")
    pending, seen = [root], set()
    total, logical, regular, links, special, entries = 0, 0, 0, 0, 0, 0
    while pending:
        path = pending.pop()
        info = path.lstat()
        identity = (info.st_dev, info.st_ino)
        if identity in seen:
            continue
        seen.add(identity)
        entries += 1
        need(entries <= 200000, "RECEIVING_ALLOCATION_INVENTORY_LIMIT")
        total += info.st_blocks * 512
        if stat.S_ISDIR(info.st_mode):
            with os.scandir(path) as rows:
                pending.extend(Path(row.path) for row in rows)
        elif stat.S_ISREG(info.st_mode):
            logical += info.st_size
            regular += 1
        elif stat.S_ISLNK(info.st_mode):
            links += 1
        else:
            special += 1  # Native controller sockets/FIFOs are counted without opening them.
    return {"allocated_bytes": total, "logical_bytes": logical,
            "inodes": entries, "files": regular, "symlinks": links, "special_files": special}


def receiving_capacity(journey, name, candidate, selection, *, installed=None):
    capacity = journey.plan.get("capacity")
    if capacity is None:
        return None
    intent = {"request": capacity["request"], "directory": capacity["directory"],
              "candidate": candidate_id(candidate), "source": selection["sourceCommit"],
              "installed": installation_id(installed) if installed else None}

    def apply(directory, attempt):
        import zipfile
        request_path = absolute(capacity["request"], file=True)
        request = read(request_path)
        native_intent = request["intent"]
        need(native_intent["operation_id"] == journey.plan["operation_id"],
             "ORIGINAL_CAPACITY_OPERATION_CHANGED")
        workspace = absolute(capacity["directory"])
        need(journey.root.is_relative_to(workspace), "RECEIVING_ALLOCATION_SCOPE_CHANGED")
        requirements = [row for row in native_intent["disk"]["paths"]
                        if Path(row["path"]).resolve() == workspace.resolve()]
        need(len(requirements) == 1, "ORIGINAL_DISK_DECLARATION_REQUIRED")
        requirement = requirements[0]
        memory_bytes = native_intent["memory"]["peak_bytes"]
        process_bytes = capacity["process_peak_bytes"]
        need(type(process_bytes) is int and 0 < process_bytes <= memory_bytes,
             "ORIGINAL_PROCESS_DECLARATION_REQUIRED")
        current = allocated_tree(workspace)
        source = allocated_tree(Path(selection["sourceRoot"]))
        wheel_blocks, python_blocks, wheel_files, python_files = 0, 0, 0, 0
        parents = set()
        block = os.statvfs(workspace).f_frsize
        for package in candidate["packages"]:
            wheel = Path(candidate["directory"]) / "wheels" / package["filename"]
            with zipfile.ZipFile(wheel) as archive:
                for item in archive.infolist():
                    if item.is_dir():
                        continue
                    wheel_files += 1
                    need(wheel_files <= 100000, "PYTHON_COHORT_MEMBER_LIMIT")
                    blocks = ((item.file_size + block - 1) // block) * block
                    wheel_blocks += blocks
                    if item.filename.endswith(".py"):
                        python_blocks += blocks
                        python_files += 1
                    member_path = PurePosixPath(item.filename)
                    need(not member_path.is_absolute() and ".." not in member_path.parts,
                         "PYTHON_COHORT_MEMBER_PATH_INVALID")
                    path = member_path.parent
                    while str(path) != ".":
                        parents.add(str(path))
                        path = path.parent
        # Original wheel entries are measured. The extra bound covers compiled
        # Python files, venv/pip bootstrap, scripts and temporary installation data.
        install_bound = wheel_blocks + 2 * python_blocks + len(parents) * block + 64 * 1024 ** 2
        install_inodes = wheel_files + python_files + len(parents) + 1024
        source_archive = Path(selection["directory"]) / "source.tgz"
        need(source_archive.is_file(), "ORIGINAL_SOURCE_ARCHIVE_REQUIRED_FOR_OFFLINE")
        source_bound = source["allocated_bytes"] + source_archive.stat().st_size + 8 * 1024 ** 2
        source_inodes = source["inodes"] + 512
        observed_installation = allocated_tree(installed["directory"]) if installed else None
        if observed_installation:
            remaining_install_bytes = observed_installation["allocated_bytes"] + 64 * 1024 ** 2
            remaining_install_inodes = observed_installation["inodes"] + 1024
        else:
            remaining_install_bytes = 2 * install_bound
            remaining_install_inodes = 2 * install_inodes
        # The primary cohort already exists. Still to come: the transfer ZIP,
        # its received wheel copy, the offline source runtime and workbook outputs.
        remaining = (remaining_install_bytes + source_bound
                     + 2 * candidate["wheelBytes"] + 16 * 1024 ** 2)
        remaining_inodes = (remaining_install_inodes + source_inodes
                            + 2 * len(candidate["packages"]) + 512)
        projected = current["allocated_bytes"] + remaining
        projected_inodes = current["inodes"] + remaining_inodes
        resources = os.statvfs(workspace)
        reserve_bytes = native_intent["disk"]["reserve_bytes"]
        reserve_inodes = native_intent["disk"]["reserve_inodes"]
        observation = {"schema": "stellar-forge.python-receiving-capacity.v1",
            "observed_at": stamp(), "phase": name, "request": str(request_path),
            "request_sha256": digest(request_path), "workspace": str(workspace),
            "current": current, "source": source,
            "cohort": {"wheel_bytes": candidate["wheelBytes"],
                       "unpacked_bytes": candidate["unpackedBytes"],
                       "allocated_wheel_entry_bytes": wheel_blocks,
                       "files": wheel_files, "python_files": python_files},
            "installation_allocation_bound": install_bound,
            "observed_primary_installation": observed_installation,
            "remaining_allocation_bound": remaining,
            "projected_allocated_bytes": projected, "projected_inodes": projected_inodes,
            "declared": {"disk_bytes": requirement["bytes"], "disk_inodes": requirement["inodes"],
                         "memory_bytes": memory_bytes, "process_peak_bytes": process_bytes,
                         "memory_backed": capacity["memory_backed"]},
            "filesystem_available_bytes": resources.f_bavail * resources.f_frsize,
            "filesystem_available_inodes": resources.f_favail}
        write(directory / "capacity-observation.json", observation)
        fits = (projected <= requirement["bytes"] and projected_inodes <= requirement["inodes"]
                and remaining + reserve_bytes <= resources.f_bavail * resources.f_frsize
                and remaining_inodes + reserve_inodes <= resources.f_favail)
        if capacity["memory_backed"]:
            fits = fits and projected + process_bytes <= memory_bytes
        need(fits, "RECEIVING_CAPACITY_ADJUSTMENT_REQUIRED", observation)
        return observation

    return journey.step(name, intent, apply)


def use_managed_application(journey, name, installed, selection, customer):
    specification = journey.plan.get("applications")
    if specification is None:
        return None
    directory = private(journey.root / "customer" / name)
    helper = absolute(specification["helper"], file=True)
    need(digest(helper) == specification["helper_sha256"],
         "MANAGED_APPLICATION_HELPER_CHANGED")
    request = {"schema": SCHEMA, "operation_id": journey.plan["operation_id"],
               "directory": str(directory), "source": selection,
               "installation": installed, "customer": customer, "node": journey.plan["node"],
               "lease_fd": journey.fd,
               "applications": {key: specification[key]
                                for key in ("sourceRoot", "sourceCommit", "files")}}
    intent = {key: value for key, value in request.items() if key != "lease_fd"}
    intent["helper_sha256"] = specification["helper_sha256"]

    def apply(step_directory, attempt):
        path = step_directory / ("applications-request-" + str(attempt) + ".json")
        write(path, request, create=True)
        result = journey.command(step_directory, attempt,
                                 [journey.plan["node"], str(helper), "--request", str(path)],
                                 timeout=600)
        need(result.get("status") == "COMPLETE"
             and result["workspaceId"] == customer["workspace_id"]
             and result["installation"]["installationId"] == installation_id(installed)
             and result["installation"]["python"] == installed["python"]
             and result["customerSource"]["sourceCommit"] == selection["sourceCommit"]
             and result["customerSource"]["sourceManifestSha256"] == selection["sourceManifestSha256"]
             and result["controller"]["sourceCommit"] == specification["sourceCommit"]
             and len(result["runs"]) == 2
             and all(row["operatorReopened"] and row["stopped"]["state"] == "STOPPED"
                     and row["stopped"]["exitCode"] == 0 for row in result["runs"]),
             "MANAGED_APPLICATION_RESULT_CHANGED")
        return result

    return journey.step(name, intent, apply)


def literal_cell(cell):
    value = cell.value
    if isinstance(value, (datetime, date, time_value)):
        value = {"iso8601": value.isoformat()}
    elif isinstance(value, bytes):
        value = {"hex": value.hex()}
    need(value is None or isinstance(value, (str, int, float, bool, dict)), "WORKBOOK_CELL_TYPE")
    need(cell.data_type != "f", "SOURCE_TEXT_BECAME_EXCEL_FORMULA")
    return {"type": cell.data_type, "value": value, "number_format": cell.number_format}


def workbook_projection(directory, customer, *, original=True):
    from stellar_forge_buyer.analysis_workbooks import WorkbookDelivery
    from openpyxl import load_workbook

    expected_manifest = customer["manifest_sha256"] if original else None
    delivery = WorkbookDelivery(directory, expected_manifest_sha256=expected_manifest)
    manifest = delivery.inspect()
    need(manifest["workspace_id"] == customer["workspace_id"]
         and manifest["name"] == customer["workbook_name"]
         and manifest["revision"] == customer["revision"]
         and manifest["row_count"] == customer["row_count"], "CUSTOMER_WORKBOOK_CHANGED")
    if original:
        need(manifest["id"] == customer["delivery_id"], "ORIGINAL_DELIVERY_CHANGED")
        need(len(manifest["files"]) == customer["payload_files"], "ORIGINAL_PAYLOAD_COUNT_CHANGED")
        need(manifest["files"]["workbook.xlsx"]["sha256"] == customer["xlsx_sha256"],
             "ORIGINAL_XLSX_CHANGED")
    need("history/definition-1.json" in manifest["files"], "ORIGINAL_WORKBOOK_HISTORY_ABSENT")
    book = load_workbook(Path(directory) / "workbook.xlsx", read_only=True, data_only=False)
    figures, found = [], set()
    try:
        for logical in manifest["sheets"]:
            rows, headers = [], [column["header"] for column in logical["columns"]]
            for name in logical["worksheets"]:
                sheet = book[name]
                iterator = sheet.iter_rows()
                first = next(iterator, ())
                need([cell.value for cell in first] == headers, "WORKBOOK_HEADINGS_CHANGED")
                for row in iterator:
                    cells = [literal_cell(cell) for cell in row]
                    rows.append(cells)
                    for cell in row:
                        if isinstance(cell.value, str):
                            for expected in customer["exact_text_cells"]:
                                if expected["sheet"] == logical["name"] and cell.value == expected["value"]:
                                    found.add((expected["sheet"], expected["value"]))
            need(len(rows) == logical["row_count"], "WORKBOOK_ROW_COUNT_CHANGED")
            figures.append({"name": logical["name"], "headers": headers, "rows": rows,
                            "source": logical["source"], "columns": logical["columns"]})
        need(sum(len(x["rows"]) for x in figures) == customer["row_count"], "COMPLETE_CUSTOMER_ROWS_REQUIRED")
        need(all((x["sheet"], x["value"]) in found for x in customer["exact_text_cells"]),
             "EXACT_CUSTOMER_TEXT_CHANGED")
    finally:
        book.close()
    return {
        "id": manifest["id"], "revision": manifest["revision"],
        "workspace_id": manifest["workspace_id"], "row_count": manifest["row_count"],
        "payload_files": len(manifest["files"]), "archive_members": len(manifest["files"]) + 1,
        "manifest_sha256": manifest["manifest_sha256"],
        "definition_sha256": manifest["definition_sha256"],
        "xlsx_sha256": manifest["files"]["workbook.xlsx"]["sha256"],
        "projection_sha256": hashlib.sha256(encoded(figures)).hexdigest(),
        "history_sha256": manifest["files"]["history/definition-1.json"]["sha256"],
        "exact_text_cells": customer["exact_text_cells"]
    }


def receive_existing(archive, directory, sha256, manifest_sha256=None):
    from stellar_forge_buyer.analysis_workbooks import WorkbookDelivery
    if directory.exists():
        delivered = WorkbookDelivery(directory, expected_manifest_sha256=manifest_sha256)
        delivered.inspect()
        delivered.verify_archive(archive)
        return delivered
    return WorkbookDelivery.receive(archive, directory, expected_sha256=sha256)


def first_line(process, timeout=30):
    selector = selectors.DefaultSelector()
    selector.register(process.stdout, selectors.EVENT_READ)
    raw, deadline = b"", time.monotonic() + timeout
    try:
        while b"\n" not in raw:
            need(process.poll() is None, "ANALYSIS_SERVICE_EXITED_BEFORE_LISTEN")
            left = deadline - time.monotonic()
            need(left > 0, "ANALYSIS_LISTEN_TIMEOUT")
            if not selector.select(min(left, 1)):
                continue
            block = os.read(process.stdout.fileno(), 4096)
            need(block, "ANALYSIS_SERVICE_EMPTY_START")
            raw += block
            need(len(raw) <= 65536, "ANALYSIS_START_SIZE_LIMIT")
        return json.loads(raw.split(b"\n", 1)[0])
    finally:
        selector.close()


def installed_modules(installed):
    expected = Path(installed["directory"]).absolute()
    need(os.path.abspath(sys.executable) == os.path.abspath(installed["python"])
         and Path(sys.prefix).absolute().is_relative_to(expected),
         "SELECTED_PRIVATE_INTERPRETER_REQUIRED")
    modules = {}
    for name in ("stellar_forge_buyer", "stellar_forge_buyer.analysis_workspace",
                 "stellar_forge_buyer.analysis_workbooks", "openpyxl"):
        module = importlib.import_module(name)
        location = Path(module.__file__).resolve()
        need(location.is_relative_to(expected.resolve()), "MODULE_OUTSIDE_PRIVATE_INSTALLATION", name)
        modules[name] = {"path": str(location), "sha256": digest(location)}
    versions = {name: importlib.metadata.version(name)
                for name in ("stellar-forge-buyer", "openpyxl", "et-xmlfile")}
    need(versions["openpyxl"] == "3.1.5"
         and versions["stellar-forge-buyer"] == installed["source"]["version"],
         "ORIGINAL_PACKAGE_VERSION_CHANGED")
    return {"python": sys.executable, "prefix": sys.prefix, "versions": versions, "modules": modules}


def simple_json(argv, *, fd, timeout=60):
    result = subprocess.run(argv, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
        stderr=subprocess.PIPE, timeout=timeout, env=environment(), pass_fds=(fd,))
    need(result.returncode == 0, "ORIGINAL_CLI_FAILED", {"exit_code": result.returncode})
    need(len(result.stdout) <= MAX_JSON, "ORIGINAL_CLI_RESULT_SIZE_LIMIT")
    return json.loads(result.stdout)


def customer_use(request):
    from stellar_forge_buyer.workbook_client import WorkspaceWorkbookClient
    customer, installed = request["customer"], request["installation"]
    environment_result = installed_modules(installed)
    directory = private(request["directory"])
    database = absolute(customer["database"])
    fd = request["lease_fd"]
    os.fstat(fd)
    command = [sys.executable, "-m", "stellar_forge_buyer.analysis_workspace",
               "--database", str(database), "--checkout", request["source"]["sourceRoot"],
               "--node", request["node"]]
    before = simple_json([*command, "status"], fd=fd)
    need(before["workspace_id"] == customer["workspace_id"], "ORIGINAL_WORKSPACE_CHANGED")
    service_log = directory / "analysis-service.stderr"
    process = None
    built = None
    with service_log.open("ab") as errors:
        os.fchmod(errors.fileno(), 0o600)
        process = subprocess.Popen([*command, "serve", "--port", "0"],
            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=errors,
            env=environment(), pass_fds=(fd,), start_new_session=True)
        try:
            launch = first_line(process)
            need(launch.get("status") == "LISTENING", "ANALYSIS_SERVICE_NOT_LISTENING")
            control = read(launch["private_launch"], 65536)
            need(control["pid"] == process.pid and control["workspace_id"] == customer["workspace_id"],
                 "ANALYSIS_ORIGINAL_PROCESS_CHANGED")
            # The original client reads the capability privately from this file.
            client = WorkspaceWorkbookClient(database, workspace_id=customer["workspace_id"])
            state = client.call("/api/state")
            need(state["workspace_id"] == customer["workspace_id"], "ANALYSIS_WORKSPACE_CHANGED")
            definition = client.call("/api/workbooks/" + customer["workbook_name"]
                                     + "?revision=" + str(customer["revision"]))
            original_job = client.call("/api/jobs/" + customer["delivery_id"])
            need(original_job["status"] == "COMPLETE" and original_job["kind"] == "workbook",
                 "COMPLETED_ORIGINAL_WORKBOOK_REQUIRED")
            original_record = next(x for x in original_job["result"]["files"] if x["name"] == "handoff.zip")
            need(original_record["sha256"] == customer["archive_sha256"],
                 "ORIGINAL_WORKBOOK_ARCHIVE_CHANGED")
            archive = directory / "original-handoff.zip"
            client.download(customer["delivery_id"], original_record, archive)
            need(digest(archive) == customer["archive_sha256"], "AUTHENTICATED_DOWNLOAD_CHANGED")
            original_dir = directory / "original-workbook"
            receive_existing(archive, original_dir, customer["archive_sha256"], customer["manifest_sha256"])
            original = workbook_projection(original_dir, customer)
            need(original["definition_sha256"] == definition["definition_sha256"],
                 "SAVED_CUSTOMER_DEFINITION_CHANGED")
            preview_hashes = []
            for index in range(len(definition["definition"]["sheets"])):
                preview = client.call("/api/workbooks/" + customer["workbook_name"]
                    + "/preview?revision=" + str(customer["revision"]) + "&sheet=" + str(index) + "&limit=20")
                preview_hashes.append(hashlib.sha256(encoded(preview)).hexdigest())
            if request["build"]:
                spec = {"id": request["build_id"], "name": customer["workbook_name"],
                        "revision": customer["revision"], "definition_sha256": definition["definition_sha256"]}
                job = client.call("/api/start/workbook", spec)
                deadline = time.monotonic() + 300
                while job["status"] == "RUNNING":
                    need(time.monotonic() < deadline, "ORIGINAL_WORKBOOK_BUILD_TIMEOUT")
                    time.sleep(0.2)
                    job = client.call("/api/jobs/" + spec["id"])
                need(job["status"] == "COMPLETE", "ORIGINAL_WORKBOOK_BUILD_FAILED",
                     {"id": job["id"], "status": job["status"], "error": job.get("error")})
                record = next(x for x in job["result"]["files"] if x["name"] == "handoff.zip")
                built_archive = directory / "rebuilt-handoff.zip"
                client.download(job["id"], record, built_archive)
                received = directory / "rebuilt-workbook"
                receive_existing(built_archive, received, record["sha256"], job["result"]["manifest_sha256"])
                built = workbook_projection(received, customer, original=False)
                need(built["projection_sha256"] == original["projection_sha256"]
                     and built["definition_sha256"] == original["definition_sha256"],
                     "REBUILT_CUSTOMER_CONTENT_CHANGED")
                built["archive_sha256"] = record["sha256"]
                built["archive_bytes"] = record["bytes"]
            elif request.get("expected_build_id"):
                prior = client.call("/api/jobs/" + request["expected_build_id"])
                need(prior["status"] == "COMPLETE" and prior["result"]["revision"] == customer["revision"]
                     and prior["result"]["definition_sha256"] == original["definition_sha256"],
                     "REOPENED_BUILD_CHANGED")
            opened = {"status": launch["status"], "address": launch["address"],
                      "workspace_id": state["workspace_id"]}
        finally:
            if process is not None:
                stop(process)
                if process.stdout:
                    process.stdout.close()
            errors.flush()
            os.fsync(errors.fileno())
    need(process.returncode == 0, "ANALYSIS_SERVICE_CLOSE_FAILED", {"exit_code": process.returncode})
    after = simple_json([*command, "status"], fd=fd)
    need(before["workspace_id"] == after["workspace_id"]
         and before["sources"] == after["sources"] and before["queries"] == after["queries"]
         and before["views"] == after["views"] and before["workbooks"] == after["workbooks"],
         "SAVED_CUSTOMER_STATE_CHANGED")
    # Exercise the installed original console script, including its actual shebang.
    entry = installed["entrypoints"]["stellar-forge-workbook"]["path"]
    inspected = simple_json([entry, "inspect", "--delivery", str(original_dir),
                             "--manifest-sha256", customer["manifest_sha256"]], fd=fd)
    need(inspected["manifest_sha256"] == original["manifest_sha256"], "INSTALLED_WORKBOOK_CLI_CHANGED")
    return {"schema": SCHEMA, "state": "COMPLETED", "environment": environment_result,
            "application": opened, "customer": original, "built": built,
            "preview_sha256": preview_hashes, "saved_state_preserved": True,
            "original_archive": str(archive), "original_archive_sha256": digest(archive),
            "service_exit_code": process.returncode}


def receive_cli(journey, installed, archive, destination, customer):
    # On a lost reply, inspect the already-published original directory and verify
    # it against the retained archive through WorkbookDelivery in the native use.
    entry = installed["entrypoints"]["stellar-forge-workbook"]["path"]
    if destination.exists():
        arguments = [entry, "inspect", "--delivery", str(destination),
                     "--manifest-sha256", customer["manifest_sha256"]]
    else:
        arguments = [entry, "receive", "--archive", str(archive),
                     "--destination", str(destination), "--sha256", customer["archive_sha256"]]
    result = journey.step("offline-workbook-receive",
        {"installation": installation_id(installed), "archive_sha256": customer["archive_sha256"],
         "destination": str(destination), "manifest_sha256": customer["manifest_sha256"]},
        lambda directory, attempt: journey.command(directory, attempt, arguments, timeout=120))
    need(result["manifest_sha256"] == customer["manifest_sha256"]
         and result["row_count"] == customer["row_count"], "OFFLINE_WORKBOOK_CHANGED")
    return result


def validate_plan(plan):
    need(plan.get("schema") == SCHEMA and isinstance(plan.get("operation_id"), str)
         and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,159}", plan["operation_id"]),
         "RECEIVING_PLAN_INVALID")
    for key in ("engine", "python", "node"):
        path = absolute(plan[key], file=(key == "engine"))
        need(path.exists(), "NATIVE_EXECUTABLE_OR_SOURCE_ABSENT", key)
    need(HEX.fullmatch(plan["engine_sha256"])
         and digest(plan["engine"]) == plan["engine_sha256"], "PYTHON_ENGINE_CHANGED")
    absolute(plan["output"])
    extras = plan.get("extras", ["workbooks"])
    need(isinstance(extras, list) and extras and "workbooks" in extras
         and all(isinstance(x, str) and re.fullmatch(r"[a-z][a-z0-9_-]{0,63}", x) for x in extras)
         and len(extras) == len(set(extras)), "DECLARED_PYTHON_EXTRAS_REQUIRED")
    source, offline = plan["source"], plan["offline_source"]
    need(isinstance(source, dict) and isinstance(offline, dict)
         and source["commit"] == offline["commit"]
         and source["manifest_sha256"] == offline["manifest_sha256"]
         and absolute(source["installation_directory"]) != absolute(offline["installation_directory"]),
         "EXACT_OFFLINE_SOURCE_SELECTION_REQUIRED")
    for selected in (source, offline):
        absolute(selected["installation_directory"])
        module = absolute(selected["installer_module"], file=True)
        need(HEX.fullmatch(selected["installer_sha256"])
             and digest(module) == selected["installer_sha256"], "INSTALLER_SOURCE_CHANGED")
    applications = plan.get("applications")
    if applications is not None:
        helper = absolute(applications["helper"], file=True)
        controller = absolute(applications["sourceRoot"])
        need(COMMIT.fullmatch(applications["sourceCommit"])
             and HEX.fullmatch(applications["helper_sha256"])
             and digest(helper) == applications["helper_sha256"],
             "MANAGED_APPLICATION_SOURCE_PIN_REQUIRED")
        files = applications["files"]
        need(isinstance(files, list) and files, "MANAGED_APPLICATION_CLOSURE_REQUIRED")
        names = set()
        for item in files:
            path = PurePosixPath(item["path"])
            need(not path.is_absolute() and ".." not in path.parts and str(path) not in names
                 and HEX.fullmatch(item["sha256"]), "MANAGED_APPLICATION_FILE_PIN_INVALID")
            names.add(str(path))
            target = controller.joinpath(*path.parts)
            need(target.resolve().is_relative_to(controller.resolve())
                 and target.is_file() and digest(target) == item["sha256"],
                 "MANAGED_APPLICATION_SOURCE_CHANGED", str(path))
    capacity = plan.get("capacity")
    if capacity is not None:
        need(isinstance(capacity, dict) and type(capacity["process_peak_bytes"]) is int
             and capacity["process_peak_bytes"] > 0 and type(capacity["memory_backed"]) is bool,
             "ORIGINAL_RECEIVING_CAPACITY_REQUIRED")
        absolute(capacity["request"])
        allocation = absolute(capacity["directory"])
        for selected in (source, offline):
            need(absolute(selected["installation_directory"]).is_relative_to(allocation),
                 "RECEIVED_RUNTIME_OUTSIDE_ALLOCATION")
    customer = plan["customer"]
    absolute(customer["database"])
    absolute(customer["archive"], file=customer.get("directory_transfer") is None)
    for key in ("archive_sha256", "manifest_sha256", "xlsx_sha256"):
        need(HEX.fullmatch(customer[key]), "CUSTOMER_PIN_REQUIRED", key)
    need(re.fullmatch(r"[a-f0-9]{32}", customer["workspace_id"])
         and re.fullmatch(r"[a-f0-9]{32}", customer["delivery_id"])
         and re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,63}", customer["workbook_name"])
         and type(customer["revision"]) is int and customer["revision"] > 0
         and type(customer["row_count"]) is int and customer["row_count"] >= 0
         and type(customer["payload_files"]) is int and customer["payload_files"] > 0,
         "CUSTOMER_EXPECTATION_INVALID")
    need(isinstance(customer["exact_text_cells"], list)
         and all(isinstance(x, dict) and isinstance(x.get("sheet"), str)
                 and isinstance(x.get("value"), str) for x in customer["exact_text_cells"]),
         "CUSTOMER_TEXT_EXPECTATION_INVALID")
    transfer = customer.get("directory_transfer")
    if transfer is not None:
        import zipfile
        need(isinstance(transfer, dict)
             and all(isinstance(transfer.get(key), str) and HEX.fullmatch(transfer[key])
                     for key in ("sha256", "manifest_sha256")),
             "CUSTOMER_TRANSFER_PIN_REQUIRED")
        transfer_path = absolute(transfer["archive"], file=True)
        need(digest(transfer_path) == transfer["sha256"], "CUSTOMER_TRANSFER_ARCHIVE_CHANGED")
        destination = Path(customer["database"]).parent
        if capacity is not None:
            need(destination.is_relative_to(absolute(capacity["directory"])),
                 "CUSTOMER_HYDRATION_OUTSIDE_ALLOCATION")
        need(destination != transfer_path.parent and not transfer_path.is_relative_to(destination),
             "CUSTOMER_TRANSFER_DESTINATION_OVERLAP")
        with zipfile.ZipFile(transfer_path) as archive:
            customer_transfer_inventory(customer, archive)
    else:
        validate_customer_handoff(customer)


def validate_customer_handoff(customer):
    absolute(customer["archive"], file=True)
    need(digest(customer["archive"]) == customer["archive_sha256"], "CUSTOMER_ARCHIVE_CHANGED")


def run_plan(plan):
    validate_plan(plan)
    root = private(plan["output"])
    need(not Path(plan["customer"]["database"]).is_relative_to(root)
         and not root.is_relative_to(Path(plan["customer"]["database"])),
         "CUSTOMER_AND_RUNTIME_DIRECTORIES_MUST_BE_SEPARATE")
    with operation_lease(root) as fd:
        journey = Journey(plan, root, fd)
        source = select_source(journey, "original-source", plan["source"])
        primary_root = Path(source["directory"]) / "python-runtime"
        candidate = cohort(journey, "primary", primary_root, source)
        first_capacity = receiving_capacity(journey, "primary-capacity", candidate, source)
        primary = install_cohort(journey, "primary", primary_root, candidate, source)
        activated = journey.engine("primary-activate", primary_root, "activate",
                       ["--installation", installation_id(primary), "--expected-installation", ""])
        activation_replay = journey.engine("primary-activate-replay", primary_root, "activate",
                       ["--installation", installation_id(primary), "--expected-installation", ""])
        need(activation_replay == activated, "ACTIVATION_REPLAY_CHANGED")
        transferred = hydrate_customer_transfer(journey, primary)
        validate_customer_handoff(plan["customer"])
        copied = receive_customer_copy(journey, primary)
        customer = {**plan["customer"], "database": copied["database"]}
        first = use_customer(journey, "primary-customer", primary, source, customer, build=True)
        managed_primary = use_managed_application(journey, "primary-applications", primary, source, customer)
        offline_capacity = receiving_capacity(journey, "offline-capacity", candidate, source, installed=primary)
        bundle_path = root / "python-workbooks.bundle.zip"
        bundled = journey.engine("offline-bundle", primary_root, "bundle",
                                 ["--candidate", candidate_id(candidate), "--output", str(bundle_path)])
        need(bundle_path.is_file(), "PYTHON_COHORT_ARCHIVE_ABSENT")
        bundle_sha256 = digest(bundle_path)
        offline_source = select_source(journey, "offline-original-source", plan["offline_source"])
        need(offline_source["sourceCommit"] == source["sourceCommit"]
             and offline_source["sourceManifestSha256"] == source["sourceManifestSha256"]
             and offline_source["profileId"] == source["profileId"], "OFFLINE_SOURCE_SELECTION_CHANGED")
        need(offline_source["directory"] != source["directory"], "FRESH_SOURCE_RUNTIME_REQUIRED")
        offline_root = Path(offline_source["directory"]) / "python-runtime"
        received = journey.engine("offline-receive", offline_root, "receive",
                                  ["--archive", str(bundle_path)])
        need(candidate_id(received) == candidate_id(candidate)
             and received["source"] == candidate["source"], "OFFLINE_COHORT_CHANGED")
        offline = install_cohort(journey, "offline", offline_root, received, offline_source)
        need(offline["python"] != primary["python"], "FRESH_PRIVATE_INTERPRETER_REQUIRED")
        journey.engine("offline-activate", offline_root, "activate",
                       ["--installation", installation_id(offline), "--expected-installation", ""])
        destination = root / "customer" / "offline-delivery"
        receive_cli(journey, offline, Path(first["original_archive"]), destination, plan["customer"])
        reopened = use_customer(journey, "offline-customer", offline, offline_source, customer,
                                expected_build_id=first["built"]["id"])
        managed_offline = use_managed_application(journey, "offline-applications", offline, offline_source, customer)
        need(reopened["customer"] == first["customer"]
             and reopened["preview_sha256"] == first["preview_sha256"],
             "OFFLINE_CUSTOMER_REOPEN_CHANGED")
        upgrade = None
        if plan.get("upgrade"):
            next_source = select_source(journey, "upgrade-source", plan["upgrade"])
            need(next_source["sourceCommit"] != source["sourceCommit"],
                 "ACTUAL_NEWER_SOURCE_REQUIRED")
            next_candidate = cohort(journey, "upgrade", primary_root, next_source)
            next_install = install_cohort(journey, "upgrade", primary_root, next_candidate, next_source)
            need(installation_id(next_install) != installation_id(primary), "NEW_SOURCE_INSTALLATION_REQUIRED")
            journey.engine("upgrade-activate", primary_root, "activate", [
                "--installation", installation_id(next_install),
                "--expected-installation", installation_id(primary)])
            next_use = use_customer(journey, "upgrade-customer", next_install, next_source, customer,
                                   expected_build_id=first["built"]["id"])
            need(next_use["customer"] == first["customer"], "UPGRADED_CUSTOMER_CHANGED")
            upgrade = {"source": next_source, "candidate": candidate_id(next_candidate),
                       "installation": installation_id(next_install), "customer": next_use["customer"]}
        need(digest(plan["customer"]["archive"]) == plan["customer"]["archive_sha256"],
             "ORIGINAL_CUSTOMER_ARCHIVE_CHANGED")
        result = {"schema": SCHEMA, "operation_id": plan["operation_id"], "state": "COMPLETED",
                  "plan_sha256": journey.plan_sha256, "completed_at": stamp(), "source": source,
                  "candidate": candidate_id(candidate), "extras": candidate["extras"],
                  "packages": candidate["packages"],
                  "wheel_bytes": candidate["wheelBytes"], "unpacked_bytes": candidate["unpackedBytes"],
                  "primary_installation": installation_id(primary),
                  "activation": {"initial": activated, "replayed": activation_replay},
                  "customer_copy": copied,
                  "offline_installation": installation_id(offline),
                  "bundle": {"path": str(bundle_path), "bytes": bundle_path.stat().st_size,
                             "sha256": bundle_sha256, "engine_result": bundled},
                  "primary": first, "offline": reopened, "offline_source": offline_source}
        if first_capacity is not None:
            result["capacity"] = {"primary": first_capacity, "offline": offline_capacity,
                                  "completed_allocation": allocated_tree(plan["capacity"]["directory"])}
        if managed_primary is not None:
            result["applications"] = {"primary": managed_primary, "offline": managed_offline}
        if transferred is not None:
            result["customer_transfer"] = transferred
        if upgrade is not None:
            result["upgrade"] = upgrade
        write(root / "result.json", result)
        journey.state.update(state="COMPLETED", result_sha256=digest(root / "result.json"),
                             completed_at=stamp())
        journey.save()
        return result


def main():
    parser = argparse.ArgumentParser(
        description="Receive and operate an exact original Python workbook runtime in one admitted native journey.")
    commands = parser.add_subparsers(dest="command", required=True)
    run = commands.add_parser("run", help="Continue the same source-pinned private receiving operation")
    run.add_argument("--plan", required=True, help="Private original source/customer/engine plan JSON")
    internal = commands.add_parser("_customer", help=argparse.SUPPRESS)
    internal.add_argument("--request", required=True)
    copy_command = commands.add_parser("_customer-copy", help=argparse.SUPPRESS)
    copy_command.add_argument("--request", required=True)
    receive_command = commands.add_parser("_customer-receive", help=argparse.SUPPRESS)
    receive_command.add_argument("--request", required=True)
    args = parser.parse_args()
    os.umask(0o077)
    signal.signal(signal.SIGTERM, interrupted)
    signal.signal(signal.SIGINT, interrupted)
    try:
        if args.command == "_customer":
            result = customer_use(read(args.request))
        elif args.command == "_customer-copy":
            result = copy_closed_customer(read(args.request))
        elif args.command == "_customer-receive":
            result = receive_customer_archive(read(args.request))
        else:
            result = run_plan(read(args.plan))
        print(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")), flush=True)
    except KeyboardInterrupt:
        print(json.dumps({"state": "INTERRUPTED", "error": "ORIGINAL_OPERATION_INTERRUPTED"}), file=sys.stderr)
        raise SystemExit(130)
    except Exception as error:
        print(json.dumps({"state": "FAILED", "error": getattr(error, "code", type(error).__name__),
                          "detail": getattr(error, "detail", None)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
