"""DBnomics economic source, native Analysis and editable workbook receiving.

The original Node DBnomics reader verifies economic source/calculation semantics.
AnalysisWorkspace, DatasetChanges, Workbooks and WorkspaceArchives own all
database, review, export and receiving operations.
"""
from __future__ import annotations
import argparse
import base64
from contextlib import closing
from copy import deepcopy
from datetime import date, datetime, time as datetime_time, timedelta, timezone
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import stat
import tempfile
import time
import uuid
import zipfile

from .analysis_workspace import AnalysisWorkspace
from .analysis_workbooks import WorkbookDelivery
from .errors import StateError
from .query_export import QueryExport
from .response_documents import validate_manifest, document_path, retained_document_files
from .result import FileDelivery
from .runtime import Runtime
from .storage import dump, lock, parse, private_directory, read, sync_directory
from .workspace_archive_receive import restore_archive

SOURCE_SCHEMA = "stellar-forge.dbnomics-analysis-source.v1"
BUNDLE_SCHEMA = "stellar-forge.dbnomics-analysis-delivery.v1"
WORKFLOW_SCHEMA = "stellar-forge.dbnomics-economic-workflow.v1"
RETURN_SCHEMA = "stellar-forge.dbnomics-workbook-return.v1"
HASH = re.compile(r"[a-f0-9]{64}")
NAME = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}")
MAX_BYTES = 256 * 1024 * 1024
TABLE = "economic_records"
REPORT = "EconomicBriefing"
WORKBOOK = "EconomicRecords"
KIND_QUERIES = (
    ("Observations", "observation", "Every original observation and missing-value state."),
    ("Calculations", "calculation", "Every selected saved calculation, with its exact calendar basis."),
    ("Series", "series", "Separate series identities and frequencies; annual and monthly populations remain distinct."),
    ("Notes", "note", "The exact selected analyst note revision."),
)
SUMMARY_COLUMNS = (
    "record_id", "series_name", "frequency", "period", "value_text", "value_state",
    "report_title", "calculated_value", "calculation", "calculation_reason",
    "base_period", "base_value_text", "note_author", "note_tag", "note_text",
)


def need(condition, code):
    if not condition:
        raise StateError("DBNOMICS_" + code)


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False,
                      separators=(",", ":")).encode("utf-8")


def sha(value):
    return hashlib.sha256(value).hexdigest()


def pin(path, maximum=MAX_BYTES):
    path = Path(path).expanduser().absolute()
    with os.fdopen(os.open(path, os.O_RDONLY | os.O_NOFOLLOW), "rb") as stream:
        before = os.fstat(stream.fileno())
        need(stat.S_ISREG(before.st_mode), "REGULAR_FILE_REQUIRED")
        need(before.st_size <= maximum, "SOURCE_FILE_LIMIT")
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
        after = os.fstat(stream.fileno())
        need(all(getattr(before, k) == getattr(after, k) for k in
                 ("st_dev", "st_ino", "st_size", "st_mtime_ns", "st_ctime_ns")), "SOURCE_CHANGED")
    return {"path": str(path), "bytes": before.st_size, "sha256": digest}


def keep(path, value, *, identity=None):
    path = Path(path)
    if path.exists():
        old = parse(read(path, MAX_BYTES))
        left = old if identity is None else {k: old.get(k) for k in identity}
        right = value if identity is None else {k: value.get(k) for k in identity}
        need(canonical(left) == canonical(right), "RETAINED_RECORD_CHANGED")
        return old
    dump(path, value, exclusive=True)
    return value


def write_exact(path, data):
    path = Path(path)
    private_directory(path.parent)
    expected = {"bytes": len(data), "sha256": sha(data)}
    if path.exists():
        actual = pin(path)
        need(all(actual[k] == v for k, v in expected.items()), "RETAINED_FILE_CHANGED")
        return {**expected, "path": str(path)}
    fd, temporary = tempfile.mkstemp(prefix=".dbnomics-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as output:
            output.write(data)
            output.flush()
            os.fsync(output.fileno())
        try:
            os.link(temporary, path)
        except FileExistsError:
            actual = pin(path)
            need(all(actual[k] == v for k, v in expected.items()), "RETAINED_FILE_CHANGED")
        os.unlink(temporary)
        sync_directory(path.parent)
    finally:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
    return {**expected, "path": str(path)}


def delivery(path, record):
    return FileDelivery({"returncode": 0, "native": {"status": "RETAINED_SOURCE_DOCUMENT",
        "file": {"path": str(path), "bytes": record["bytes"], "sha256": record["sha256"],
                 "contentType": record.get("contentType", "application/octet-stream")}}}, path, MAX_BYTES)


def copy_delivery(context, destination):
    destination = Path(destination)
    private_directory(destination.parent)
    with context as (source, record):
        expected = {k: record[k] for k in ("bytes", "sha256")}
        if destination.exists():
            actual = pin(destination)
            need(all(actual[k] == v for k, v in expected.items()), "DELIVERY_CHANGED")
            return {**record, "path": str(destination)}
        fd, temporary = tempfile.mkstemp(prefix=".dbnomics-delivery-", dir=destination.parent)
        try:
            count, checksum = 0, hashlib.sha256()
            with os.fdopen(fd, "wb") as output:
                while block := source.read(1024 * 1024):
                    count += len(block)
                    need(count <= expected["bytes"], "DELIVERY_CHANGED")
                    output.write(block)
                    checksum.update(block)
                need(count == expected["bytes"] and checksum.hexdigest() == expected["sha256"], "DELIVERY_CHANGED")
                output.flush()
                os.fsync(output.fileno())
            try:
                os.link(temporary, destination)
            except FileExistsError:
                actual = pin(destination)
                need(all(actual[k] == v for k, v in expected.items()), "DELIVERY_CHANGED")
            os.unlink(temporary)
            sync_directory(destination.parent)
        finally:
            try:
                os.unlink(temporary)
            except FileNotFoundError:
                pass
    return {**record, "path": str(destination)}


def checked_download(context, expected):
    with context as (source, record):
        need(all(record[k] == expected[k] for k in ("bytes", "sha256")), "RECEIVED_EXPORT_CHANGED")
        count, checksum = 0, hashlib.sha256()
        while block := source.read(1024 * 1024):
            count += len(block)
            checksum.update(block)
        need(count == expected["bytes"] and checksum.hexdigest() == expected["sha256"], "RECEIVED_EXPORT_CHANGED")
    return {"bytes": count, "sha256": checksum.hexdigest()}


def native(checkout, node, args):
    result = Runtime(checkout, node=node, timeout=900, output_limit=8 * 1024 * 1024).run(args, target="dbnomics")
    need(result.returncode == 0 and not result.error and isinstance(result.output, dict),
         "NATIVE_COMMAND_FAILED:" + str(result.error or result.output))
    return result.output


class DbnomicsAnalysisSource:
    """Complete Analysis document with the original economic handover inside it."""
    def __init__(self, directory, *, expected_document_id=None):
        self.directory = private_directory(Path(directory).expanduser().absolute(), create=False)
        self.manifest_file = self.directory / "source.json"
        raw = read(self.manifest_file, 1024 * 1024)
        self.manifest = validate_manifest(parse(raw))
        self.document_id = sha(raw)
        need(expected_document_id is None or self.document_id == expected_document_id, "SOURCE_DOCUMENT_CHANGED")
        matches = [x for x in self.manifest["originals"] if x["role"] == "analysis-receipt"]
        need(len(matches) == 1, "ANALYSIS_RECEIPT_REQUIRED")
        self.receipt = self.file(matches[0]).json(limit=32 * 1024 * 1024)
        need(self.receipt.get("schema") == SOURCE_SCHEMA, "ANALYSIS_RECEIPT_CHANGED")
        selection = self.receipt.get("selection")
        need(isinstance(selection, dict) and self.receipt.get("identity", {}).get("selectionSha256") == sha(canonical(selection)),
             "ANALYSIS_SELECTION_CHANGED")
        columns = self.receipt.get("columns")
        need(isinstance(columns, list) and len(columns) == len(set(columns)) and
             all(isinstance(k, str) and re.fullmatch(r"[a-z][a-z0-9_]*", k) for k in columns), "ANALYSIS_COLUMNS_CHANGED")
        need(self.receipt.get("mapping") == {"collection": "", "columns": [
            {"name": k, "pointer": "/" + k, "type": "text", "missing": "error"} for k in columns]},
             "ANALYSIS_MAPPING_CHANGED")
        need(all(k in columns for k in self.receipt.get("workbookColumns", [])) and
             "record_id" in self.receipt.get("workbookColumns", []) and "review_note" in columns,
             "WORKBOOK_IDENTITY_COLUMNS_REQUIRED")

    @classmethod
    def create(cls, workspace, directory, selection, *, operation_id, checkout, node="node"):
        need(isinstance(operation_id, str) and NAME.fullmatch(operation_id), "SOURCE_OPERATION_REQUIRED")
        selected = private_directory(Path(directory).expanduser().absolute())
        request = selected / "native-selection.json"
        keep(request, {"id": operation_id, "selection": selection})
        result = native(checkout, node, ["analysis-source", "--workspace", str(Path(workspace).expanduser().absolute()),
            "--input", str(request), "--out", str(selected / "source")])
        source = cls(selected / "source", expected_document_id=result["documentId"])
        keep(selected / "native-source.json", result, identity=("id", "documentId", "selection", "rowCount", "counts"))
        source.verify()
        return source

    @classmethod
    def from_bundle(cls, file, directory, *, expected_sha256):
        pinned = pin(file, MAX_BYTES * 2)
        need(pinned["sha256"] == expected_sha256, "ANALYSIS_HANDOVER_CHANGED")
        bundle = parse(read(Path(file), MAX_BYTES * 2))
        need(bundle.get("schema") == BUNDLE_SCHEMA and isinstance(bundle.get("files"), list), "ANALYSIS_HANDOVER_INVALID")
        root = private_directory(Path(directory).expanduser().absolute())
        seen, total = set(), 0
        for entry in bundle["files"]:
            name = entry.get("path")
            need(isinstance(name, str) and (name == "source.json" or re.fullmatch(r"files/[a-f0-9]{64}\.(json|csv|bin)", name)),
                 "ANALYSIS_HANDOVER_PATH_INVALID")
            need(name not in seen and type(entry.get("bytes")) is int and 0 <= entry["bytes"] <= MAX_BYTES,
                 "ANALYSIS_HANDOVER_FILE_INVALID")
            seen.add(name)
            try:
                data = base64.b64decode(entry["contentBase64"], validate=True)
            except (ValueError, TypeError):
                raise StateError("DBNOMICS_ANALYSIS_HANDOVER_BYTES_INVALID") from None
            need(len(data) == entry["bytes"] and sha(data) == entry["sha256"], "ANALYSIS_HANDOVER_CHANGED")
            total += len(data)
            need(total <= MAX_BYTES + 1024 * 1024, "ANALYSIS_HANDOVER_LIMIT")
            write_exact(root / name, data)
        source = cls(root, expected_document_id=bundle.get("documentId"))
        required = {"source.json", source.manifest["document"]["path"],
                    *(x["path"] for x in source.manifest["originals"])}
        need(seen == required and source.receipt["selection"] == bundle.get("selection"), "ANALYSIS_HANDOVER_MEMBERS_CHANGED")
        source.verify()
        return source

    def file(self, item):
        return delivery(self.directory / item["path"], item)

    @property
    def mapping(self):
        return deepcopy(self.receipt["mapping"])

    def verify(self):
        items = {x["path"]: x for x in [self.manifest["document"], *self.manifest["originals"]]}
        for item in items.values():
            with self.file(item).open():
                pass
        with self.file(self.manifest["document"]).open() as (stream, _):
            import io
            previous = csv.field_size_limit(32 * 1024 * 1024)
            try:
                wrapper = io.TextIOWrapper(stream, encoding="utf-8", newline="")
                rows = csv.DictReader(wrapper)
                need(rows.fieldnames == self.receipt["columns"], "SOURCE_HEADER_CHANGED")
                count, ids, counts = 0, set(), {}
                for row in rows:
                    need(None not in row and None not in row.values(), "SOURCE_ROW_WIDTH_CHANGED")
                    need(row["record_id"] not in ids and HASH.fullmatch(row["record_id"]), "SOURCE_RECORD_ID_CHANGED")
                    ids.add(row["record_id"])
                    counts[row["record_kind"]] = counts.get(row["record_kind"], 0) + 1
                    count += 1
                need(count == self.receipt["rowCount"] and
                     all(counts.get(k, 0) == n for k, n in self.receipt["counts"].items()), "SOURCE_RECORD_COUNT_CHANGED")
            finally:
                if "wrapper" in locals(): wrapper.detach()
                csv.field_size_limit(previous)
        return {"document_id": self.document_id, "rows": count, "counts": self.receipt["counts"],
                "files": len(items), "bytes": sum(x["bytes"] for x in items.values())}

    def register(self, workspace, name=TABLE, *, expected_revision=0):
        return workspace.responses.register(name, "document", manifest=self.manifest_file,
                                             format="csv", expected_revision=expected_revision)

    def original(self, destination):
        archive = next(x for x in self.receipt["archives"] if x["role"] == "dbnomics-source-handover")
        target = Path(destination).expanduser().absolute()
        refs = {x["role"]: x for x in self.manifest["originals"]}
        if target.exists():
            result = pin(target)
            need(all(result[k] == archive[k] for k in ("bytes", "sha256")), "ORIGINAL_HANDOVER_CHANGED")
            return result
        private_directory(target.parent)
        fd, temporary = tempfile.mkstemp(prefix=".dbnomics-source-", dir=target.parent)
        try:
            count, checksum = 0, hashlib.sha256()
            with os.fdopen(fd, "wb") as output:
                for part in archive["parts"]:
                    item = refs[part["role"]]
                    need(all(item[k] == part[k] for k in ("bytes", "sha256")), "ORIGINAL_PART_CHANGED")
                    with self.file(item).open() as (stream, _):
                        while block := stream.read(1024 * 1024):
                            output.write(block);count += len(block);checksum.update(block)
                need(count == archive["bytes"] and checksum.hexdigest() == archive["sha256"], "ORIGINAL_HANDOVER_CHANGED")
                output.flush();os.fsync(output.fileno())
            try:
                os.link(temporary, target)
            except FileExistsError:
                actual = pin(target)
                need(all(actual[k] == archive[k] for k in ("bytes", "sha256")), "ORIGINAL_HANDOVER_CHANGED")
            os.unlink(temporary);sync_directory(target.parent)
        finally:
            try: os.unlink(temporary)
            except FileNotFoundError: pass
        return {**archive, "path": str(target)}



def workbook_projection(directory, file, *, expected_sha256, workbook, operation,
                        source_document_id, workbook_columns, side):
    """Project every selected worksheet value; retain both original XLSX files.

    This is a declared file import for the native response mapper. DatasetChanges
    remains the owner of identity matching, scalar differences and review.
    """
    from openpyxl import load_workbook

    root = private_directory(Path(directory).expanduser().absolute())
    selected = pin(file, 128 * 1024 * 1024)
    need(selected["sha256"] == expected_sha256, "WORKBOOK_RETURN_CHANGED")
    inspected = workbook.inspect()
    logical = inspected["sheets"]
    need(len(logical) == 1 and logical[0]["name"] == "Records", "ORIGINAL_RECORDS_WORKSHEET_REQUIRED")
    columns = logical[0]["columns"]
    headers = [x["header"] for x in columns]
    need(headers == workbook_columns and all(x["source"] == x["header"] for x in columns),
         "ORIGINAL_WORKBOOK_COLUMNS_CHANGED")
    names = logical[0]["worksheets"]
    need(isinstance(names, list) and names and len(names) == len(set(names)), "ORIGINAL_WORKSHEET_NAMES_CHANGED")
    with delivery(file, selected).open() as (stream, _), zipfile.ZipFile(stream) as archive:
        members = archive.infolist()
        need(len(members) <= 10000 and sum(x.file_size for x in members) <= MAX_BYTES and
             len({x.filename for x in members}) == len(members) and
             all(not x.flag_bits & 1 for x in members), "WORKBOOK_PACKAGE_LIMIT")
    rows, locations, sheet_rows, extras = [], [], {}, []
    with delivery(file, selected).open() as (stream, _):
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            need(all(name in book.sheetnames for name in names), "WORKBOOK_RECORDS_SHEET_MISSING")
            extras = [{"name": s.title, "state": s.sheet_state} for s in book.worksheets if s.title not in names]
            for name in names:
                sheet = book[name]
                # Read the actual worksheet XML rather than trusting a stale
                # producer-supplied dimension range in a returned workbook.
                sheet.reset_dimensions()
                iterator = sheet.iter_rows()
                try:
                    first = next(iterator)
                except StopIteration:
                    raise StateError("DBNOMICS_WORKBOOK_HEADER_MISSING:" + name) from None
                actual = [cell.value for cell in first]
                while actual and actual[-1] is None:
                    actual.pop()
                need(len(actual) == len(headers) and all(isinstance(k, str) for k in actual) and
                     len(actual) == len(set(actual)) and set(actual) == set(headers),
                     "WORKBOOK_HEADER_CHANGED:" + name)
                indices = {name: i for i, name in enumerate(actual)}
                count = 0
                for line, cells in enumerate(iterator, start=2):
                    need(line <= 1048576, "WORKBOOK_ROW_ADDRESS_LIMIT")
                    need(len(cells) <= 16384, "WORKBOOK_COLUMN_LIMIT")
                    if not any(cell.value is not None for cell in cells):
                        continue
                    need(not any(cell.value is not None for cell in cells[len(actual):]),
                         "WORKBOOK_UNNAMED_VALUES:" + name + ":" + str(line))
                    row, types = {}, {}
                    for field in headers:
                        i = indices[field]
                        cell = cells[i] if i < len(cells) else None
                        value = cell.value if cell is not None else None
                        kind = cell.data_type if cell is not None else "n"
                        if value is None:
                            kind, value = "blank", None
                        elif kind == "f":
                            need(isinstance(value, str), "WORKBOOK_FORMULA_REPRESENTATION")
                            kind = "formula"
                        elif kind == "e":
                            kind = "error"
                        elif isinstance(value, bool):
                            kind = "boolean"
                        elif isinstance(value, datetime):
                            kind, value = "datetime", value.isoformat()
                        elif isinstance(value, date):
                            kind, value = "date", value.isoformat()
                        elif isinstance(value, datetime_time):
                            kind, value = "time", value.isoformat()
                        elif isinstance(value, timedelta):
                            kind, value = "duration", str(value)
                        elif isinstance(value, str):
                            kind = "string"
                        elif type(value) is int:
                            kind = "number"
                        elif type(value) is float and math.isfinite(value):
                            kind = "number"
                        else:
                            raise StateError("DBNOMICS_WORKBOOK_CELL_UNSUPPORTED:" + name + ":" + str(line) + ":" + field)
                        row[field], types[field] = value, kind
                    need(isinstance(row["record_id"], str) and HASH.fullmatch(row["record_id"]),
                         "WORKBOOK_RECORD_ID_REQUIRED:" + name + ":" + str(line))
                    row["excel_cell_types"] = canonical(types).decode("utf-8")
                    rows.append(row)
                    locations.append({"record_id": row["record_id"], "worksheet": name, "row": line})
                    count += 1
                    need(len(rows) <= 100000, "WORKBOOK_COMPLETE_ROW_LIMIT")
                sheet_rows[name] = count
        finally:
            book.close()
    original = side == "baseline"
    need(not original or len(rows) == logical[0]["row_count"], "ORIGINAL_WORKBOOK_ROWS_CHANGED")
    # A changed/new identity is a native added/removed record. Duplicate keys
    # remain a native comparison error; no fallback row identifier is invented.
    mapping = {"collection": "/rows", "columns": [
        {"name": k, "pointer": "/" + k, "type": "native", "missing": "error"}
        for k in [*headers, "excel_cell_types"]]}
    projected = canonical({"rows": rows})
    need(len(projected) <= 32 * 1024 * 1024, "WORKBOOK_COMPLETE_PROJECTION_LIMIT")
    def retain(data, extension, content_type, role=None):
        record = {"path": "files/" + sha(data) + "." + extension, "bytes": len(data),
                  "sha256": sha(data), "contentType": content_type}
        if role is not None: record["role"] = role
        write_exact(root / record["path"], data)
        return record
    originals = []
    with delivery(file, selected).open() as (stream, _):
        index = 0
        while block := stream.read(32 * 1024 * 1024):
            index += 1
            originals.append(retain(block, "bin", "application/octet-stream", "workbook-part-" + str(index).zfill(3)))
    originals.append(retain(workbook.raw, "json", "application/json", "original-workbook-manifest"))
    receipt = {"schema": RETURN_SCHEMA, "side": side, "return_id": operation["id"],
        "created_at": operation["created_at"], "source_document_id": source_document_id,
        "original_workbook_id": inspected["id"], "original_workbook_manifest_sha256": inspected["manifest_sha256"],
        "original_definition_sha256": inspected["definition_sha256"], "file": selected,
        "worksheets": sheet_rows, "other_worksheets": extras, "row_count": len(rows),
        "columns": [*headers, "excel_cell_types"], "mapping": mapping, "locations": locations,
        "workbook_parts": [{k: item[k] for k in ("role", "bytes", "sha256")} for item in originals if item["role"].startswith("workbook-part-")],
        "semantics": "All selected worksheet rows, including hidden rows, are read. Formulas are retained as formula text and never evaluated. "
            "Blank/string/number/boolean/date/datetime/time/duration/formula/error kinds are kept in excel_cell_types. Styles, comments, "
            "charts and other worksheets remain in the original XLSX bytes; this review compares the named Records cell values. "
            "Original provider records remain in their separate economic_records view."}
    originals.append(retain(canonical(receipt), "json", "application/json", "workbook-projection-receipt"))
    document = retain(projected, "json", "application/json")
    descriptor = {"schema": "stellar-forge.analysis-source-document.v1",
        "title": ("Original" if original else "Returned") + " DBnomics Records workbook " + operation["id"],
        "format": "json", "document": document, "originals": originals, "complete": True,
        "provenance": {"kind": "operator-file", "publisher": "DBnomics economic workbook operator",
            "capturedAt": operation["created_at"], "scope": source_document_id + " / " + inspected["id"] + " / " + side,
            "description": receipt["semantics"]}}
    validate_manifest(descriptor)
    raw = canonical(descriptor)
    write_exact(root / "source.json", raw)
    return {"directory": str(root), "manifest": str(root / "source.json"), "document_id": sha(raw),
            "receipt": receipt, "mapping": mapping, "rows": len(rows)}

class DbnomicsEconomicWorkflow:
    def __init__(self, directory, *, checkout=None, node="node", wait_seconds=900):
        self.root = private_directory(Path(directory).expanduser().absolute())
        self.records = private_directory(self.root / "records")
        self.deliveries = private_directory(self.root / "deliveries")
        self.checkout = str(Path(checkout).expanduser().resolve()) if checkout else None
        self.node, self.wait_seconds = node, wait_seconds
        need(1 <= wait_seconds <= 3600, "WAIT_INVALID")
        self.operation = None

    def read_operation(self):
        self.operation = parse(read(self.root / "operation.json", 1024 * 1024))
        need(self.operation.get("schema") == WORKFLOW_SCHEMA, "WORKFLOW_CHANGED")
        return self.operation

    def receipt(self, name, value, *, identity=None):
        return keep(self.records / (name + ".json"), value, identity=identity)

    def saved(self, name):
        return parse(read(self.records / (name + ".json"), MAX_BYTES))

    def workspace(self):
        return AnalysisWorkspace(self.root / "analysis", checkout=self.checkout, node=self.node)

    def source(self):
        selected = self.saved("source")
        source = DbnomicsAnalysisSource(selected["directory"], expected_document_id=selected["document_id"])
        source.verify()
        return source

    def bind(self, ws):
        self.receipt("workspace", {"workspace_id": ws.identity})

    def job(self, ws, name, kind, body, *, records=None, jobs=None):
        records = records or self.records
        jobs = jobs or self.operation["jobs"]
        request = {"id": jobs[name], **body}
        selected = keep(records / ("request-" + name + ".json"), {"kind": kind, "request": request})
        current_request = selected["request"]
        deadline = time.monotonic() + self.wait_seconds
        while True:
            job = ws.start(selected["kind"], current_request)
            while job["status"] == "RUNNING":
                need(time.monotonic() < deadline, "ORIGINAL_JOB_STILL_RUNNING:" + job["id"])
                time.sleep(0.2)
                job = ws.job(job["id"])
            if job["status"] != "INTERRUPTED" or kind not in {
                    "query", "query-export", "report", "report-export", "workbook", "compare", "change-export"}:
                break
            # These are completed-local-lifetime jobs. The original engine has
            # positively recorded their interruption and owns no running worker.
            # Retain that native job and one exact successor before dispatch.
            follow_path = records / ("interrupted-" + name + "-" + job["id"] + ".json")
            binding = {"kind": kind, "previous_id": job["id"], "request": current_request,
                       "native_status": job["status"], "native_finished_at": job.get("finished_at")}
            if follow_path.exists():
                follow = parse(read(follow_path, 2 * 1024 * 1024))
                need(all(follow.get(k) == v for k, v in binding.items()), "NATIVE_INTERRUPTION_CHANGED")
            else:
                follow = {**binding, "next_id": uuid.uuid4().hex}
                dump(follow_path, follow, exclusive=True)
            current_request = {**selected["request"], "id": follow["next_id"]}
            need(time.monotonic() < deadline, "ORIGINAL_LOCAL_CONTINUATION_PENDING:" + follow["next_id"])
        need(job["status"] == "COMPLETE", "NATIVE_JOB_" + job["status"] + ":" + job["id"] + ":" + str(job.get("error")))
        keep(records / ("job-" + name + ".json"), {k: job.get(k) for k in ("id", "kind", "request", "result", "finished_at")})
        return job

    def prepare(self, source):
        need(isinstance(source, DbnomicsAnalysisSource), "ECONOMIC_SOURCE_REQUIRED")
        source.verify()
        need(source.receipt["rowCount"] <= 100000, "COMPLETE_BRIEFING_ROW_LIMIT")
        selected = {"document_id": source.document_id, "selection": source.receipt["selection"],
                    "source_identity": source.receipt["identity"]}
        path = self.root / "operation.json"
        if path.exists():
            self.read_operation()
            need(self.operation["source"] == selected, "WORKFLOW_SOURCE_CHANGED")
        else:
            names = ("preview", "import", "all_query", "query_export", "report", "report_html",
                     "report_json", "report_csv", "workbook", "workspace_export")
            self.operation = {"schema": WORKFLOW_SCHEMA, "id": uuid.uuid4().hex,
                "created_at": datetime.now(timezone.utc).isoformat(), "source": selected,
                "jobs": {name: uuid.uuid4().hex for name in names},
                "report_request": uuid.uuid4().hex, "report_publication": uuid.uuid4().hex,
                "workbook_request": uuid.uuid4().hex, "restore_operation": uuid.uuid4().hex}
            dump(path, self.operation, exclusive=True)
        self.receipt("source", {"directory": str(source.directory), **selected})
        with closing(self.workspace()) as ws:
            self.bind(ws)
            registered = source.register(ws)
            self.receipt("registration", registered)
            preview = self.job(ws, "preview", "response-preview", {"source": TABLE,
                "source_revision": registered["revision"], "mapping": source.mapping})
            result = {"operation_id": self.operation["id"], "workspace_id": ws.identity,
                "document_id": source.document_id, "source": registered, "preview_id": preview["id"],
                "preview_sha256": preview["result"]["preview_sha256"], "preview": preview["result"],
                "rows": source.receipt["rowCount"], "counts": source.receipt["counts"],
                "selection": source.receipt["selection"], "terms": source.receipt["terms"],
                "workbook_columns": source.receipt["workbookColumns"]}
            return self.receipt("review", result)

    def save_query(self, ws, name, sql):
        spec = {"sql": sql, "parameters_text": "[]", "max_rows": 100000}
        try:
            current = ws.saved_query(name, revision=1)
            need(all(current[k] == v for k, v in spec.items()), "SAVED_QUERY_CHANGED")
        except StateError as error:
            if error.code != "ANALYSIS_QUERY_NOT_FOUND": raise
            current = ws.save_query(name, spec, expected_revision=0)
        self.receipt("query-" + name, {k: current[k] for k in (*spec, "name", "revision")})
        return current

    def finish(self, *, preview_sha256, review_note):
        self.read_operation()
        review = self.saved("review")
        need(preview_sha256 == review["preview_sha256"] and isinstance(review_note, str) and
             1 <= len(review_note) <= 8000, "EXACT_IMPORT_REVIEW_REQUIRED")
        decision = self.receipt("decision", {"preview_id": review["preview_id"],
            "preview_sha256": preview_sha256, "allow_incomplete": False, "note": review_note})
        if (self.records / "complete.json").exists():
            complete = self.saved("complete")
            for group in ("query_export", "briefing", "workbook", "workspace_archive"):
                for item in complete[group]["files"].values():
                    actual = pin(item["path"])
                    need(all(actual[k] == item[k] for k in ("bytes", "sha256")), "DELIVERY_CHANGED")
            return complete
        source = self.source()
        with closing(self.workspace()) as ws:
            self.bind(ws)
            imported = self.job(ws, "import", "response-import", {k: decision[k] for k in
                ("preview_id", "preview_sha256", "allow_incomplete")})
            version = imported["result"]["version"]
            prior = next((v for v in ws.state()["views"] if v["name"] == TABLE), None)
            if prior is None:
                prior = ws.database.activate(TABLE, version["version_id"], expected_revision=0)
            need(prior["version_id"] == version["version_id"], "ECONOMIC_VIEW_CHANGED")
            self.receipt("activation", {"view": TABLE, "version_id": version["version_id"],
                "revision": prior["revision"], "row_count": version["row_count"]})
            columns = source.receipt["columns"]
            quote = lambda name: '"' + name.replace('"', '""') + '"'
            select = ",".join(quote(k) for k in columns)
            ordering = '"record_kind","report_id","series_id",CASE WHEN "period_order"=\'\' THEN NULL ELSE CAST("period_order" AS INTEGER) END,"period","record_id"'
            all_query = self.save_query(ws, "EconomicRecords", "SELECT " + select + ' FROM "' + TABLE + '" ORDER BY ' + ordering)
            query = self.job(ws, "all_query", "query", {k: all_query[k] for k in ("sql", "parameters_text", "max_rows")} |
                             {"expected_views": ws.state()["views"]})
            exported = self.job(ws, "query_export", "query-export", {
                "query_id": query["id"], "result_sha256": query["result"]["result_sha256"],
                "label": source.manifest["title"], "formats": ["jsonl", "csv", "sqlite"],
                "timeout": 300, "spreadsheet_safe": True})
            query_files = {name: copy_delivery(ws.exports.download(exported["id"], name), self.deliveries / "query" / name)
                           for name in ("rows.jsonl", "rows.csv", "rows.sqlite3", "manifest.json")}
            typed = ws.workbooks.source({"kind": "query-export", "id": exported["id"]})
            need(typed["row_count"] == source.receipt["rowCount"], "COMPLETE_TYPED_ROWS_CHANGED")
            sections = []
            for name, kind, description in KIND_QUERIES:
                sql = "SELECT " + ",".join(quote(k) for k in SUMMARY_COLUMNS) + ' FROM "' + TABLE + '" WHERE "record_kind"='
                sql += "'" + kind + "' ORDER BY " + ordering
                saved = self.save_query(ws, "Economic" + name, sql)
                sections.append({"title": name, "note": description, "query_name": saved["name"], "query_revision": saved["revision"]})
            intro = ("Original DBnomics capture " + source.receipt["selection"]["captureId"] +
                     ". The source is described by its complete provider series and dimensions; frequencies are kept separate. " +
                     source.receipt["valueSemantics"] + "\n\n" + review_note)
            definition = ws.reports.save(REPORT, {"id": self.operation["report_request"], "expected_revision": 0,
                "title": source.manifest["title"], "introduction": intro, "sections": sections})
            self.receipt("report-definition", definition)
            metadata = ws.reports.definition(REPORT, revision=definition["revision"])
            report = self.job(ws, "report", "report", {"name": REPORT, "revision": definition["revision"],
                "definition_sha256": metadata["sha256"], "expected_views": ws.state()["views"]})
            publication = ws.reports.publish(report["id"], {"id": self.operation["report_publication"],
                "result_sha256": report["result"]["result_sha256"], "expected_revision": 0, "note": review_note})
            self.receipt("report-publication", publication)
            cursor = None
            while True:
                publications = ws.reports.publications(name=REPORT, before=cursor, limit=25)
                matched = next((item for item in publications["publications"] if item["publication"]["id"] == publication["id"]), None)
                if matched is not None:
                    publication_sha = matched["sha256"]
                    break
                cursor = publications["next_before"]
                need(cursor is not None, "ORIGINAL_PUBLICATION_ABSENT")
            briefing_files = {}
            for format in ("html", "json", "csv"):
                out = self.job(ws, "report_" + format, "report-export", {"publication_id": publication["id"],
                    "publication_sha256": publication_sha, "format": format})
                path, file = ws.report_export_file(out["id"])
                name = "economic-briefing." + ("zip" if format == "csv" else format)
                briefing_files[name] = copy_delivery(delivery(path, file).open(), self.deliveries / "briefing" / name)
            definition = ws.workbooks.save(WORKBOOK, {"id": self.operation["workbook_request"], "expected_revision": 0,
                "title": source.manifest["title"], "notes": source.receipt["valueSemantics"] + " " +
                    "Filter record_kind to distinguish original observations, calculations, series and analyst notes. " +
                    "Keep record_id when returning the Records sheet. Use review_note for an operator annotation. " +
                    "Original typed exports preserve dimensions and attributes that are omitted from this display.",
                "sheets": [{"name": "Records", "source": {"kind": "query-export", "id": exported["id"], "sha256": typed["sha256"]},
                    "columns": [{"source": k, "header": k, "type": "text", "null": "text", "null_text": "NULL",
                                 "precision": "preserve"} for k in source.receipt["workbookColumns"]]}]})
            self.receipt("workbook-definition", definition)
            workbook = self.job(ws, "workbook", "workbook", {"name": WORKBOOK, "revision": definition["revision"],
                "definition_sha256": definition["definition_sha256"]})
            workbook_files = {name: copy_delivery(ws.workbooks.download(workbook["id"], name), self.deliveries / "workbook" / name)
                              for name in ("workbook.xlsx", "handoff.zip", "manifest.json")}
            archive = self.archive(ws, "workspace_export")
            result = {"schema": WORKFLOW_SCHEMA, "operation_id": self.operation["id"], "workspace_id": ws.identity,
                "source": self.operation["source"], "document_id": source.document_id,
                "source_files": retained_document_files(ws.database.directory, source.document_id),
                "source_rows": source.receipt["rowCount"], "counts": source.receipt["counts"],
                "import": {"job_id": imported["id"], "version": version, "decision": decision},
                "query": {"id": query["id"], "result_sha256": query["result"]["result_sha256"],
                          "preview_truncated": query["result"]["truncated"]},
                "query_export": {"id": exported["id"], "source": typed, "files": query_files},
                "briefing": {"id": report["id"], "result_sha256": report["result"]["result_sha256"],
                    "publication": publication, "publication_sha256": publication_sha, "files": briefing_files},
                "workbook": {"id": workbook["id"], "definition": definition, "files": workbook_files},
                "workspace_archive": archive, "restore_operation": self.operation["restore_operation"]}
            return self.receipt("complete", result)

    def archive(self, ws, name, *, records=None, jobs=None):
        plan = ws.archives.review(include_previous_backups=False)
        need(plan["ready"], "PORTABLE_SOURCE_MISSING")
        job = self.job(ws, name, "workspace-export", {"include_previous_backups": False}, records=records, jobs=jobs)
        folder = self.deliveries / ("workspace-" + job["id"])
        files = {name: copy_delivery(ws.archives.download(job["id"], name), folder / name)
                 for name in ("workspace.zip", "archive-manifest.json")}
        return {"id": job["id"], "files": files}


    def original_workbook(self):
        complete = self.saved("complete")
        files = complete["workbook"]["files"]
        for record in files.values():
            actual = pin(record["path"])
            need(all(actual[k] == record[k] for k in ("bytes", "sha256")), "ORIGINAL_WORKBOOK_CHANGED")
        target = self.root / "original-workbook"
        if target.exists():
            workbook = WorkbookDelivery(target, expected_manifest_sha256=files["manifest.json"]["sha256"])
            workbook.verify_archive(files["handoff.zip"]["path"])
        else:
            workbook = WorkbookDelivery.receive(files["handoff.zip"]["path"], target,
                expected_sha256=files["handoff.zip"]["sha256"], max_bytes=MAX_BYTES)
        inspected = workbook.inspect()
        need(inspected["id"] == complete["workbook"]["id"] and
             inspected["definition_sha256"] == complete["workbook"]["definition"]["definition_sha256"],
             "ORIGINAL_WORKBOOK_IDENTITY_CHANGED")
        return workbook

    def return_directory(self, identifier):
        need(isinstance(identifier, str) and NAME.fullmatch(identifier), "RETURN_ID_REQUIRED")
        return private_directory(self.root / "returns" / identifier)

    def return_operation(self, identifier):
        folder = self.return_directory(identifier)
        operation = parse(read(folder / "operation.json", 1024 * 1024))
        need(operation.get("schema") == RETURN_SCHEMA and operation["id"] == identifier and
             operation["workflow_id"] == self.operation["id"], "RETURN_OPERATION_CHANGED")
        return folder, operation

    def return_workbook(self, identifier, file, *, expected_sha256):
        self.read_operation()
        source, workbook = self.source(), self.original_workbook()
        need(isinstance(expected_sha256, str) and HASH.fullmatch(expected_sha256), "WORKBOOK_RETURN_HASH_REQUIRED")
        folder = self.return_directory(identifier)
        binding = {"schema": RETURN_SCHEMA, "id": identifier, "workflow_id": self.operation["id"],
            "source_document_id": source.document_id, "source_selection": source.receipt["selection"],
            "original_workbook_id": workbook.manifest["id"], "original_manifest_sha256": sha(workbook.raw)}
        path = folder / "operation.json"
        if path.exists():
            operation = parse(read(path, 1024 * 1024))
            need(all(operation.get(k) == v for k, v in binding.items()), "RETURN_OPERATION_CHANGED")
            selected = operation["returned_file"]
            need(selected["sha256"] == expected_sha256, "RETURN_OPERATION_CHANGED")
        else:
            selected = pin(file, 128 * 1024 * 1024)
            need(selected["sha256"] == expected_sha256, "WORKBOOK_RETURN_CHANGED")
            names = ("baseline_preview", "baseline_import", "returned_preview", "returned_import",
                     "comparison", "preview_changes", "changes_json", "changes_csv", "workspace_export",
                     "notes_preview", "notes_import", "notes_workspace_export")
            operation = {**binding, "returned_file": selected, "created_at": datetime.now(timezone.utc).isoformat(),
                "view": "workbook_" + sha(identifier.encode())[:20],
                "jobs": {name: uuid.uuid4().hex for name in names},
                "review_id": uuid.uuid4().hex, "restore_operation": uuid.uuid4().hex,
                "notes_restore_operation": uuid.uuid4().hex,
                "notes_source_operation": "workbook-notes-source-" + uuid.uuid4().hex,
                "notes_operation": "workbook-notes-" + uuid.uuid4().hex,
                "source_receive_operation": "workbook-source-" + uuid.uuid4().hex}
            dump(path, operation, exclusive=True)
        records = private_directory(folder / "records")
        retained_path = folder / "returned.xlsx"
        if retained_path.exists():
            retained = pin(retained_path, 128 * 1024 * 1024)
            need(all(retained[k] == selected[k] for k in ("bytes", "sha256")), "WORKBOOK_RETURN_CHANGED")
        else:
            original = pin(file, 128 * 1024 * 1024)
            need(all(original[k] == selected[k] for k in ("bytes", "sha256")), "WORKBOOK_RETURN_CHANGED")
            retained = copy_delivery(delivery(file, original).open(), retained_path)
        if (records / "comparison.json").exists():
            old = parse(read(records / "comparison.json", MAX_BYTES))
            need(old["returned_file"]["sha256"] == retained["sha256"], "RETURN_OPERATION_CHANGED")
            return self.return_review(identifier)
        projections = {}
        baseline = workbook.directory / "workbook.xlsx"
        for side, source_file, checksum in (
            ("baseline", baseline, workbook.manifest["files"]["workbook.xlsx"]["sha256"]),
            ("returned", Path(retained["path"]), retained["sha256"]),
        ):
            projections[side] = workbook_projection(folder / side, source_file, expected_sha256=checksum,
                workbook=workbook, operation=operation, source_document_id=source.document_id,
                workbook_columns=source.receipt["workbookColumns"], side=side)
            keep(records / ("projection-" + side + ".json"), projections[side])
        with closing(self.workspace()) as ws:
            self.bind(ws)
            versions = {}
            for side in ("baseline", "returned"):
                projection = projections[side]
                name = operation["view"] + "_" + side
                registered = ws.responses.register(name, "document", manifest=projection["manifest"],
                    format="json", expected_revision=0)
                keep(records / ("registration-" + side + ".json"), registered)
                preview = self.job(ws, side + "_preview", "response-preview", {"source": name,
                    "source_revision": registered["revision"], "mapping": projection["mapping"]},
                    records=records, jobs=operation["jobs"])
                imported = self.job(ws, side + "_import", "response-import", {
                    "preview_id": preview["id"], "preview_sha256": preview["result"]["preview_sha256"],
                    "allow_incomplete": False}, records=records, jobs=operation["jobs"])
                versions[side] = imported["result"]["version"]
                need(versions[side]["row_count"] == projection["rows"], "COMPLETE_WORKBOOK_ROWS_CHANGED")
            current = next((x for x in ws.state()["views"] if x["name"] == operation["view"]), None)
            if current is None:
                current = ws.database.activate(operation["view"], versions["baseline"]["version_id"], expected_revision=0)
            need(current["version_id"] == versions["baseline"]["version_id"], "RETURN_REVIEW_VIEW_CHANGED")
            comparison = self.job(ws, "comparison", "compare", {"view": operation["view"],
                "expected_revision": current["revision"], "base_version_id": versions["baseline"]["version_id"],
                "candidate_version_id": versions["returned"]["version_id"], "keys": ["record_id"]},
                records=records, jobs=operation["jobs"])
            report = comparison["result"]
            exported = self.job(ws, "preview_changes", "change-export", {"comparison_id": comparison["id"],
                "result_sha256": report["result_sha256"], "format": "json", "selection_sha256": sha(canonical(None))},
                records=records, jobs=operation["jobs"])
            export_path, metadata = ws.change_export_file(exported["id"])
            full_changes = copy_delivery(delivery(export_path, metadata).open(), folder / "review" / "changes.json")
            result = {"schema": RETURN_SCHEMA, "return_id": identifier, "workflow_id": self.operation["id"],
                "workspace_id": ws.identity, "original_workbook_id": workbook.manifest["id"],
                "source_document_id": source.document_id, "returned_file": retained,
                "projections": {side: {"document_id": value["document_id"], "rows": value["rows"],
                    "worksheets": value["receipt"]["worksheets"], "other_worksheets": value["receipt"]["other_worksheets"]}
                    for side, value in projections.items()},
                "versions": versions, "review_view": current, "comparison_id": comparison["id"],
                "result_sha256": report["result_sha256"], "report": report, "changes_file": full_changes,
                "review_id": operation["review_id"], "source_values_view": TABLE,
                "semantics": projections["returned"]["receipt"]["semantics"]}
            return keep(records / "comparison.json", result)

    def return_review(self, identifier, *, after=0, limit=100):
        self.read_operation()
        folder, operation = self.return_operation(identifier)
        records = folder / "records"
        result = parse(read(records / "comparison.json", MAX_BYTES))
        actual = pin(result["returned_file"]["path"])
        need(all(actual[k] == result["returned_file"][k] for k in ("bytes", "sha256")), "WORKBOOK_RETURN_CHANGED")
        with closing(self.workspace()) as ws:
            self.bind(ws)
            page = ws.changes.page(result["comparison_id"], after=after, limit=limit)
            need(page["report"]["result_sha256"] == result["result_sha256"], "WORKBOOK_COMPARISON_CHANGED")
            return {**result, "page": page}

    def accept_return(self, identifier, *, result_sha256, note):
        self.read_operation()
        folder, operation = self.return_operation(identifier)
        records = folder / "records"
        comparison = parse(read(records / "comparison.json", MAX_BYTES))
        need(result_sha256 == comparison["result_sha256"] and isinstance(note, str) and
             1 <= len(note.encode()) <= 8000, "EXACT_RETURN_REVIEW_REQUIRED")
        review = keep(records / "review.json", {"id": operation["review_id"],
            "result_sha256": result_sha256, "note": note})
        with closing(self.workspace()) as ws:
            self.bind(ws)
            selection = ws.changes.select(comparison["comparison_id"], review)
            keep(records / "selection.json", selection)
            activation = self.saved("activation")
            economic = next(x for x in ws.state()["views"] if x["name"] == TABLE)
            need(all(economic[k] == activation[k] for k in ("version_id", "revision")), "ECONOMIC_VIEW_CHANGED")
            exports = {}
            for format in ("json", "csv"):
                exported = self.job(ws, "changes_" + format, "change-export", {
                    "comparison_id": comparison["comparison_id"], "result_sha256": result_sha256,
                    "selection_sha256": sha(canonical(selection)), "format": format},
                    records=records, jobs=operation["jobs"])
                path, record = ws.change_export_file(exported["id"])
                exports[format] = copy_delivery(delivery(path, record).open(), folder / "accepted" / ("changes." + format))
            archive = self.archive(ws, "workspace_export", records=records, jobs=operation["jobs"])
            result = {"schema": RETURN_SCHEMA, "return_id": identifier, "workflow_id": self.operation["id"],
                "workspace_id": ws.identity, "comparison_id": comparison["comparison_id"],
                "result_sha256": result_sha256, "selection": selection, "selection_sha256": sha(canonical(selection)),
                "source_values_view": economic, "changes": exports, "workspace_archive": archive,
                "restore_operation": operation["restore_operation"]}
            return keep(records / "accepted.json", result)

    def adopt_notes(self, identifier, project, *, author):
        self.read_operation()
        need(self.checkout and isinstance(author, str) and 1 <= len(author) <= 160, "SOURCE_OPERATOR_REQUIRED")
        folder, operation = self.return_operation(identifier)
        records = folder / "records"
        accepted = parse(read(records / "accepted.json", MAX_BYTES))
        source = self.source()
        with closing(self.workspace()) as ws:
            self.bind(ws)
            changed, cursor = [], 0
            while True:
                page = ws.changes.page(accepted["comparison_id"], after=cursor, limit=100)
                need(page["report"]["result_sha256"] == accepted["result_sha256"] and
                     page["selection_sha256"] == accepted["selection_sha256"], "ACCEPTED_RETURN_CHANGED")
                changed.extend(page["rows"])
                cursor = page["next_after"]
                if cursor is None: break
            columns = page["report"]["columns"]
        indices = {name: i for i, name in enumerate(columns)}
        additions = []
        for row in changed:
            if row["kind"] == "moved":
                continue
            need(row["kind"] == "changed" and set(row["changed_columns"]) <= {"review_note", "excel_cell_types"},
                 "NOTES_REQUIRE_UNCHANGED_SOURCE_CELLS")
            before_types = json.loads(row["before"][indices["excel_cell_types"]]["text"])
            after_types = json.loads(row["after"][indices["excel_cell_types"]]["text"])
            need(all(before_types.get(k) == after_types.get(k) for k in
                     set(before_types) | set(after_types) if k != "review_note"), "NOTES_REQUIRE_UNCHANGED_SOURCE_TYPES")
            cell = row["after"][indices["review_note"]]
            if cell["type"] == "null" or cell["type"] == "string" and not cell["text"].strip():
                continue
            need(cell["type"] == "string" and after_types.get("review_note") == "string" and
                 len(cell["text"]) <= 16000, "PLAIN_ANALYST_NOTE_REQUIRED")
            identity = row["after"][indices["record_id"]]["text"]
            observation = row["after"][indices["observation_id"]]
            observation_id = observation["text"] if observation["type"] == "string" and observation["text"] else None
            additions.append({"id": "wb-" + sha(canonical({"return": identifier, "record": identity, "text": cell["text"]}))[:48],
                "observationId": observation_id, "author": author, "tag": "workbook-review", "text": cell["text"]})
        need(additions, "RETURN_HAS_NO_NEW_ANALYST_NOTES")
        handover = source.original(folder / "original-source-handover.json")
        bundle = parse(read(Path(handover["path"]), MAX_BYTES))
        selected = source.receipt["selection"]
        matching = [x for x in bundle["notes"] if x["captureId"] == selected["captureId"] and
                    x["revision"] == selected["notesRevision"]]
        need(len(matching) == (0 if selected["notesRevision"] == 0 else 1), "ORIGINAL_NOTES_EDITION_CHANGED")
        entries = deepcopy(matching[0]["entries"]) if matching else []
        existing = {entry["id"] for entry in entries}
        need(not any(entry["id"] in existing for entry in additions), "WORKBOOK_NOTE_ID_ALREADY_PRESENT")
        project = str(Path(project).expanduser().absolute())
        request = {"id": operation["notes_operation"], "captureId": selected["captureId"],
            "expectedRevision": selected["notesRevision"], "entries": [*entries, *additions]}
        keep(records / "notes-adoption-binding.json", {"project": project, "project_id": source.receipt["identity"]["project"]["id"],
            "return_id": identifier, "selection_sha256": accepted["selection_sha256"], "author": author,
            "returned_sha256": operation["returned_file"]["sha256"]})
        request_path = records / "notes-request.json"
        keep(request_path, request)
        status = native(self.checkout, self.node, ["project-status", "--workspace", project])
        need(status["project"]["id"] == source.receipt["identity"]["project"]["id"], "ORIGINAL_ECONOMIC_PROJECT_REQUIRED")
        result = native(self.checkout, self.node, ["save-notes", "--workspace", project, "--input", str(request_path)])
        need(result["captureId"] == selected["captureId"] and result["captureSha256"] == selected["captureSha256"] and
             result["operationId"] == operation["notes_operation"] and
             result["revision"] == selected["notesRevision"] + 1, "ORIGINAL_NOTES_RESULT_CHANGED")
        # Retain the actual new native notes edition in the same Analysis archive.
        # Frozen initial reports and their original note revisions remain intact.
        updated_selection = {**selected, "notesRevision": result["revision"]}
        adopted_source = DbnomicsAnalysisSource.create(project, folder / "adopted-economic-source",
            updated_selection, operation_id=operation["notes_source_operation"], checkout=self.checkout, node=self.node)
        with closing(self.workspace()) as ws:
            self.bind(ws)
            name = operation["view"] + "_notes"
            registered = adopted_source.register(ws, name)
            keep(records / "notes-source-registration.json", registered)
            preview = self.job(ws, "notes_preview", "response-preview", {"source": name,
                "source_revision": registered["revision"], "mapping": adopted_source.mapping},
                records=records, jobs=operation["jobs"])
            imported = self.job(ws, "notes_import", "response-import", {"preview_id": preview["id"],
                "preview_sha256": preview["result"]["preview_sha256"], "allow_incomplete": False},
                records=records, jobs=operation["jobs"])
            version = imported["result"]["version"]
            current = next((x for x in ws.state()["views"] if x["name"] == name), None)
            if current is None:
                current = ws.database.activate(name, version["version_id"], expected_revision=0)
            need(current["version_id"] == version["version_id"], "ADOPTED_NOTES_VIEW_CHANGED")
            archive = self.archive(ws, "notes_workspace_export", records=records, jobs=operation["jobs"])
        return keep(records / "notes-adopted.json", {"project": project, "project_id": status["project"]["id"],
            "return_id": identifier, "notes_added": len(additions), "result": result,
            "analysis_source": {"document_id": adopted_source.document_id, "selection": updated_selection,
                "source_identity": adopted_source.receipt["identity"], "row_count": adopted_source.receipt["rowCount"],
                "view": current}, "workspace_archive": archive, "restore_operation": operation["notes_restore_operation"]})

    def reopen(self, *, identifier=None, target=None, source_target=None):
        self.read_operation()
        complete = self.saved("complete")
        adopted, receiving_name = None, "reopened"
        receiving_input = "receiving-input.json"
        if identifier is None:
            result, records, restore_id = complete, self.records, self.operation["restore_operation"]
            default = self.root / "received"
            receive_id = "analysis-source-" + self.operation["id"]
        else:
            folder, operation = self.return_operation(identifier)
            records = folder / "records"
            result = parse(read(records / "accepted.json", MAX_BYTES))
            restore_id = operation["restore_operation"]
            default, receive_id = folder / "received", operation["source_receive_operation"]
            if (records / "notes-adopted.json").exists():
                adopted = parse(read(records / "notes-adopted.json", MAX_BYTES))
                result = {**result, "workspace_archive": adopted["workspace_archive"]}
                restore_id = adopted["restore_operation"]
                default = folder / "received-with-notes"
                receive_id = operation["notes_source_operation"] + "-receive"
                receiving_name, receiving_input = "reopened-with-notes", "receiving-with-notes-input.json"
        target = str(Path(target or default).expanduser().absolute())
        selected = result["workspace_archive"]["files"]["workspace.zip"]
        keep(records / receiving_input, {"target": target, "operation": restore_id,
            "archive_sha256": selected["sha256"], "workspace_id": complete["workspace_id"]})
        restored = restore_archive(selected["path"], target, operation=restore_id, expected_sha256=selected["sha256"])
        with closing(AnalysisWorkspace(target)) as ws:
            need(ws.identity == complete["workspace_id"], "RECEIVED_WORKSPACE_ID_CHANGED")
            original = DbnomicsAnalysisSource(document_path(ws.database.directory, complete["document_id"]).parent,
                                              expected_document_id=complete["document_id"])
            source_check = original.verify()
            state = ws.state()
            economic = next(x for x in state["views"] if x["name"] == TABLE)
            need(economic["version_id"] == complete["import"]["version"]["version_id"], "RECEIVED_ECONOMIC_VIEW_CHANGED")
            files = {}
            for name, expected in complete["query_export"]["files"].items():
                files["query/" + name] = checked_download(ws.exports.download(complete["query_export"]["id"], name), expected)
            for name, expected in complete["workbook"]["files"].items():
                files["workbook/" + name] = checked_download(ws.workbooks.download(complete["workbook"]["id"], name), expected)
            for name, expected in complete["briefing"]["files"].items():
                format = "csv" if name.endswith(".zip") else name.rsplit(".", 1)[1]
                job = self.saved("job-report_" + format)
                path, record = ws.report_export_file(job["id"])
                files["briefing/" + name] = checked_download(delivery(path, record).open(), expected)
            reviewed = None
            if identifier is not None:
                page = ws.changes.page(result["comparison_id"], limit=1)
                need(page["selection_sha256"] == result["selection_sha256"] and
                     page["report"]["result_sha256"] == result["result_sha256"], "RECEIVED_REVIEW_CHANGED")
                reviewed = {"comparison_id": result["comparison_id"], "selection": page["selection"],
                            "counts": page["report"]["counts"], "view": page["current_view"]}
            notes_reopened = None
            if adopted is not None:
                updated = adopted["analysis_source"]
                original = DbnomicsAnalysisSource(document_path(ws.database.directory, updated["document_id"]).parent,
                                                  expected_document_id=updated["document_id"])
                notes_reopened = {"source": original.verify(), "selection": original.receipt["selection"],
                    "view": next(x for x in state["views"] if x["name"] == updated["view"]["name"])}
                need(notes_reopened["view"]["version_id"] == updated["view"]["version_id"], "RECEIVED_NOTES_VIEW_CHANGED")
            handover = original.original(Path(target).parent / (Path(target).name + "-economic-source.json"))
            source_identity = original.receipt["identity"]
            source_selection = original.receipt["selection"]
        received_source = None
        if source_target is not None:
            need(self.checkout is not None, "ORIGINAL_NODE_RUNTIME_REQUIRED")
            source_target = str(Path(source_target).expanduser().absolute())
            keep(records / (receiving_name + "-source-input.json"), {"target": source_target,
                "operation_id": receive_id, "source_file_sha256": handover["sha256"],
                "original_project": source_identity["project"]})
            native(self.checkout, self.node, ["init", "--workspace", source_target])
            imported = native(self.checkout, self.node, ["import-sources", "--workspace", source_target,
                "--id", receive_id, "--file", handover["path"]])
            status = native(self.checkout, self.node, ["project-status", "--workspace", source_target])
            capture = next((x for x in status["captures"] if x["captureId"] == source_identity["captureId"]), None)
            need(capture and capture["captureSha256"] == source_identity["captureSha256"], "RECEIVED_SOURCE_CAPTURE_CHANGED")
            received_notes = native(self.checkout, self.node, ["notes", "--workspace", source_target,
                "--id", source_selection["captureId"], "--revision", str(source_selection["notesRevision"])])
            need(received_notes["revision"] == source_selection["notesRevision"], "RECEIVED_NOTES_REVISION_CHANGED")
            received_source = {"project": status["project"], "captures": status["captures"], "reports": status["reports"],
                               "notes": received_notes, "import": imported, "directory": source_target}
        reopened = {"schema": WORKFLOW_SCHEMA, "operation_id": self.operation["id"], "return_id": identifier,
            "restore": restored, "workspace_id": complete["workspace_id"], "source_check": source_check,
            "source_files": files, "economic_view": economic, "review": reviewed,
            "original_handover": handover, "economic_receiving": received_source, "notes_reopened": notes_reopened}
        # A later original source-project import is a completed stage of the
        # same archive receiving operation; retain its own complete result.
        record_name = receiving_name + ("-with-economic-source" if source_target is not None else "")
        return keep(records / (record_name + ".json"), reopened,
                    identity=("operation_id", "return_id", "workspace_id", "source_check", "source_files", "economic_view", "review", "notes_reopened"))

    def inspect(self):
        self.read_operation()
        result = {"operation": self.operation, "source": self.saved("source"),
                  "review": self.saved("review"), "returns": []}
        if (self.records / "complete.json").exists(): result["complete"] = self.saved("complete")
        folder = self.root / "returns"
        if folder.exists():
            for path in sorted(folder.iterdir()):
                if path.is_dir() and (path / "operation.json").is_file():
                    item = parse(read(path / "operation.json", 1024 * 1024))
                    for name in ("comparison", "accepted", "notes-adopted", "reopened", "reopened-with-notes",
                                 "reopened-with-economic-source", "reopened-with-notes-with-economic-source"):
                        record = path / "records" / (name + ".json")
                        if record.exists(): item[name] = parse(read(record, MAX_BYTES))
                    result["returns"].append(item)
        return result


def main(argv=None):
    parser = argparse.ArgumentParser(description="Retain DBnomics economic sources, build native Analysis briefings and review complete returned workbooks.", allow_abbrev=False)
    parser.add_argument("--operation", required=True, help="Original private economic-analysis workflow directory")
    parser.add_argument("--checkout", help="Original complete merchant repository/source root for native DBnomics calls")
    parser.add_argument("--node", default="node")
    parser.add_argument("--wait-seconds", type=int, default=900)
    commands = parser.add_subparsers(dest="command", required=True)
    def source_options(command):
        sources = command.add_mutually_exclusive_group(required=True)
        sources.add_argument("--source", help="Complete Analysis source directory containing source.json")
        sources.add_argument("--bundle", help="Complete original analysis-source-export JSON")
        sources.add_argument("--project", help="Existing native DBnomics project")
        command.add_argument("--sha256", help="Exact complete bundle SHA256")
        command.add_argument("--selection", help="JSON analysis selection for the existing DBnomics project")
        command.add_argument("--source-operation", help="Original stable native analysis-source operation ID")
    prepare = commands.add_parser("prepare", allow_abbrev=False)
    source_options(prepare)
    receive = commands.add_parser("receive", allow_abbrev=False)
    source_options(receive)
    receive.add_argument("--review-note", required=True)
    receive.add_argument("--source-receiving-project")
    finish = commands.add_parser("finish", allow_abbrev=False)
    finish.add_argument("--preview-sha256", required=True)
    finish.add_argument("--review-note", required=True)
    returned = commands.add_parser("return-workbook", allow_abbrev=False)
    returned.add_argument("--id", required=True)
    returned.add_argument("--file", required=True)
    returned.add_argument("--sha256", required=True)
    review = commands.add_parser("return-review", allow_abbrev=False)
    review.add_argument("--id", required=True)
    review.add_argument("--after", type=int, default=0)
    review.add_argument("--limit", type=int, default=100)
    accept = commands.add_parser("accept-return", allow_abbrev=False)
    accept.add_argument("--id", required=True)
    accept.add_argument("--result-sha256", required=True)
    accept.add_argument("--review-note", required=True)
    notes = commands.add_parser("adopt-notes", allow_abbrev=False)
    notes.add_argument("--id", required=True)
    notes.add_argument("--project", required=True)
    notes.add_argument("--author", required=True)
    reopened = commands.add_parser("reopen", allow_abbrev=False)
    reopened.add_argument("--id")
    reopened.add_argument("--target")
    reopened.add_argument("--source-project")
    commands.add_parser("inspect", allow_abbrev=False)
    server = commands.add_parser("serve", allow_abbrev=False)
    server.add_argument("--port", type=int, default=0)
    args = parser.parse_args(argv)
    workflow = DbnomicsEconomicWorkflow(args.operation, checkout=args.checkout, node=args.node,
                                        wait_seconds=args.wait_seconds)
    if args.command in {"prepare", "receive"}:
        if args.source:
            source = DbnomicsAnalysisSource(args.source)
        elif args.bundle:
            need(args.sha256 is not None, "EXACT_ANALYSIS_HANDOVER_SHA_REQUIRED")
            source = DbnomicsAnalysisSource.from_bundle(args.bundle, workflow.root / "source-received",
                                                       expected_sha256=args.sha256)
        else:
            need(args.checkout and args.selection and args.source_operation, "NATIVE_SOURCE_SELECTION_REQUIRED")
            selection = parse(read(Path(args.selection).expanduser().absolute(), 1024 * 1024))
            source = DbnomicsAnalysisSource.create(args.project, workflow.root / "source-selected", selection,
                operation_id=args.source_operation, checkout=args.checkout, node=args.node)
        result = workflow.prepare(source)
        if args.command == "receive":
            completed = workflow.finish(preview_sha256=result["preview_sha256"], review_note=args.review_note)
            reopened = workflow.reopen(source_target=args.source_receiving_project)
            result = {"review": result, "complete": completed, "reopened": reopened}
    elif args.command == "finish":
        result = workflow.finish(preview_sha256=args.preview_sha256, review_note=args.review_note)
    elif args.command == "return-workbook":
        result = workflow.return_workbook(args.id, args.file, expected_sha256=args.sha256)
    elif args.command == "return-review":
        result = workflow.return_review(args.id, after=args.after, limit=args.limit)
    elif args.command == "accept-return":
        result = workflow.accept_return(args.id, result_sha256=args.result_sha256, note=args.review_note)
    elif args.command == "adopt-notes":
        result = workflow.adopt_notes(args.id, args.project, author=args.author)
    elif args.command == "reopen":
        result = workflow.reopen(identifier=args.id, target=args.target, source_target=args.source_project)
    elif args.command == "inspect":
        result = workflow.inspect()
    else:
        from .analysis_workspace import serve
        workflow.read_operation()
        with closing(workflow.workspace()) as workspace:
            workflow.bind(workspace)
            server, control, url = serve(workspace, port=args.port)
            print(json.dumps({"status": "LISTENING", "address": url.split("#", 1)[0],
                              "private_launch": str(control), "workspace_id": workspace.identity}), flush=True)
            try:
                server.serve_forever()
            except KeyboardInterrupt:
                pass
            finally:
                server.server_close()
        return
    print(canonical(result).decode("utf-8"), flush=True)


if __name__ == "__main__":
    main()
