"""Saved Excel deliveries from original complete customer query/report results.

The existing AnalysisWorkspace owns the database lease, executor and job log.
This module reads retained results; it never executes SQL or an acquisition.
"""
from __future__ import annotations

import argparse
import base64
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import stat
import tempfile
import threading
import time
import zipfile

from .analysis_reports import AnalysisReports
from .errors import BuyerError, StateError
from .query_export import QueryExport
from .storage import digest, lock, parse, private_directory, read, sync_directory, write

_SCHEMA = "stellar-forge.customer-workbook.v1"
_DEFINITION = "stellar-forge.customer-workbook-definition.v1"
_NAME = re.compile(r"[A-Za-z][A-Za-z0-9_-]{0,63}")
_ID = re.compile(r"[a-f0-9]{32}")
_HASH = re.compile(r"[a-f0-9]{64}")
_BAD_SHEET = re.compile(r"[\\/*?:\[\]]")
_TYPES = {"native", "text", "integer", "number", "decimal-text", "boolean", "date", "datetime-utc", "datetime-local", "binary-base64"}
_ROWS_PER_SHEET = 1_048_575  # One additional row is reserved for column headings.
_MAX_COLUMNS = 16_384
_MAX_MANIFEST = 4 * 1024 * 1024
_MIMES = {"workbook.xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          "handoff.zip": "application/zip", "manifest.json": "application/json"}


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def _now():
    return datetime.now(timezone.utc).isoformat()


def _error(code):
    raise StateError("ANALYSIS_WORKBOOK_" + code)


def _name(value):
    if not isinstance(value, str) or not _NAME.fullmatch(value):
        _error("NAME_INVALID")
    return value


def _id(value):
    if not isinstance(value, str) or not _ID.fullmatch(value):
        _error("REQUEST_ID_REQUIRED")
    return value


def _hash(value):
    if not isinstance(value, str) or not _HASH.fullmatch(value):
        _error("DIGEST_REQUIRED")
    return value


def _text(value, maximum, *, required=False):
    if not isinstance(value, str) or len(value.encode("utf-8")) > maximum or "\0" in value or required and not value.strip():
        _error("TEXT_INVALID")
    return value


def _unpack(row):
    if row is None:
        _error("NOT_FOUND")
    raw = row["record"].encode()
    if digest(raw) != row["sha256"]:
        _error("RETAINED_RECORD_CHANGED")
    return parse(raw)


def _page(before, limit):
    if type(limit) is not int or not 1 <= limit <= 100 or before is not None and (type(before) is not int or before < 1):
        _error("PAGE_INVALID")


def _source_key(source, column):
    """Select report cells by position while preserving original column labels."""
    if not isinstance(column, dict) or column.get("source") not in source["columns"]:
        _error("SOURCE_COLUMN_CHANGED")
    if source["kind"] != "report":
        return column["source"]
    index = column.get("source_index")
    if index is None:
        matches = [index for index, name in enumerate(source["columns"]) if name == column["source"]]
        if len(matches) != 1:
            _error("SOURCE_COLUMN_INDEX_REQUIRED")
        return matches[0]
    if (type(index) is not int or not 0 <= index < len(source["columns"])
            or source["columns"][index] != column["source"]):
        _error("SOURCE_COLUMN_CHANGED")
    return index


def _relative(value):
    if not isinstance(value, str) or "\\" in value or "\0" in value:
        _error("FILE_INVALID")
    path = PurePosixPath(value)
    if path.is_absolute() or not path.parts or any(part in {".", "..", ""} for part in path.parts) or str(path) != value:
        _error("FILE_INVALID")
    if value in {"workbook.xlsx", "definition.json"}:
        return value
    if len(path.parts) == 2 and path.parts[0] == "history" and re.fullmatch(r"definition-[1-9][0-9]*\.json", path.parts[1]):
        return value
    if (len(path.parts) == 3 and path.parts[0] == "sources"
            and re.fullmatch(r"(?:query-export|report)-[a-f0-9]{32}", path.parts[1])
            and path.parts[2] in {"manifest.json", "rows.jsonl", "rows.csv", "rows.sqlite3", "report.json"}):
        return value
    _error("FILE_INVALID")


@contextmanager
def _open_private(path):
    fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
    with os.fdopen(fd, "rb") as stream:
        info = os.fstat(stream.fileno())
        if (not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_mode & 0o077
                or hasattr(os, "getuid") and info.st_uid != os.getuid()):
            _error("PRIVATE_FILE_REQUIRED")
        yield stream


def _file_record(path):
    with _open_private(path) as stream:
        return {"bytes": os.fstat(stream.fileno()).st_size, "sha256": hashlib.file_digest(stream, "sha256").hexdigest()}



def _publish_received_directory(source, target):
    """Atomically publish a complete directory without replacing any destination."""
    import ctypes
    import sys

    libc = ctypes.CDLL(None, use_errno=True)
    if sys.platform.startswith("linux") and hasattr(libc, "renameat2"):
        rename = libc.renameat2
        rename.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p, ctypes.c_uint]
        rename.restype = ctypes.c_int
        result = rename(-100, os.fsencode(source), -100, os.fsencode(target), 1)  # RENAME_NOREPLACE
    elif sys.platform == "darwin" and hasattr(libc, "renamex_np"):
        rename = libc.renamex_np
        rename.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
        rename.restype = ctypes.c_int
        result = rename(os.fsencode(source), os.fsencode(target), 4)  # RENAME_EXCL
    else:
        _error("ATOMIC_DIRECTORY_PUBLICATION_UNAVAILABLE")
    if result:
        code = ctypes.get_errno()
        raise OSError(code, os.strerror(code), os.fspath(target))
    sync_directory(target.parent)


def _literal(value):
    """Check before openpyxl's string binding can truncate an Excel cell."""
    try:
        length = len(value.encode("utf-16-le")) // 2
    except UnicodeError:
        _error("CELL_UNICODE_INVALID")
    if length > 32_767:
        _error("CELL_TEXT_EXCEEDS_EXCEL_LIMIT")
    if any(ord(c) < 32 and c not in "\t\n\r" or ord(c) in {0xFFFE, 0xFFFF} for c in value):
        _error("CELL_XML_CHARACTER_INVALID")
    return value


def _display(value):
    if isinstance(value, bytes):
        return "base64:" + base64.b64encode(value).decode("ascii")
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and not math.isfinite(value):
        return "NaN" if math.isnan(value) else "Infinity" if value > 0 else "-Infinity"
    return str(value)


def project_cell(value, column):
    """Return (value, Excel number format, conversion), with explicit precision."""
    kind = column["type"]
    if value is None:
        return (None if column["null"] == "blank" else _literal(column["null_text"]), None, "null")
    if kind == "text":
        return _literal(_display(value)), "@", "empty-string" if value == "" else None
    if kind == "binary-base64":
        if not isinstance(value, bytes):
            _error("BINARY_CELL_REQUIRED")
        return _literal(base64.b64encode(value).decode("ascii")), "@", "binary-text"
    if kind == "decimal-text":
        if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
            _error("DECIMAL_CELL_REQUIRED")
        text = str(value)
        try:
            number = Decimal(text)
        except InvalidOperation:
            _error("DECIMAL_CELL_REQUIRED")
        if not number.is_finite():
            _error("FINITE_DECIMAL_REQUIRED")
        return _literal(text), "@", "decimal-text"
    if kind == "boolean":
        if type(value) is bool:
            return value, None, None
        if type(value) is int and value in {0, 1}:
            return bool(value), None, "boolean"
        if isinstance(value, str) and value.casefold() in {"true", "false"}:
            return value.casefold() == "true", None, "boolean"
        _error("BOOLEAN_CELL_REQUIRED")
    if kind in {"date", "datetime-utc", "datetime-local"}:
        if not isinstance(value, str):
            _error("ISO_DATE_TEXT_REQUIRED")
        try:
            if kind == "date":
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                    _error("ISO_DATE_TEXT_REQUIRED")
                parsed = date.fromisoformat(value)
                if parsed < date(1900, 1, 1):
                    return _literal(value), "@", "date-text"
                return parsed, "yyyy-mm-dd", "date"
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            if "T" not in value and " " not in value:
                _error("ISO_DATETIME_TEXT_REQUIRED")
            if kind == "datetime-utc":
                if parsed.tzinfo is None:
                    _error("DATETIME_TIMEZONE_REQUIRED")
                parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
            elif parsed.tzinfo is not None:
                _error("LOCAL_DATETIME_MUST_BE_NAIVE")
            if parsed.year < 1900 or parsed.microsecond % 1000:
                return _literal(value), "@", "datetime-text"
            return parsed, 'yyyy-mm-dd hh:mm:ss.000" UTC"' if kind == "datetime-utc" else "yyyy-mm-dd hh:mm:ss.000", "datetime"
        except (ValueError, OverflowError):
            _error("ISO_DATE_TEXT_REQUIRED")
    if kind == "integer":
        if type(value) is int:
            pass
        elif isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value):
            value = int(value)
        else:
            _error("INTEGER_CELL_REQUIRED")
    elif kind == "number":
        if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
            _error("NUMERIC_CELL_REQUIRED")
        text = str(value)
        try:
            number = Decimal(text)
        except InvalidOperation:
            _error("NUMERIC_CELL_REQUIRED")
        if not number.is_finite():
            _error("FINITE_NUMBER_REQUIRED")
        if column["precision"] == "preserve" and (len(number.as_tuple().digits) > 15 or number.adjusted() > 307 or number.adjusted() < -307):
            return _literal(text), "@", "precision-text"
        value = float(number)
        if not math.isfinite(value) or value == 0 and number != 0:
            _error("NUMBER_OUTSIDE_EXCEL_RANGE")
    if isinstance(value, bytes):
        return _literal(_display(value)), "@", "binary-text"
    if isinstance(value, str):
        return _literal(value), "@", "empty-string" if value == "" else None
    if isinstance(value, bool):
        return value, None, None
    if type(value) is int:
        if len(str(abs(value))) > 15 and column["precision"] == "preserve":
            return str(value), "@", "precision-text"
        return value, "0", "numeric-choice" if len(str(abs(value))) > 15 else None
    if isinstance(value, float):
        if not math.isfinite(value):
            return _display(value), "@", "nonfinite-text"
        decimal = Decimal(repr(value))
        if column["precision"] == "preserve" and (len(decimal.as_tuple().digits) > 15 or decimal.adjusted() > 307 or decimal.adjusted() < -307):
            return repr(value), "@", "precision-text"
        return value, "General", "numeric-choice" if column["precision"] == "excel-number" else None
    _error("CELL_TYPE_UNSUPPORTED")


def _sheet_layout(sheets):
    used = {"readme", "sources"}
    result = []
    for sheet in sheets:
        names = []
        count = max(1, (sheet["source"]["row_count"] + _ROWS_PER_SHEET - 1) // _ROWS_PER_SHEET)
        for index in range(count):
            suffix = "" if index == 0 else f" ({index + 1})"
            name = sheet["name"][:31 - len(suffix)] + suffix
            if name.casefold() in used:
                _error("SHEET_NAMES_COLLIDE_AFTER_SPLIT")
            used.add(name.casefold())
            names.append(name)
        result.append(names)
    return result


class WorkbookDelivery:
    """Open, inspect, copy or receive a complete workbook without a live workspace."""

    def __init__(self, directory, *, expected_manifest_sha256=None):
        self.directory = Path(directory).absolute()
        if self.directory.resolve() != self.directory or not self.directory.is_dir():
            _error("CANONICAL_DELIVERY_DIRECTORY_REQUIRED")
        self.raw = read(self.directory / "manifest.json", _MAX_MANIFEST)
        if expected_manifest_sha256 is not None and digest(self.raw) != _hash(expected_manifest_sha256):
            _error("MANIFEST_CHANGED")
        self.manifest = parse(self.raw)
        manifest = self.manifest
        if (manifest.get("schema") != _SCHEMA or manifest.get("complete") is not True
                or not isinstance(manifest.get("files"), dict) or not {"workbook.xlsx", "definition.json"} <= set(manifest["files"])):
            _error("COMPLETE_DELIVERY_REQUIRED")
        _id(manifest.get("id"))
        _hash(manifest.get("definition_sha256"))
        for name, record in manifest["files"].items():
            _relative(name)
            if not isinstance(record, dict) or type(record.get("bytes")) is not int or record["bytes"] < 0:
                _error("FILE_METADATA_INVALID")
            _hash(record.get("sha256"))

    @contextmanager
    def open(self, name):
        if name == "manifest.json":
            with io.BytesIO(self.raw) as stream:
                yield stream
            return
        if name not in self.manifest["files"]:
            _error("FILE_INVALID")
        path = self.directory / _relative(name)
        if path.resolve() != path:
            _error("FILE_LOCATION_CHANGED")
        with _open_private(path) as stream:
            expected = self.manifest["files"][name]
            if (os.fstat(stream.fileno()).st_size != expected["bytes"]
                    or hashlib.file_digest(stream, "sha256").hexdigest() != expected["sha256"]):
                _error("FILE_CHANGED")
            stream.seek(0)
            yield stream

    def inspect(self):
        for name in self.manifest["files"]:
            with self.open(name):
                pass
        with self.open("definition.json") as stream:
            raw = stream.read(_MAX_MANIFEST + 1)
        if len(raw) > _MAX_MANIFEST or digest(raw) != self.manifest["definition_sha256"]:
            _error("DEFINITION_CHANGED")
        current = parse(raw)
        while current.get("refresh_origin"):
            origin = current["refresh_origin"]
            revision = origin.get("base_revision")
            if type(revision) is not int or not 1 <= revision < current["revision"]:
                _error("REFRESH_HISTORY_CHANGED")
            name = "history/definition-" + str(revision) + ".json"
            if origin.get("base_definition_file") != name:
                _error("REFRESH_HISTORY_CHANGED")
            with self.open(name) as stream:
                previous_raw = stream.read(_MAX_MANIFEST + 1)
            if len(previous_raw) > _MAX_MANIFEST or digest(previous_raw) != origin.get("base_definition_sha256"):
                _error("REFRESH_HISTORY_CHANGED")
            previous = parse(previous_raw)
            if (previous.get("name") != current["name"] or previous.get("revision") != revision
                    or previous.get("workspace_id") != current["workspace_id"]):
                _error("REFRESH_HISTORY_CHANGED")
            current = previous
        return {**self.manifest, "manifest_sha256": digest(self.raw)}

    def verify_archive(self, path):
        """A recovered ZIP must contain the complete original manifest and files."""
        with _open_private(path) as stream, zipfile.ZipFile(stream) as package:
            names = [entry.filename for entry in package.infolist()]
            if len(names) != len(set(names)) or set(names) != set(self.manifest["files"]) | {"manifest.json"}:
                _error("ARCHIVE_MEMBERS_INVALID")
            meta = package.getinfo("manifest.json")
            if meta.file_size != len(self.raw) or package.read(meta) != self.raw:
                _error("ARCHIVE_MANIFEST_CHANGED")
            for name, expected in self.manifest["files"].items():
                entry = package.getinfo(name)
                if entry.file_size != expected["bytes"] or stat.S_ISLNK(entry.external_attr >> 16):
                    _error("ARCHIVE_MEMBER_CHANGED")
                with package.open(entry) as member:
                    if hashlib.file_digest(member, "sha256").hexdigest() != expected["sha256"]:
                        _error("ARCHIVE_MEMBER_CHANGED")

    def copy_to(self, destination):
        target = Path(destination).absolute()
        parent = private_directory(target.parent)
        if parent / target.name != target:
            _error("CANONICAL_DELIVERY_DIRECTORY_REQUIRED")
        target.mkdir(mode=0o700)
        try:
            for name in self.manifest["files"]:
                path = target / name
                private_directory(path.parent)
                with self.open(name) as source, path.open("xb") as output:
                    os.fchmod(output.fileno(), 0o600)
                    shutil.copyfileobj(source, output, 1024 * 1024)
                    output.flush()
                    os.fsync(output.fileno())
            write(target / "manifest.json", self.raw)
            sync_directory(target)
        except BaseException:
            shutil.rmtree(target)
            raise
        return WorkbookDelivery(target, expected_manifest_sha256=digest(self.raw))

    @classmethod
    def receive(cls, archive, destination, *, expected_sha256=None, max_bytes=1024 * 1024 * 1024):
        """Receive into private staging, then publish one verified complete directory."""
        if type(max_bytes) is not int or max_bytes < 1:
            _error("RECEIVING_LIMIT_REQUIRED")
        target = Path(destination).absolute()
        parent = private_directory(target.parent)
        target = parent / target.name
        if os.path.lexists(target):
            raise FileExistsError(str(target))
        with _open_private(Path(archive).absolute()) as source:
            archive_sha = hashlib.file_digest(source, "sha256").hexdigest()
            if expected_sha256 is not None and archive_sha != _hash(expected_sha256):
                _error("ARCHIVE_CHANGED")
            source.seek(0)
            # The original operation can reclaim its own unfinished staging after
            # process termination. The retained archive remains the only input.
            control = private_directory(parent / (".workbook-receive-" + digest(os.fsencode(target.name))[:24]))
            with lock(control):
                if os.path.lexists(target):
                    raise FileExistsError(str(target))
                intent_path, staging = control / "source.json", control / "content"
                intent = {"schema": "stellar-forge.workbook-receiving.v1", "destination": str(target),
                          "archive_sha256": archive_sha}
                if intent_path.exists():
                    with _open_private(intent_path) as retained:
                        old = retained.read(64 * 1024 + 1)
                    if len(old) > 64 * 1024 or parse(old) != intent:
                        _error("RECEIVING_OPERATION_CHANGED")
                else:
                    if os.path.lexists(staging):
                        _error("RECEIVING_STATE_INVALID")
                    write(intent_path, _json(intent).encode(), exclusive=True)
                if os.path.lexists(staging):
                    private_directory(staging, create=False)
                    shutil.rmtree(staging)
                staging.mkdir(mode=0o700)
                try:
                    with zipfile.ZipFile(source) as package:
                        infos = package.infolist()
                        names = [entry.filename for entry in infos]
                        if len(names) != len(set(names)) or "manifest.json" not in names:
                            _error("ARCHIVE_MEMBERS_INVALID")
                        meta = package.getinfo("manifest.json")
                        if meta.file_size > _MAX_MANIFEST:
                            _error("MANIFEST_TOO_LARGE")
                        raw = package.read(meta)
                        manifest = parse(raw)
                        if manifest.get("schema") != _SCHEMA or manifest.get("complete") is not True or not isinstance(manifest.get("files"), dict):
                            _error("COMPLETE_DELIVERY_REQUIRED")
                        if set(names) != set(manifest["files"]) | {"manifest.json"}:
                            _error("ARCHIVE_MEMBERS_INVALID")
                        if sum(entry.file_size for entry in infos) > max_bytes:
                            _error("RECEIVING_BYTE_LIMIT_EXCEEDED")
                        for name, expected in manifest["files"].items():
                            _relative(name)
                            entry = package.getinfo(name)
                            if entry.file_size != expected["bytes"] or stat.S_ISLNK(entry.external_attr >> 16):
                                _error("ARCHIVE_MEMBER_CHANGED")
                            path = staging / name
                            private_directory(path.parent)
                            sha, size = hashlib.sha256(), 0
                            with package.open(entry) as input_stream, path.open("xb") as output:
                                os.fchmod(output.fileno(), 0o600)
                                for block in iter(lambda: input_stream.read(1024 * 1024), b""):
                                    sha.update(block)
                                    size += len(block)
                                    if size > expected["bytes"]:
                                        _error("ARCHIVE_MEMBER_CHANGED")
                                    output.write(block)
                                output.flush()
                                os.fsync(output.fileno())
                            if size != expected["bytes"] or sha.hexdigest() != expected["sha256"]:
                                _error("ARCHIVE_MEMBER_CHANGED")
                        write(staging / "manifest.json", raw)
                        for directory in sorted((p for p in staging.rglob("*") if p.is_dir()),
                                                key=lambda p: len(p.parts), reverse=True):
                            sync_directory(directory)
                        sync_directory(staging)
                    result = cls(staging)
                    result.inspect()
                    _publish_received_directory(staging, target)
                    sync_directory(control)
                    return cls(target)
                except BaseException:
                    if staging.exists():
                        shutil.rmtree(staging)
                    raise


class AnalysisWorkbooks:
    def __init__(self, database, *, identity, reports=None, initialize=True):
        self.database, self.identity, self.reports = database, identity, reports
        from .workbook_refresh import WorkbookRefresh
        self.refresh = WorkbookRefresh(self)
        self.directory = database.directory / "analysis-workbooks"
        self._mutex = threading.Lock()
        self._signals = {}
        self._closed = False
        if initialize:
            private_directory(self.directory)
            with database._connection() as db:
                db.execute("BEGIN IMMEDIATE")
                db.execute("CREATE TABLE IF NOT EXISTS sf_analysis_workbooks(name TEXT PRIMARY KEY,revision INTEGER NOT NULL,record TEXT NOT NULL,sha256 TEXT NOT NULL)")
                db.execute("CREATE TABLE IF NOT EXISTS sf_analysis_workbook_history(sequence INTEGER PRIMARY KEY,name TEXT NOT NULL,revision INTEGER NOT NULL,request_id TEXT NOT NULL UNIQUE,request_sha256 TEXT NOT NULL,record TEXT NOT NULL,sha256 TEXT NOT NULL,UNIQUE(name,revision))")
                db.execute("CREATE TABLE IF NOT EXISTS sf_analysis_workbook_control(id TEXT PRIMARY KEY REFERENCES sf_analysis_jobs(id),cancel_requested INTEGER NOT NULL DEFAULT 0,progress TEXT NOT NULL)")
                db.commit()

    @staticmethod
    def state(db):
        if db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='sf_analysis_workbooks'").fetchone() is None:
            return []
        return [{"name": row["name"], "revision": row["revision"], "sha256": row["sha256"],
                 "title": _unpack(row)["title"], "sheets": len(_unpack(row)["sheets"])}
                for row in db.execute("SELECT * FROM sf_analysis_workbooks ORDER BY name")]

    def source(self, request):
        if not isinstance(request, dict) or request.get("kind") not in {"query-export", "report"}:
            _error("SOURCE_REQUIRED")
        kind, identifier = request["kind"], _id(request.get("id"))
        with self.database._connection(readonly=True) as db:
            db.execute("BEGIN")
            if kind == "query-export":
                row = db.execute("SELECT request,result FROM sf_analysis_jobs WHERE id=? AND kind='query-export' AND status='COMPLETE'", (identifier,)).fetchone()
                if row is None or row["result"] is None:
                    _error("COMPLETE_TYPED_QUERY_EXPORT_REQUIRED")
                result = parse(row["result"].encode())
                export = QueryExport(self.database.directory / "analysis-query-exports" / identifier)
                raw = read(export.directory / "manifest.json", 16 * 1024 * 1024)
                sha = digest(raw)
                manifest = export.manifest
                if result.get("manifest_sha256") != sha:
                    _error("SOURCE_MANIFEST_CHANGED")
                if "rows.jsonl" not in manifest["outputs"]:
                    _error("ORIGINAL_TYPED_JSONL_REQUIRED")
                item = {"kind": kind, "id": identifier, "sha256": sha, "label": result.get("label", identifier),
                        "columns": manifest["columns"], "row_count": manifest["rowCount"],
                        "binding_sha256": manifest["bindingSha256"]}
            else:
                if self.reports is None:
                    _error("ORIGINAL_REPORT_READER_REQUIRED")
                row = db.execute("SELECT * FROM sf_analysis_report_publications WHERE id=?", (identifier,)).fetchone()
                publication = _unpack(row)
                report = AnalysisReports._report(db, publication["job_id"])
                if report["result_sha256"] != publication["result_sha256"]:
                    _error("SOURCE_REPORT_CHANGED")
                section = request.get("section")
                if type(section) is not int or not 0 <= section < len(report["sections"]):
                    _error("REPORT_SECTION_REQUIRED")
                selected = report["sections"][section]
                item = {"kind": kind, "id": identifier, "sha256": row["sha256"], "section": section,
                        "job_id": report["job_id"], "result_sha256": report["result_sha256"],
                        "label": publication["title"] + " / " + selected["title"],
                        "columns": selected["columns"], "row_count": selected["row_count"]}
        if request.get("sha256") is not None and request["sha256"] != item["sha256"]:
            _error("SOURCE_CHANGED")
        return item

    def sources(self, *, kind="query-export", before=None, limit=25):
        _page(before, limit)
        with self.database._connection(readonly=True) as db:
            if kind == "query-export":
                rows = db.execute("SELECT sequence,id FROM sf_analysis_jobs WHERE kind='query-export' AND status='COMPLETE' AND (? IS NULL OR sequence<?) ORDER BY sequence DESC LIMIT ?", (before, before, limit + 1)).fetchall()
            elif kind == "report" and self.reports is not None:
                rows = db.execute("SELECT sequence,id,job_id FROM sf_analysis_report_publications WHERE (? IS NULL OR sequence<?) ORDER BY sequence DESC LIMIT ?", (before, before, limit + 1)).fetchall()
            else:
                _error("SOURCE_KIND_INVALID")
        items = []
        for row in rows[:limit]:
            if kind == "report":
                with self.database._connection(readonly=True) as db:
                    sections = len(AnalysisReports._report(db, row["job_id"])["sections"])
                items.extend(self.source({"kind": kind, "id": row["id"], "section": section}) for section in range(sections))
            else:
                try:
                    items.append(self.source({"kind": kind, "id": row["id"]}))
                except StateError as exc:
                    if exc.code != "ANALYSIS_WORKBOOK_ORIGINAL_TYPED_JSONL_REQUIRED":
                        raise
                    items.append({"kind": kind, "id": row["id"], "available": False, "error": exc.code})
        return {"items": items, "next_before": rows[limit - 1]["sequence"] if len(rows) > limit else None}

    def compose(self, name, request):
        """Normalize a complete definition before its compare-and-swap save."""
        name = _name(name)
        title = _text(request.get("title"), 300, required=True)
        notes = _text(request.get("notes", ""), 16_000)
        sheets = request.get("sheets")
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= 32:
            _error("ONE_TO_32_SHEETS_REQUIRED")
        result = []
        for sheet in sheets:
            if not isinstance(sheet, dict):
                _error("SHEET_REQUIRED")
            label = _text(sheet.get("name"), 124, required=True)
            if len(label) > 31 or _BAD_SHEET.search(label) or label.startswith("'") or label.endswith("'") or any(ord(c) < 32 for c in label):
                _error("EXCEL_SHEET_NAME_INVALID")
            source = self.source(sheet.get("source"))
            columns = sheet.get("columns")
            if not isinstance(columns, list) or not 1 <= len(columns) <= _MAX_COLUMNS:
                _error("COLUMNS_REQUIRED")
            projected, headers = [], set()
            for column in columns:
                source_key = _source_key(source, column)
                header = _literal(_text(column.get("header", column["source"]), 2000, required=True))
                if header.casefold() in headers:
                    _error("COLUMN_HEADERS_MUST_BE_UNIQUE")
                headers.add(header.casefold())
                kind = column.get("type", "native")
                null, precision = column.get("null", "text"), column.get("precision", "preserve")
                if kind not in _TYPES or null not in {"blank", "text"} or precision not in {"preserve", "excel-number"}:
                    _error("COLUMN_PROJECTION_INVALID")
                projected.append({"source": column["source"], "header": header, "type": kind, "null": null,
                                  "null_text": _literal(_text(column.get("null_text", "NULL"), 1000)), "precision": precision,
                                  **({"source_index": source_key} if source["kind"] == "report" and "source_index" in column else {})})
            result.append({"name": label, "source": source, "columns": projected})
        _sheet_layout(result)
        return {"schema": _DEFINITION, "name": name, "title": title, "notes": notes, "sheets": result}

    def save(self, name, request, *, refresh_origin=None):
        request_id = _id(request.get("id"))
        expected = request.get("expected_revision")
        if type(expected) is not int or expected < 0:
            _error("REVISION_REQUIRED")
        normalized = self.compose(name, request)
        if refresh_origin is not None:
            normalized["refresh_origin"] = refresh_origin
        request_sha = digest(_json({"id": request_id, "expected_revision": expected, **normalized}).encode())
        with self.database._connection() as db:
            db.execute("BEGIN IMMEDIATE")
            old = db.execute("SELECT * FROM sf_analysis_workbook_history WHERE request_id=?", (request_id,)).fetchone()
            if old is not None:
                if old["request_sha256"] != request_sha:
                    _error("REQUEST_ID_REUSED")
                return {**_unpack(old), "definition_sha256": old["sha256"]}
            head = db.execute("SELECT revision FROM sf_analysis_workbooks WHERE name=?", (name,)).fetchone()
            if (head[0] if head else 0) != expected:
                _error("DEFINITION_CHANGED")
            record = {**normalized, "revision": expected + 1, "saved_at": _now(), "workspace_id": self.identity}
            raw = _json(record)
            sha = digest(raw.encode())
            db.execute("INSERT INTO sf_analysis_workbooks VALUES(?,?,?,?) ON CONFLICT(name) DO UPDATE SET revision=excluded.revision,record=excluded.record,sha256=excluded.sha256", (name, record["revision"], raw, sha))
            db.execute("INSERT INTO sf_analysis_workbook_history(name,revision,request_id,request_sha256,record,sha256) VALUES(?,?,?,?,?,?)", (name, record["revision"], request_id, request_sha, raw, sha))
            db.commit()
            return {**record, "definition_sha256": sha}

    def definition(self, name, *, revision=None):
        name = _name(name)
        if revision is not None and (type(revision) is not int or revision < 1):
            _error("REVISION_REQUIRED")
        with self.database._connection(readonly=True) as db:
            db.execute("BEGIN")
            if revision is None:
                row = db.execute("SELECT * FROM sf_analysis_workbooks WHERE name=?", (name,)).fetchone()
            else:
                row = db.execute("SELECT * FROM sf_analysis_workbook_history WHERE name=? AND revision=?", (name, revision)).fetchone()
            record = _unpack(row)
            return {"definition": record, "definition_sha256": row["sha256"],
                    "head_revision": db.execute("SELECT revision FROM sf_analysis_workbooks WHERE name=?", (name,)).fetchone()[0],
                    "history": [dict(r) for r in db.execute("SELECT revision,sha256 FROM sf_analysis_workbook_history WHERE name=? ORDER BY revision DESC LIMIT 100", (name,))]}

    def prepare(self, request):
        result = self.definition(request.get("name"), revision=request.get("revision"))
        if request.get("definition_sha256") != result["definition_sha256"]:
            _error("DEFINITION_CHANGED")
        return {"name": result["definition"]["name"], "revision": result["definition"]["revision"],
                "definition_sha256": result["definition_sha256"]}

    @contextmanager
    def _rows(self, source):
        current = self.source(source)
        if current != source:
            _error("SOURCE_CHANGED")
        if source["kind"] == "query-export":
            iterator = QueryExport(self.database.directory / "analysis-query-exports" / source["id"]).rows()
            try:
                yield iterator
            finally:
                iterator.close()
        else:
            with self.database._connection(readonly=True) as db:
                db.execute("BEGIN")
                publication, report = self.reports._export_input(db, {"publication_id": source["id"], "publication_sha256": source["sha256"]})
                selected = report["sections"][source["section"]]
                def values():
                    for cells in AnalysisReports._rows(db, report, selected):
                        result = []
                        for cell in cells:
                            kind, text = cell["type"], cell["text"]
                            result.append(None if kind == "null" else int(text) if kind == "integer" else float(text) if kind in {"number", "sqlite-real"} else base64.b64decode(text, validate=True) if kind == "blob" else text)
                        yield result
                yield values()

    def preview(self, name, *, revision=None, sheet=0, limit=20):
        result = self.definition(name, revision=revision)
        sheets = result["definition"]["sheets"]
        if type(sheet) is not int or not 0 <= sheet < len(sheets) or type(limit) is not int or not 1 <= limit <= 100:
            _error("PREVIEW_PAGE_INVALID")
        selected, rows = sheets[sheet], []
        keys = [_source_key(selected["source"], column) for column in selected["columns"]]
        with self._rows(selected["source"]) as stream:
            for row in stream:
                values = []
                for column, key in zip(selected["columns"], keys):
                    value, number_format, conversion = project_cell(row[key], column)
                    values.append({"value": value.isoformat() if isinstance(value, (date, datetime)) else str(value) if type(value) is int and abs(value) >= 2**53 else value,
                                   "excel_type": "blank" if value is None else "text" if isinstance(value, str) else "date" if isinstance(value, (date, datetime)) else "boolean" if isinstance(value, bool) else "number",
                                   "format": number_format, "conversion": conversion})
                rows.append(values)
                if len(rows) >= limit:
                    break
        return {"definition_sha256": result["definition_sha256"], "sheet": selected, "rows": rows,
                "row_count": selected["source"]["row_count"], "preview_only": len(rows) < selected["source"]["row_count"]}

    def _retain_source(self, source, target):
        target = target / "sources" / (source["kind"] + "-" + source["id"])
        if target.exists():
            return
        private_directory(target.parent)
        if source["kind"] == "query-export":
            QueryExport(self.database.directory / "analysis-query-exports" / source["id"]).copy_to(target)
            if digest(read(target / "manifest.json", 16 * 1024 * 1024)) != source["sha256"]:
                _error("SOURCE_CHANGED")
        else:
            private_directory(target)
            with self.database._connection(readonly=True) as db, (target / "report.json").open("xb") as output:
                os.fchmod(output.fileno(), 0o600)
                db.execute("BEGIN")
                publication, report = self.reports._export_input(db, {"publication_id": source["id"], "publication_sha256": source["sha256"]})
                for block in self.reports._json_blocks(db, publication, report):
                    output.write(block)
                output.flush()
                os.fsync(output.fileno())

    def build(self, definition, job_id, destination, *, cancel=lambda: False, progress=lambda value: None):
        """Build an entire portable delivery from pinned native results, without DB writes."""
        try:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            _error("INSTALL_WORKBOOKS_EXTRA")
        job_id = _id(job_id)
        if definition.get("schema") != _DEFINITION or type(definition.get("revision")) is not int:
            _error("SAVED_DEFINITION_REQUIRED")
        raw = _json(definition).encode()
        definition_sha = digest(raw)
        layout = _sheet_layout(definition["sheets"])
        target = Path(destination).absolute()
        private_directory(target.parent)
        target.mkdir(mode=0o700)
        workbook = Workbook(write_only=True, iso_dates=True)
        workbook.properties.title = definition["title"]
        workbook.properties.creator = "Stellar Forge"
        workbook.properties.description = "Customer workbook with original typed results and source bindings in its handoff."
        header_fill = PatternFill("solid", fgColor="E7E7E7")
        header_font = Font(name="Calibri", size=11, bold=True, color="202020")
        outputs, total_rows = [], 0

        def stopped():
            if cancel():
                _error("CANCELLED")

        def cell(sheet, value, number_format=None, heading=False):
            item = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                item.data_type = "s"  # '=' remains original text, never a formula.
            if number_format:
                item.number_format = number_format
            if heading:
                item.fill, item.font = header_fill, header_font
                item.alignment = Alignment(vertical="top", wrap_text=True)
            return item

        def new_sheet(name, columns, row_count):
            sheet = workbook.create_sheet(name)
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = True
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row_count + 1}"
            for index, column in enumerate(columns, 1):
                sheet.column_dimensions[get_column_letter(index)].width = min(48, max(14, len(column["header"]) + 3))
            sheet.append([cell(sheet, _literal(column["header"]), heading=True) for column in columns])
            return sheet

        try:
            write(target / "definition.json", raw)
            definition_history = self.refresh.retain_history(definition, target, cancel=cancel)
            for ordinal, logical in enumerate(definition["sheets"]):
                stopped()
                source = logical["source"]
                keys = [_source_key(source, column) for column in logical["columns"]]
                if self.source(source) != source:
                    _error("SOURCE_CHANGED")
                self._retain_source(source, target)
                expected = source["row_count"]
                counts, row_count, physical_index = {}, 0, 0
                sheet = new_sheet(layout[ordinal][0], logical["columns"], min(expected, _ROWS_PER_SHEET))
                with self._rows(source) as stream:
                    for values in stream:
                        stopped()
                        if row_count and row_count % _ROWS_PER_SHEET == 0:
                            physical_index += 1
                            sheet = new_sheet(layout[ordinal][physical_index], logical["columns"], min(expected - row_count, _ROWS_PER_SHEET))
                        cells = []
                        for index, (column, key) in enumerate(zip(logical["columns"], keys)):
                            value, format, conversion = project_cell(values[key], column)
                            cells.append(cell(sheet, value, format))
                            if conversion:
                                key = str(index) + ":" + conversion
                                counts[key] = counts.get(key, 0) + 1
                        sheet.append(cells)
                        row_count += 1
                        total_rows += 1
                        if row_count % 1000 == 0:
                            progress({"sheet": logical["name"], "sheet_rows": row_count, "total_rows": total_rows, "expected_rows": sum(s["source"]["row_count"] for s in definition["sheets"])})
                if row_count != expected:
                    _error("SOURCE_ROW_COUNT_CHANGED")
                outputs.append({"name": logical["name"], "worksheets": layout[ordinal], "row_count": row_count,
                                "source": source, "columns": logical["columns"], "conversions": counts})
                progress({"sheet": logical["name"], "sheet_rows": row_count, "total_rows": total_rows})
            readme = workbook.create_sheet("Readme")
            readme.column_dimensions["A"].width = 28
            readme.column_dimensions["B"].width = 110
            for key, value in [("Workbook", definition["title"]), ("Notes", definition["notes"]),
                               ("Saved definition", definition["name"] + " / revision " + str(definition["revision"])),
                               ("Definition SHA-256", definition_sha), ("Delivery", job_id),
                               ("Text and formulas", "All source strings are literal Excel text, including strings beginning with =. Source formulas are not evaluated."),
                               ("Precision", "Preserve keeps values beyond Excel's numeric precision as text. Excel-number is an explicit column choice. Decimal-text preserves original decimal spelling and scale."),
                               ("Null and empty text", "Each column selects a null label or an empty cell. Excel can display nulls and empty strings alike; the handoff retains their original typed values."),
                               ("Dates", "ISO dates use the 1900 date system. UTC timestamps require a timezone. Values before 1900 or beyond millisecond precision remain original text."),
                               ("Complete rows", "Each data sheet includes every selected source row. Larger results continue on numbered sheets, with a header repeated on each sheet."),
                               ("Offline handoff", "Open manifest.json and the original source exports alongside this workbook to recover exact input types, query bindings, and projections.")]:
                readme.append([cell(readme, key, heading=True), cell(readme, _literal(value))])
            if definition.get("refresh_origin"):
                origin = definition["refresh_origin"]
                readme.append([cell(readme, "Edition note", heading=True), cell(readme, _literal(origin["reason"]))])
                readme.append([cell(readme, "Previous definition", heading=True), cell(readme, origin["base_definition_file"])])
            meta = new_sheet("Sources", [{"header": h} for h in ["Sheet", "Source kind", "Source ID", "Source SHA-256", "Rows", "Input column", "Input ordinal", "Excel heading", "Type", "Null", "Precision"]], sum(len(s["columns"]) for s in outputs))
            for logical in outputs:
                for column in logical["columns"]:
                    values = [logical["name"], logical["source"]["kind"], logical["source"]["id"], logical["source"]["sha256"], logical["row_count"], column["source"], _source_key(logical["source"], column) + 1 if logical["source"]["kind"] == "report" else "", column["header"], column["type"], "blank" if column["null"] == "blank" else column["null_text"], column["precision"]]
                    meta.append([cell(meta, value) for value in values])
            stopped()
            workbook_path = target / "workbook.xlsx"
            with workbook_path.open("xb") as output:
                os.fchmod(output.fileno(), 0o600)
                workbook.save(output)
                output.flush()
                os.fsync(output.fileno())
            files = {}
            for path in sorted(target.rglob("*")):
                if path.is_file():
                    relative = path.relative_to(target).as_posix()
                    _relative(relative)
                    files[relative] = _file_record(path)
            manifest = {"schema": _SCHEMA, "complete": True, "id": job_id, "workspace_id": self.identity,
                        "name": definition["name"], "revision": definition["revision"], "title": definition["title"],
                        "definition_sha256": definition_sha, "created_at": _now(), "row_count": total_rows,
                        "sheets": outputs, "files": files, "definition_history": definition_history, "excel": {"date_system": "1900", "date_encoding": "ISO8601", "strings": "literal", "data_rows_per_sheet": _ROWS_PER_SHEET}}
            manifest_raw = _json(manifest).encode()
            if len(manifest_raw) > _MAX_MANIFEST:
                _error("MANIFEST_TOO_LARGE")
            write(target / "manifest.json", manifest_raw)
            stopped()
            with (target / "handoff.zip").open("xb") as output:
                os.fchmod(output.fileno(), 0o600)
                with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6, allowZip64=True) as package:
                    for name in [*files, "manifest.json"]:
                        stopped()
                        info = zipfile.ZipInfo(name)
                        info.compress_type = zipfile.ZIP_DEFLATED
                        info.external_attr = 0o100600 << 16
                        with _open_private(target / name) as source, package.open(info, "w", force_zip64=True) as member:
                            shutil.copyfileobj(source, member, 1024 * 1024)
                output.flush()
                os.fsync(output.fileno())
            sync_directory(target)
            return self._completed(target, expected_definition=definition_sha, inspect=False)
        except BaseException:
            shutil.rmtree(target)
            raise
        finally:
            workbook.close()
            # openpyxl removes its worksheet staging files on successful save.
            # Close and remove only staging paths owned by this failed workbook.
            for sheet in workbook.worksheets:
                writer = getattr(sheet, "_writer", None)
                if writer is not None and isinstance(writer.out, str) and os.path.exists(writer.out):
                    try:
                        sheet.close()
                    finally:
                        Path(writer.out).unlink(missing_ok=True)

    def _completed(self, path, *, expected_definition, inspect=True):
        delivery = WorkbookDelivery(path)
        manifest = delivery.inspect() if inspect else delivery.manifest
        if manifest["definition_sha256"] != expected_definition:
            _error("DEFINITION_CHANGED")
        if manifest["id"] != path.name:
            _error("DELIVERY_ID_CHANGED")
        if inspect:
            try:
                delivery.verify_archive(path / "handoff.zip")
            except zipfile.BadZipFile:
                _error("COMPLETE_ARCHIVE_REQUIRED")
        files = [{"name": name, **(_file_record(path / name) if name == "handoff.zip" else
                   {"bytes": len(delivery.raw), "sha256": digest(delivery.raw)} if name == "manifest.json" else manifest["files"][name]), "mime": mime}
                 for name, mime in _MIMES.items()]
        return {"schema": _SCHEMA, "id": manifest["id"], "name": manifest["name"], "revision": manifest["revision"],
                "title": manifest["title"], "definition_sha256": expected_definition, "directory": str(path),
                "manifest_sha256": digest(delivery.raw), "row_count": manifest["row_count"],
                "sheets": manifest["sheets"], "files": files, "completed_at": manifest["created_at"]}

    def _progress(self, job_id, values):
        with self.database._connection() as db:
            db.execute("INSERT INTO sf_analysis_workbook_control(id,progress) VALUES(?,?) ON CONFLICT(id) DO UPDATE SET progress=excluded.progress", (job_id, _json({**values, "updated_at": _now()})))

    def progress(self, job_id):
        with self.database._connection(readonly=True) as db:
            row = db.execute("SELECT progress,cancel_requested FROM sf_analysis_workbook_control WHERE id=?", (_id(job_id),)).fetchone()
        return {**parse(row["progress"].encode()), "cancel_requested": bool(row["cancel_requested"])} if row else {}

    def cancel(self, job_id):
        job_id = _id(job_id)
        with self.database._connection() as db:
            row = db.execute("SELECT status FROM sf_analysis_jobs WHERE id=? AND kind='workbook'", (job_id,)).fetchone()
            if row is None:
                _error("NOT_FOUND")
            if row["status"] == "RUNNING":
                db.execute("INSERT INTO sf_analysis_workbook_control(id,cancel_requested,progress) VALUES(?,1,'{}') ON CONFLICT(id) DO UPDATE SET cancel_requested=1", (job_id,))
        with self._mutex:
            if job_id in self._signals:
                self._signals[job_id].set()
        return self.progress(job_id)

    def close(self):
        with self._mutex:
            self._closed = True
            for signal in self._signals.values():
                signal.set()

    def _finish(self, job_id, status, *, result=None, error=None):
        with self.database._connection() as db:
            db.execute("UPDATE sf_analysis_jobs SET status=?,result=?,error=?,finished_at=? WHERE id=? AND kind='workbook'", (status, _json(result) if result else None, error, _now(), job_id))

    def run(self, job_id, spec):
        job_id = _id(job_id)
        path = self.directory / job_id
        with self._mutex:
            signal = self._signals.setdefault(job_id, threading.Event())
            if self._closed or self.progress(job_id).get("cancel_requested"):
                signal.set()
        try:
            definition = self.definition(spec["name"], revision=spec["revision"])
            if definition["definition_sha256"] != spec["definition_sha256"]:
                _error("DEFINITION_CHANGED")
            result = self.build(definition["definition"], job_id, path, cancel=signal.is_set,
                                progress=lambda value: self._progress(job_id, value))
            self._finish(job_id, "COMPLETE", result=result)
        except BaseException as exc:
            try:
                result = self._completed(path, expected_definition=spec["definition_sha256"])
            except (BuyerError, OSError, ValueError, TypeError, KeyError):
                code = getattr(exc, "code", None) or type(exc).__name__
                self._finish(job_id, "CANCELLED" if code == "ANALYSIS_WORKBOOK_CANCELLED" else "FAILED", error=str(code)[:256])
            else:
                self._finish(job_id, "COMPLETE", result=result)
        finally:
            with self._mutex:
                self._signals.pop(job_id, None)

    def recover_completed(self):
        with self.database._connection(readonly=True) as db:
            rows = db.execute("SELECT id,request FROM sf_analysis_jobs WHERE kind='workbook' AND status IN ('INTERRUPTED','FAILED','CANCELLED')").fetchall()
        for row in rows:
            spec = parse(row["request"].encode())
            path = self.directory / row["id"]
            if not (path / "handoff.zip").is_file():
                continue
            try:
                result = self._completed(path, expected_definition=spec["definition_sha256"])
            except (BuyerError, OSError, ValueError, TypeError, KeyError):
                continue
            self._finish(row["id"], "COMPLETE", result=result)

    def history(self, *, before=None, limit=25):
        _page(before, limit)
        with self.database._connection(readonly=True) as db:
            rows = db.execute("SELECT sequence,id,status,created_at,finished_at,request,result,error FROM sf_analysis_jobs WHERE kind='workbook' AND (? IS NULL OR sequence<?) ORDER BY sequence DESC LIMIT ?", (before, before, limit + 1)).fetchall()
        items = []
        for row in rows[:limit]:
            request = parse(row["request"].encode())
            item = {key: row[key] for key in ("sequence", "id", "status", "created_at", "finished_at", "error")}
            item.update(name=request["name"], revision=request["revision"], progress=self.progress(row["id"]))
            if row["result"]:
                result = parse(row["result"].encode())
                item.update(row_count=result["row_count"], sheets=len(result["sheets"]), files=result["files"], title=result["title"], manifest_sha256=result["manifest_sha256"])
            items.append(item)
        return {"items": items, "next_before": rows[limit - 1]["sequence"] if len(rows) > limit else None}

    @contextmanager
    def download(self, job_id, name):
        job_id = _id(job_id)
        if name not in _MIMES:
            _error("FILE_INVALID")
        with self.database._connection(readonly=True) as db:
            row = db.execute("SELECT result FROM sf_analysis_jobs WHERE id=? AND kind='workbook' AND status='COMPLETE'", (job_id,)).fetchone()
        if row is None or row["result"] is None:
            _error("COMPLETE_DELIVERY_REQUIRED")
        result = parse(row["result"].encode())
        path = self.directory / job_id
        delivery = WorkbookDelivery(path, expected_manifest_sha256=result["manifest_sha256"])
        record = next((f for f in result["files"] if f["name"] == name), None)
        if record is None:
            _error("FILE_INVALID")
        if name == "handoff.zip":
            with _open_private(path / name) as stream:
                if os.fstat(stream.fileno()).st_size != record["bytes"] or hashlib.file_digest(stream, "sha256").hexdigest() != record["sha256"]:
                    _error("FILE_CHANGED")
                stream.seek(0)
                yield stream, record
        else:
            with delivery.open(name) as stream:
                yield stream, record


def main():
    parser = argparse.ArgumentParser(description="Inspect, receive or copy a complete customer workbook handoff offline.")
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect")
    inspect.add_argument("--delivery", required=True)
    inspect.add_argument("--manifest-sha256")
    copy = commands.add_parser("copy")
    copy.add_argument("--delivery", required=True)
    copy.add_argument("--destination", required=True)
    copy.add_argument("--manifest-sha256")
    receive = commands.add_parser("receive")
    receive.add_argument("--archive", required=True)
    receive.add_argument("--destination", required=True)
    receive.add_argument("--sha256")
    receive.add_argument("--max-bytes", type=int, default=1024 * 1024 * 1024)
    args = parser.parse_args()
    if args.command == "receive":
        delivery = WorkbookDelivery.receive(args.archive, args.destination, expected_sha256=args.sha256, max_bytes=args.max_bytes)
    else:
        delivery = WorkbookDelivery(args.delivery, expected_manifest_sha256=args.manifest_sha256)
        if args.command == "copy":
            delivery = delivery.copy_to(args.destination)
    print(_json(delivery.inspect()))


if __name__ == "__main__":
    main()

