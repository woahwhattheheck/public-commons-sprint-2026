"""Complete existing DBnomics source -> Analysis -> workbook -> reviewed return.

Run this against the already completed native DBnomics receiving project.
No new provider capture or purchase is requested. An optional real analyst note
is written to one explicitly selected observation, then reviewed and adopted.
"""
from __future__ import annotations
import argparse
from contextlib import closing
import json
from pathlib import Path
import tempfile
import os

from stellar_forge_buyer.dbnomics import (
    DbnomicsAnalysisSource, DbnomicsEconomicWorkflow, canonical, keep, need,
    native, pin, write_exact,
)
from stellar_forge_buyer.storage import private_directory, read, parse


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, allow_abbrev=False)
    parser.add_argument("--checkout", required=True)
    parser.add_argument("--node", default="node")
    parser.add_argument("--project", required=True)
    parser.add_argument("--operation", required=True)
    parser.add_argument("--receiving-id", default="dbnomics-insee-ipc-20261012")
    parser.add_argument("--title", default="Insee annual and monthly economic briefing, including imputed rents")
    parser.add_argument("--review-note", required=True)
    parser.add_argument("--annotation", help="Actual analyst commentary to return in the original Records sheet")
    parser.add_argument("--observation-id", help="Exact original observation identity for the supplied commentary")
    parser.add_argument("--author", default="Economic analyst")
    args = parser.parse_args(argv)
    root = private_directory(Path(args.operation).expanduser().absolute())
    project = Path(args.project).expanduser().absolute()
    need(bool(args.annotation) == bool(args.observation_id), "EXACT_ANNOTATION_OBSERVATION_REQUIRED")
    status = native(args.checkout, args.node, ["project-status", "--workspace", str(project)])
    # This uses the complete original native receiving result, not a new source.
    receiving = parse(read(project / "receiving" / args.receiving_id / "result.json", 32 * 1024 * 1024))
    captures = status["captures"]
    need(len(captures) == 1, "SELECT_THE_ORIGINAL_COMPLETE_RECEIVING_PROJECT")
    capture = captures[0]
    need(receiving.get("schema") == "stellar-forge.dbnomics-receiving-result.v1" and
         receiving.get("operationId") == args.receiving_id and
         receiving.get("projectId") == status["project"]["id"] and
         receiving.get("capture", {}).get("id") == capture["captureId"] and
         receiving.get("capture", {}).get("sha256") == capture["captureSha256"] and
         receiving.get("reopened", {}).get("state") == "PUBLISHED", "ORIGINAL_RECEIVING_RESULT_REQUIRED")
    reports = [item for item in status["reports"] if item["captureId"] == capture["captureId"]]
    need(reports and all(item["captureId"] == capture["captureId"] for item in reports), "COMPLETE_NATIVE_REPORTS_REQUIRED")
    selection_path = root / "selection.json"
    if selection_path.exists():
        selection = parse(read(selection_path, 1024 * 1024))
    else:
        notes = native(args.checkout, args.node, ["notes", "--workspace", str(project), "--id", capture["captureId"]])
        selection = {"captureId": capture["captureId"], "captureSha256": capture["captureSha256"],
            "reportIds": [x["id"] for x in reports], "notesRevision": notes["revision"], "title": args.title}
        keep(selection_path, selection)
    keep(root / "original-receiving.json", {"project_id": status["project"]["id"],
        "receiving_id": args.receiving_id, "receiving_result_sha256": pin(project / "receiving" / args.receiving_id / "result.json")["sha256"],
        "selection": selection})
    source = DbnomicsAnalysisSource.create(project, root / "source-selected", selection,
        operation_id="economic-analysis-" + args.receiving_id, checkout=args.checkout, node=args.node)
    workflow = DbnomicsEconomicWorkflow(root / "workflow", checkout=args.checkout, node=args.node)
    review = workflow.prepare(source)
    complete = workflow.finish(preview_sha256=review["preview_sha256"], review_note=args.review_note)
    base_reopen = workflow.reopen(source_target=root / "received-economic-project")
    result = {"source": source.verify(), "review": review, "complete": complete, "base_reopen": base_reopen}
    if args.annotation is not None:
        from openpyxl import load_workbook
        need(1 <= len(args.annotation) <= 16000, "ANALYST_COMMENT_REQUIRED")
        request = keep(root / "annotation.json", {"observation_id": args.observation_id,
            "text": args.annotation, "author": args.author, "source_document_id": source.document_id})
        returned = root / "returned-economic-workbook.xlsx"
        if not returned.exists():
            book = load_workbook(complete["workbook"]["files"]["workbook.xlsx"]["path"], data_only=False, keep_links=False)
            try:
                sheet = book["Records"]
                header = {cell.value: i for i, cell in enumerate(next(sheet.iter_rows()), start=1)}
                matches = []
                for row in sheet.iter_rows(min_row=2):
                    if row[header["record_kind"] - 1].value == "observation" and row[header["observation_id"] - 1].value == args.observation_id:
                        matches.append(row[header["review_note"] - 1])
                need(len(matches) == 1, "EXACT_ORIGINAL_OBSERVATION_NOT_FOUND")
                matches[0].value = args.annotation
                matches[0].data_type = "s"
                fd, temporary = tempfile.mkstemp(prefix=".returned-economic-", suffix=".xlsx", dir=root)
                os.close(fd)
                try:
                    book.save(temporary)
                    with open(temporary, "rb") as stream:
                        write_exact(returned, stream.read())
                finally:
                    os.unlink(temporary)
            finally:
                book.close()
        returned_file = pin(returned)
        compared = workflow.return_workbook("analyst-note-return", returned, expected_sha256=returned_file["sha256"])
        accepted = workflow.accept_return("analyst-note-return", result_sha256=compared["result_sha256"],
            note="Reviewed the complete returned Records edition against its original workbook. " + args.review_note)
        adopted = workflow.adopt_notes("analyst-note-return", project, author=args.author)
        reopened = workflow.reopen(identifier="analyst-note-return", source_target=root / "returned-source-project")
        result["returned"] = {"file": returned_file, "comparison": compared, "accepted": accepted,
                              "notes_adopted": adopted, "reopened": reopened}
    # The complete returned-workbook result remains distinct from a previously
    # completed base run, so adding a later real annotation cannot disappear
    # behind that earlier immutable receiving receipt.
    result_path = root / ("result-with-return.json" if args.annotation is not None else "result.json")
    result = keep(result_path, result, identity=("source", "review", "complete"))
    print(canonical({**result, "result_file": pin(result_path)}).decode("utf-8"), flush=True)


if __name__ == "__main__":
    main()
