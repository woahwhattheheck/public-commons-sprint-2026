"""XLSX file I/O for the original supplier review engine.

This module only carries source-bound review cells. Reconciliation, decision
validation, revision handling and finalization belong to supplier-report.mjs.
It uses the existing Analysis openpyxl workbook dependency.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import math
import re
import stat
import sys
import zipfile
import xml.etree.ElementTree as ET
from typing import Any

MODEL_SCHEMA = "conapi.supplier-review-workbook.v1"
RETURN_SCHEMA = "conapi.supplier-review-workbook-return.v1"
IO_SCHEMA = "conapi.supplier-review-workbook-io.v1"
MAX_JSON_BYTES = 96 * 1024 * 1024
MAX_XLSX_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_MEMBERS = 4096
MAX_ROWS = 1_048_576
CHUNK_CHARS = 24_000
REVIEW_HEADER_ROW = 10
REVIEW_HEADERS = [
    "Item ID", "Key", "Change", "Baseline", "Incoming", "Saved decision",
    "Decision", "Reason", "Source side", "Source row",
]
VALUE_HEADERS = ["Item ID", "Field", "Part", "Value kind", "Encoding", "Value"]
SOURCE_HEADERS = ["Item ID", "Source side", "Source row", "Field", "Part", "Encoding", "Value"]
ACTIONS = ["use_candidate", "keep_baseline", "exclude", "replace", "choose_source"]
KINDS = ["text", "null"]
ENCODINGS = ["text", "json"]
XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class WorkbookProblem(ValueError):
    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(message)


def problem(code: str, message: str):
    raise WorkbookProblem(code, message)


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def object_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            problem("WORKBOOK_JSON_INVALID", "Repeated JSON property: " + str(key))
        result[key] = value
    return result


def parse_json(text: str):
    def invalid_constant(value):
        problem("WORKBOOK_JSON_INVALID", "Non-finite JSON value: " + value)
    try:
        return json.loads(text, object_pairs_hook=object_pairs, parse_constant=invalid_constant)
    except WorkbookProblem:
        raise
    except (ValueError, TypeError) as exc:
        problem("WORKBOOK_JSON_INVALID", str(exc))


def json_text(value) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def string(value, label: str, maximum: int = 32767, empty: bool = False) -> str:
    if not isinstance(value, str) or (not empty and not value) or len(value) > maximum:
        problem("WORKBOOK_MODEL_INVALID", label + " must be a text value within its stated limit.")
    if XML_INVALID.search(value):
        problem("WORKBOOK_MODEL_INVALID", label + " contains a character Excel cannot retain in a cell.")
    return value


def integer(value, label: str, minimum: int = 0, maximum: int = MAX_ROWS) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or not minimum <= value <= maximum:
        problem("WORKBOOK_MODEL_INVALID", label + " must be an integer within its stated limit.")
    return value


def decode64(value: Any, maximum: int, label: str) -> bytes:
    if not isinstance(value, str) or len(value) > ((maximum + 2) // 3) * 4:
        problem("WORKBOOK_SIZE", label + " is larger than the supported input.")
    try:
        data = base64.b64decode(value, validate=True)
    except (ValueError, TypeError) as exc:
        problem("WORKBOOK_BASE64_INVALID", label + ": " + str(exc))
    if len(data) > maximum or base64.b64encode(data).decode("ascii") != value:
        problem("WORKBOOK_BASE64_INVALID", label + " is not canonical base64 within the supported size.")
    return data


def model_validate(model):
    if not isinstance(model, dict) or model.get("schema") != MODEL_SCHEMA:
        problem("WORKBOOK_MODEL_INVALID", "Select an original supplier-review workbook model.")
    for name in ("workbookId", "requestId", "comparisonId", "baselineEditionId", "candidateEditionId"):
        string(model.get(name), name, 4096)
    integer(model.get("revision"), "revision", maximum=2**53 - 1)
    for name in ("basisText", "basisSha256"):
        if not isinstance(model.get(name), str):
            problem("WORKBOOK_MODEL_INVALID", name + " is missing.")
    basis = model["basisText"].encode("utf-8")
    if len(basis) > MAX_JSON_BYTES or sha(basis) != model["basisSha256"]:
        problem("WORKBOOK_BASIS_CHANGED", "The complete original workbook basis does not match its hash.")
    basis_value = parse_json(model["basisText"])
    if not isinstance(basis_value, dict):
        problem("WORKBOOK_BASIS_CHANGED", "The complete original workbook basis is not an object.")
    for name in ("workbookId", "requestId", "comparisonId", "revision"):
        if basis_value.get(name) != model[name]:
            problem("WORKBOOK_BASIS_CHANGED", "The original workbook basis selects another " + name + ".")
    if basis_value.get("schema") != MODEL_SCHEMA:
        problem("WORKBOOK_BASIS_CHANGED", "The original workbook basis uses another schema.")
    items = model.get("items")
    fields = model.get("fields")
    replacements = model.get("replacementRows")
    sources = model.get("sourceRows")
    if not all(isinstance(value, list) for value in (items, fields, replacements, sources)):
        problem("WORKBOOK_MODEL_INVALID", "The complete review, field, replacement and source rows are required.")
    if len(items) + REVIEW_HEADER_ROW > MAX_ROWS or len(replacements) + 1 > MAX_ROWS or len(sources) + 1 > MAX_ROWS:
        problem("WORKBOOK_SIZE", "This review exceeds Excel's worksheet row limit. Select a smaller original review edition.")
    ids, field_names = set(), set()
    for item in items:
        if not isinstance(item, dict):
            problem("WORKBOOK_MODEL_INVALID", "Each review item must be an object.")
        item_id = string(item.get("id"), "Item ID", 4096)
        if item_id in ids:
            problem("WORKBOOK_MODEL_INVALID", "Repeated review item: " + item_id)
        ids.add(item_id)
        for name in ("key", "change", "baseline", "candidate", "savedDecision"):
            string(item.get(name, ""), name, 32767, empty=True)
    for field in fields:
        if not isinstance(field, dict):
            problem("WORKBOOK_MODEL_INVALID", "Each canonical field must be an object.")
        name = string(field.get("name"), "Field", 4096)
        if name in field_names:
            problem("WORKBOOK_MODEL_INVALID", "Repeated canonical field: " + name)
        field_names.add(name)
    for row in replacements:
        if not isinstance(row, dict) or row.get("itemId") not in ids or row.get("field") not in field_names:
            problem("WORKBOOK_MODEL_INVALID", "A replacement row selects an unknown original item or field.")
        integer(row.get("part"), "Replacement part", minimum=1)
        if row.get("kind") not in KINDS or row.get("encoding") not in ENCODINGS:
            problem("WORKBOOK_MODEL_INVALID", "A replacement row has an unsupported value kind or encoding.")
        string(row.get("text"), "Replacement value", CHUNK_CHARS, empty=True)
    for row in sources:
        if not isinstance(row, dict) or row.get("itemId") not in ids:
            problem("WORKBOOK_MODEL_INVALID", "A source row selects an unknown original item.")
        if row.get("side") not in ("baseline", "candidate"):
            problem("WORKBOOK_MODEL_INVALID", "A source row selects an unknown side.")
        integer(row.get("rowNumber"), "Original source row", minimum=1, maximum=2**53 - 1)
        integer(row.get("part"), "Original value part", minimum=1)
        string(row.get("field"), "Original source field", 4096)
        string(row.get("text"), "Original source value", CHUNK_CHARS, empty=True)
        if row.get("encoding") not in ENCODINGS:
            problem("WORKBOOK_MODEL_INVALID", "An original source value has an unsupported encoding.")
    return model


def write_workbook(model) -> bytes:
    model_validate(model)
    try:
        import openpyxl
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill, Protection
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        problem("WORKBOOK_RUNTIME_MISSING", "Use the installed Analysis workbooks environment with openpyxl: " + str(exc))
    workbook = openpyxl.Workbook(write_only=True)
    workbook.properties.title = "Supplier review"
    workbook.properties.subject = model["comparisonId"]
    workbook.properties.creator = "Supplier review"
    font = Font(name="Calibri", size=11, color="202020")
    header_font = Font(name="Calibri", size=11, bold=True, color="202020")
    muted = PatternFill("solid", fgColor="F0F0F0")
    header_fill = PatternFill("solid", fgColor="D8D8D8")
    editable_fill = PatternFill("solid", fgColor="FFFFFF")
    alignment = Alignment(vertical="top", wrap_text=True)

    def sheet(name, widths, freeze):
        ws = workbook.create_sheet(name)
        ws.freeze_panes = freeze
        ws.sheet_view.showGridLines = False
        ws.sheet_format.defaultRowHeight = 18
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        # Original cells are shaded. Sorting and adding value fragments stay
        # available; exact identity validation is performed on every return.
        ws.protection.sheet = False
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
        ws.protection.autoFilter = False
        ws.protection.sort = False
        return ws

    def cell(ws, value="", *, editable=False, heading=False, numeric=False):
        result = WriteOnlyCell(ws, value=value)
        if not numeric:
            result.data_type = "s"
            result.number_format = "@"
        result.font = header_font if heading else font
        result.fill = header_fill if heading else editable_fill if editable else muted
        result.alignment = alignment
        result.protection = Protection(locked=not editable)
        return result

    def append(ws, values, *, editable=(), heading=False, numeric=()):
        ws.append([cell(ws, value, editable=i in editable, heading=heading, numeric=i in numeric)
                   for i, value in enumerate(values)])

    def choices(ws, expression, address):
        validation = DataValidation(type="list", formula1='"' + ",".join(expression) + '"',
                                    allow_blank=True, showErrorMessage=True, showInputMessage=True)
        validation.errorTitle = "Choose a supported value"
        validation.error = "Use a listed value. Leave Decision empty to keep the current review."
        validation.promptTitle = "Supplier review"
        validation.prompt = "The returned file will be reviewed in the original workspace."
        validation.add(address)
        ws.data_validations.append(validation)

    review = sheet("Review", [44, 36, 18, 48, 48, 22, 22, 54, 18, 16], "G11")
    end = max(REVIEW_HEADER_ROW + len(model["items"]), REVIEW_HEADER_ROW + 1)
    review.auto_filter.ref = "A10:J" + str(end)
    choices(review, ACTIONS, "G11:G" + str(end))
    choices(review, ["baseline", "candidate"], "I11:I" + str(end))
    append(review, ["Supplier review", "Enter decisions and reasons in the white cells."], heading=True)
    append(review, ["Comparison", model["comparisonId"]])
    append(review, ["Original request", model["requestId"]])
    append(review, ["Workbook", model["workbookId"]])
    append(review, ["Original revision", model["revision"], "Items", len(model["items"])], numeric=(1, 3))
    append(review, ["Baseline edition", model["baselineEditionId"]])
    append(review, ["Incoming edition", model["candidateEditionId"]])
    append(review, ["Reviewer", "", "Return note", ""], editable=(1, 3))
    append(review, ["Use replace with the Corrections sheet. Source values contains complete original field values."])
    append(review, REVIEW_HEADERS, heading=True)
    for item in model["items"]:
        append(review, [item["id"], item.get("key", ""), item.get("change", ""),
                        item.get("baseline", ""), item.get("candidate", ""),
                        item.get("savedDecision", ""), "", "", "", ""], editable=(6, 7, 8, 9))

    corrections = sheet("Corrections", [44, 28, 10, 18, 14, 88], "F2")
    corrections.auto_filter.ref = "A1:F" + str(max(2, 1 + len(model["replacementRows"])))
    choice_end = max(2, 1 + len(model["replacementRows"]))
    choices(corrections, KINDS, "D2:D" + str(choice_end))
    choices(corrections, ENCODINGS, "E2:E" + str(choice_end))
    append(corrections, VALUE_HEADERS, heading=True)
    for row in model["replacementRows"]:
        append(corrections, [row["itemId"], row["field"], str(row["part"]),
                             row["kind"], row["encoding"], row["text"]], editable=(2, 3, 4, 5))

    original = sheet("Source values", [44, 18, 16, 28, 10, 14, 88], "G2")
    original.auto_filter.ref = "A1:G" + str(max(2, 1 + len(model["sourceRows"])))
    append(original, SOURCE_HEADERS, heading=True)
    for row in model["sourceRows"]:
        append(original, [row["itemId"], row["side"], str(row["rowNumber"]), row["field"],
                          str(row["part"]), row["encoding"], row["text"]])

    basis = sheet("Basis", [16, 100, 72], "A3")
    basis.sheet_state = "hidden"
    encoded = base64.b64encode(model["basisText"].encode("utf-8")).decode("ascii")
    chunks = [encoded[offset:offset + CHUNK_CHARS] for offset in range(0, len(encoded), CHUNK_CHARS)]
    append(basis, [MODEL_SCHEMA, model["workbookId"], model["basisSha256"]])
    append(basis, ["Chunks", len(chunks)], numeric=(1,))
    for index, chunk in enumerate(chunks, 1):
        append(basis, [index, chunk], numeric=(0,))

    guide = sheet("Read me", [30, 110], "A2")
    append(guide, ["Supplier review", "How to return a complete review"], heading=True)
    for title, content in [
        ("Decisions", "Leave Decision empty to keep the current review. Enter one supported decision and a reason for each change you want to return."),
        ("Reviewer", "Enter your name in Review B8. The original engine records this attribution when the reviewed return is applied."),
        ("Original values", "Source values contains the exact original field text. Long values use ordered Part rows; join their text in order. Encoding json means the cell contains an exact JSON value."),
        ("Source choice", "For choose_source, enter baseline or candidate and the exact Source row from Source values. This resolves an original ambiguous row without inventing another record."),
        ("Corrections", "For replace, edit the item's complete canonical fields in Corrections. The original rules validate all fields together. A decimal value stays text when the original engine represents it as an exact decimal string."),
        ("Missing and blank", "Value kind null with an empty value means explicit null. Value kind text with an empty value means a blank string. They remain distinct."),
        ("Precision", "Value cells use Excel Text format. Keep exact decimal and identifier digits as text. Numeric cells in returned decisions or values are rejected to prevent Excel's 15-digit rounding from becoming a source correction."),
        ("Long values", "Part numbers start at 1 and must be contiguous per item and field. You may insert more rows for one long value, retaining the item/field and using the same kind and encoding. Each part holds at most 24,000 characters."),
        ("Complete return", "Keep every original Review item and canonical Corrections field. Sorting and filtering are supported. Do not rename the original sheets or identity columns."),
        ("Current review", "Upload this workbook in the original supplier workspace. The application compares returned decisions with the exact export basis and current saved decisions. Resolve any overlapping changes before applying."),
        ("Saved result", "Applying creates a complete original supplier workspace through decideMany. Original source, rules and review history remain in that result. Finalized CSV/SQL/HTML use the original finalizer."),
        ("Cell formulas", "Paste intended values as text. Formula cells, macros, embedded objects and external workbook links are not review values and are rejected on return."),
        ("Source", model.get("sourceDescription", "Original source editions and full workspace are retained in this workbook's Basis sheet.")),
    ]:
        append(guide, [title, content])
    output = io.BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    data = output.getvalue()
    if len(data) > MAX_XLSX_BYTES:
        problem("WORKBOOK_SIZE", "The complete XLSX is larger than the supported return size.")
    return data


def inspect_archive(data: bytes):
    if not data or len(data) > MAX_XLSX_BYTES:
        problem("WORKBOOK_SIZE", "Choose a nonempty XLSX within 64 MiB.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError):
        problem("WORKBOOK_ARCHIVE_INVALID", "The selected file is not an XLSX archive.")
    with archive:
        members = archive.infolist()
        if len(members) > MAX_MEMBERS:
            problem("WORKBOOK_SIZE", "The XLSX contains too many archive members.")
        names, total = set(), 0
        for member in members:
            name = member.filename
            if name in names or name.startswith("/") or "\\" in name or any(part in ("", ".", "..") for part in name.rstrip("/").split("/")):
                problem("WORKBOOK_ARCHIVE_INVALID", "The XLSX has repeated or invalid member paths.")
            names.add(name)
            mode = member.external_attr >> 16
            if stat.S_ISLNK(mode) or member.flag_bits & 1:
                problem("WORKBOOK_ARCHIVE_INVALID", "Encrypted members and symbolic links are not supported in a review XLSX.")
            if member.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                problem("WORKBOOK_ARCHIVE_INVALID", "The XLSX uses an unsupported compression method.")
            total += member.file_size
            if total > MAX_EXPANDED_BYTES:
                problem("WORKBOOK_SIZE", "The expanded XLSX exceeds 256 MiB.")
            lower = name.lower()
            if "vbaproject" in lower or lower.startswith(("xl/embeddings/", "xl/externallinks/")):
                problem("WORKBOOK_ACTIVE_CONTENT", "Return review values without macros, embedded objects or external workbook links.")
            if lower.endswith((".xml", ".rels")):
                payload = archive.read(member)
                upper = payload.upper()
                if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                    problem("WORKBOOK_ARCHIVE_INVALID", "XML document types and entities are not part of this review format.")
                if lower.endswith(".rels"):
                    try:
                        relationships = ET.fromstring(payload)
                    except ET.ParseError:
                        problem("WORKBOOK_ARCHIVE_INVALID", "An XLSX relationship document is invalid.")
                    for relation in relationships:
                        if relation.get("TargetMode") == "External":
                            target = relation.get("Target", "")
                            kind = relation.get("Type", "")
                            if not kind.endswith("/hyperlink") or not re.match(r"^(https?://|mailto:)", target, re.I):
                                problem("WORKBOOK_ACTIVE_CONTENT", "External data relationships are not review values.")
        if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
            problem("WORKBOOK_ARCHIVE_INVALID", "The selected archive is not an Excel workbook.")


def cell_text(cell, label: str, *, blank: bool = True, numeric: bool = False) -> str:
    value = cell.value
    if value is None:
        if blank:
            return ""
        problem("WORKBOOK_CELL_REQUIRED", label + " is required.")
    if cell.data_type == "f":
        problem("WORKBOOK_FORMULA", label + " contains a formula. Paste its intended value as text.")
    if numeric and isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if not isinstance(value, str):
        problem("WORKBOOK_VALUE_TYPE", label + " must be text. Format the cell as Text and re-enter the exact value.")
    if len(value) > 32767 or XML_INVALID.search(value):
        problem("WORKBOOK_VALUE_TYPE", label + " contains an unsupported cell value.")
    return value


def sheet_rows(workbook, name, columns, start=1):
    if name not in workbook.sheetnames:
        problem("WORKBOOK_SHEET_MISSING", "The returned workbook is missing " + name + ".")
    ws = workbook[name]
    # Do not trust producer-supplied used-range dimensions to bound original data.
    ws.reset_dimensions()
    for index, cells in enumerate(ws.iter_rows(min_row=start), start):
        if index > MAX_ROWS:
            problem("WORKBOOK_SIZE", name + " exceeds Excel's row limit.")
        if len(cells) > columns and any(cell.value is not None for cell in cells[columns:]):
            problem("WORKBOOK_COLUMNS_CHANGED", name + " has data outside its original columns.")
        padded = list(cells[:columns])
        if len(padded) < columns:
            # EmptyCell is the public read-only representation of absent cells.
            from openpyxl.cell.read_only import EmptyCell
            padded.extend(EmptyCell() for _ in range(columns - len(padded)))
        yield index, padded


def require_header(cells, expected, sheet_name):
    actual = [cell_text(cell, sheet_name + " header", blank=False) for cell in cells]
    if actual != expected:
        problem("WORKBOOK_COLUMNS_CHANGED", sheet_name + " headers differ from the original review format.")


def read_workbook(data: bytes):
    inspect_archive(data)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ImportError as exc:
        problem("WORKBOOK_RUNTIME_MISSING", "Use the installed Analysis workbooks environment with openpyxl: " + str(exc))
    except (ValueError, KeyError, TypeError, OSError, zipfile.BadZipFile) as exc:
        problem("WORKBOOK_ARCHIVE_INVALID", str(exc))
    try:
        basis_rows = iter(sheet_rows(workbook, "Basis", 3))
        try:
            _, first = next(basis_rows)
            _, second = next(basis_rows)
        except StopIteration:
            problem("WORKBOOK_BASIS_CHANGED", "The complete original workbook basis is missing.")
        head = [cell_text(cell, "Basis identity", blank=False) for cell in first]
        if head[0] != MODEL_SCHEMA or not re.fullmatch(r"[0-9a-f]{64}", head[2]):
            problem("WORKBOOK_BASIS_CHANGED", "The workbook basis identity is invalid.")
        if cell_text(second[0], "Basis chunks") != "Chunks":
            problem("WORKBOOK_BASIS_CHANGED", "The workbook basis chunk declaration is invalid.")
        count_text = cell_text(second[1], "Basis chunk count", numeric=True, blank=False)
        if not re.fullmatch(r"[1-9][0-9]*", count_text):
            problem("WORKBOOK_BASIS_CHANGED", "The workbook basis chunk count is invalid.")
        count = int(count_text)
        if count > ((MAX_JSON_BYTES + 2) // 3 * 4 + CHUNK_CHARS - 1) // CHUNK_CHARS:
            problem("WORKBOOK_SIZE", "The complete workbook basis is too large.")
        chunks = []
        for _, row in basis_rows:
            if all(cell.value is None for cell in row):
                continue
            part = cell_text(row[0], "Basis part", numeric=True, blank=False)
            if part != str(len(chunks) + 1) or len(chunks) >= count:
                problem("WORKBOOK_BASIS_CHANGED", "The workbook basis is incomplete or reordered.")
            chunks.append(cell_text(row[1], "Basis part text", blank=False))
            if row[2].value is not None:
                problem("WORKBOOK_BASIS_CHANGED", "The workbook basis has an unexpected value.")
        if len(chunks) != count:
            problem("WORKBOOK_BASIS_CHANGED", "The workbook basis is incomplete.")
        basis_data = decode64("".join(chunks), MAX_JSON_BYTES, "Original workbook basis")
        if sha(basis_data) != head[2]:
            problem("WORKBOOK_BASIS_CHANGED", "The returned workbook basis differs from its original hash.")
        try:
            basis_text = basis_data.decode("utf-8", errors="strict")
        except UnicodeError:
            problem("WORKBOOK_BASIS_CHANGED", "The workbook basis is not UTF-8.")
        basis = parse_json(basis_text)
        if not isinstance(basis, dict) or basis.get("schema") != MODEL_SCHEMA or basis.get("workbookId") != head[1]:
            problem("WORKBOOK_BASIS_CHANGED", "The original workbook identity does not match its complete basis.")
        rows, reviewer, note, overview = [], "", "", {}
        header_seen = False
        for index, cells in sheet_rows(workbook, "Review", len(REVIEW_HEADERS)):
            if index < REVIEW_HEADER_ROW:
                if index == 8:
                    if cell_text(cells[0], "Reviewer label") != "Reviewer" or cell_text(cells[2], "Return note label") != "Return note":
                        problem("WORKBOOK_COLUMNS_CHANGED", "The reviewer and return-note fields have moved.")
                    reviewer = cell_text(cells[1], "Reviewer")
                    note = cell_text(cells[3], "Return note")
                elif index in (2, 3, 4, 5, 6, 7):
                    overview[index] = cell_text(cells[1], "Review identity", numeric=index == 5, blank=False)
                continue
            if index == REVIEW_HEADER_ROW:
                require_header(cells, REVIEW_HEADERS, "Review")
                header_seen = True
                continue
            if all(cell.value is None for cell in cells):
                continue
            values = [cell_text(cell, "Review " + str(index) + " / " + REVIEW_HEADERS[i])
                      for i, cell in enumerate(cells)]
            rows.append(dict(zip(("itemId", "key", "change", "baseline", "candidate", "savedDecision",
                                  "action", "reason", "side", "rowNumberText"), values)))
        if not header_seen:
            problem("WORKBOOK_COLUMNS_CHANGED", "The original Review header is missing.")
        expected_overview = {2: basis.get("comparisonId"), 3: basis.get("requestId"), 4: head[1],
                             5: str(basis.get("revision")), 6: basis.get("baselineEditionId"),
                             7: basis.get("candidateEditionId")}
        if overview != expected_overview:
            problem("WORKBOOK_BASIS_CHANGED", "Visible comparison and source identities differ from the retained workbook basis.")
        correction_rows = []
        header_seen = False
        for index, cells in sheet_rows(workbook, "Corrections", len(VALUE_HEADERS)):
            if index == 1:
                require_header(cells, VALUE_HEADERS, "Corrections")
                header_seen = True
                continue
            if all(cell.value is None for cell in cells):
                continue
            values = [cell_text(cell, "Corrections " + str(index) + " / " + VALUE_HEADERS[i])
                      for i, cell in enumerate(cells)]
            correction_rows.append(dict(zip(("itemId", "field", "partText", "kind", "encoding", "text"), values)))
        if not header_seen:
            problem("WORKBOOK_COLUMNS_CHANGED", "The original Corrections header is missing.")
        # Source values are useful working evidence but remain read-only. The
        # JS layer compares their complete values with its original model.
        source_rows = []
        header_seen = False
        for index, cells in sheet_rows(workbook, "Source values", len(SOURCE_HEADERS)):
            if index == 1:
                require_header(cells, SOURCE_HEADERS, "Source values")
                header_seen = True
                continue
            if all(cell.value is None for cell in cells):
                continue
            values = [cell_text(cell, "Source values " + str(index) + " / " + SOURCE_HEADERS[i])
                      for i, cell in enumerate(cells)]
            source_rows.append(dict(zip(("itemId", "side", "rowNumberText", "field", "partText", "encoding", "text"), values)))
        if not header_seen:
            problem("WORKBOOK_COLUMNS_CHANGED", "The original Source values header is missing.")
        return {"schema": RETURN_SCHEMA, "workbookId": head[1], "basisSha256": head[2],
                "basisText": basis_text, "reviewer": reviewer, "note": note,
                "rows": rows, "replacementRows": correction_rows, "sourceRows": source_rows,
                "file": {"bytes": len(data), "sha256": sha(data)}}
    finally:
        workbook.close()


def main():
    body = sys.stdin.buffer.read(MAX_JSON_BYTES + 1)
    if len(body) > MAX_JSON_BYTES:
        problem("WORKBOOK_SIZE", "Workbook conversion input exceeds 96 MiB.")
    try:
        request = parse_json(body.decode("utf-8", errors="strict"))
    except UnicodeError:
        problem("WORKBOOK_JSON_INVALID", "Workbook conversion input must be UTF-8.")
    if not isinstance(request, dict) or request.get("schema") != IO_SCHEMA:
        problem("WORKBOOK_MODEL_INVALID", "Select the original workbook conversion protocol.")
    if request.get("action") == "write":
        data = write_workbook(request.get("model"))
        result = {"schema": IO_SCHEMA, "action": "write", "file": {
            "name": "Supplier_Review.xlsx",
            "contentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "bytes": len(data), "sha256": sha(data),
            "contentBase64": base64.b64encode(data).decode("ascii")}}
    elif request.get("action") == "read":
        data = decode64(request.get("contentBase64"), MAX_XLSX_BYTES, "Returned workbook")
        expected = request.get("sha256")
        if expected is not None and expected != sha(data):
            problem("WORKBOOK_FILE_CHANGED", "The selected return differs from the retained file hash.")
        result = read_workbook(data)
    else:
        problem("WORKBOOK_MODEL_INVALID", "Choose write or read for workbook conversion.")
    sys.stdout.write(json_text(result) + "\n")


if __name__ == "__main__":
    try:
        main()
    except WorkbookProblem as exc:
        sys.stderr.write(json_text({"error": exc.code, "message": str(exc)}) + "\n")
        sys.exit(2)
    except Exception as exc:
        sys.stderr.write(json_text({"error": "WORKBOOK_IO_FAILED", "message": str(exc)}) + "\n")
        sys.exit(1)
